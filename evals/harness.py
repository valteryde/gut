"""Gut eval harness — runs the real agent loop headless against a local
OpenAI-compatible model and scores each run mechanically.

The daemon is imported in-process: pyautogui is stubbed, every screen,
browser and desktop tool answers "no screen here", cleanup is off, and the
LLM endpoint is pointed straight at the local server (LM Studio, vLLM,
llama.cpp, Ollama's /v1) — no docker, no LiteLLM, no desktop. Everything
else is the production code path: plan lint, evidence gate, workers,
footers, digests, the wrap-up checker.

    evals/.venv/bin/python evals/harness.py                 # whole suite, both modes
    evals/.venv/bin/python evals/harness.py --only single-price --mode orchestrate
    evals/.venv/bin/python evals/harness.py --llm http://192.168.50.63:1234 \
        --model qwen/qwen3-vl-8b --repeat 2

Results land in evals/results/<stamp>/ as one JSON per run (metrics, the
run's checks, the events and the full message context) plus summary.json;
a table prints at the end. Metrics are mechanical — counts from the tool
ledger and the events, file digests of what was delivered — never an LLM
judging an LLM.

Setup once:  python3.13 -m venv evals/.venv && evals/.venv/bin/pip install
             fastapi httpx pillow websockets openpyxl python-docx
"""
from __future__ import annotations

import argparse
import asyncio
import fnmatch
import json
import os
import re
import shutil
import sys
import tempfile
import time
from pathlib import Path
from unittest.mock import MagicMock

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parent
RESULTS = ROOT / "results"
HEADLESS_NOTE = ("(no screen, browser or desktop in this environment — use "
                 "run_command, or a worker for anything online)")
# Tools that need the X desktop or Chrome. Answered with HEADLESS_NOTE so
# the loop stays honest about what it can and cannot do here.
HEADLESS_TOOLS = {
    "screenshot", "wait", "click", "mouse_move", "scroll", "type_text",
    "key", "open_url", "browser_text", "browser_dom", "browser_click",
    "browser_type", "browser_eval", "desktop_tree", "desktop_act",
    "desktop_click", "desktop_type", "office_eval", "list_windows",
    "focus_window", "send_image",
}


def parse_args() -> argparse.Namespace:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--llm", default=os.environ.get(
        "EVAL_LLM", "http://192.168.50.63:1234"),
        help="OpenAI-compatible server base (with or without /v1)")
    ap.add_argument("--model", default=os.environ.get(
        "EVAL_MODEL", "qwen/qwen3-vl-8b"))
    ap.add_argument("--mode", default="both",
                    choices=["orchestrate", "classic", "both"])
    ap.add_argument("--tasks", default=str(ROOT / "tasks.json"))
    ap.add_argument("--only", nargs="*", default=None,
                    help="task ids to run (default: all)")
    ap.add_argument("--repeat", type=int, default=1)
    ap.add_argument("--max-steps", type=int, default=60,
                    help="AGENT_MAX_STEPS for the main loop")
    ap.add_argument("--worker-steps", type=int, default=25,
                    help="SUBAGENT_MAX_STEPS per worker")
    ap.add_argument("--timeout", type=int, default=1200,
                    help="wall-clock cap per run, seconds")
    ap.add_argument("--ctx", default="32k",
                    help="MODEL_CONTEXT_LIMITS cap for every model, e.g. 32k")
    ap.add_argument("--openserp", default=os.environ.get("OPENSERP_URL", ""),
                    help="openserp base for web_search (default: DDG fallback)")
    ap.add_argument("--keep", action="store_true",
                    help="keep the per-run home dirs (default: copy sent "
                         "files into results and delete)")
    ap.add_argument("--verbose", "-v", action="store_true",
                    help="print tool calls live")
    return ap.parse_args()


# ── daemon import (env must be set before) ─────────────────────────────────

def import_daemon(args: argparse.Namespace):
    base = args.llm.rstrip("/")
    if base.endswith("/v1"):
        base = base[:-3]
    eval_root = Path(tempfile.mkdtemp(prefix="gut-eval-"))
    os.environ.update({
        "HOME": str(eval_root / "home"),
        "GUT_DATA_DIR": str(eval_root / "data"),
        "LITELLM_URL": base,
        "LITELLM_MASTER_KEY": "sk-eval",
        "DEFAULT_MODEL": args.model,
        "GUT_CLEANUP": "off",
        "AGENT_MAX_STEPS": str(args.max_steps),
        "SUBAGENT_MAX_STEPS": str(args.worker_steps),
        "MODEL_CONTEXT_LIMITS": f"*={args.ctx}" if args.ctx else "",
        "OPENSERP_URL": args.openserp,
        "ACTION_SETTLE_SECS": "0",
        "LLM_MAX_RETRIES": "3",
        "ASK_USER_TIMEOUT": "1",
    })
    (eval_root / "home").mkdir(parents=True)
    # The agent's python3 must have openpyxl/docx — the venv's does.
    venv_bin = ROOT / ".venv" / "bin"
    if venv_bin.is_dir():
        os.environ["PATH"] = f"{venv_bin}:{os.environ.get('PATH', '')}"
    pag = MagicMock()
    pag.size.return_value = (1024, 768)
    sys.modules["pyautogui"] = pag
    sys.path.insert(0, str(REPO / "backend"))
    import agent_daemon as ad  # noqa: E402
    return ad, eval_root


def patch_headless(ad, events: list, verbose: bool) -> None:
    """Replace everything that needs a display with honest stubs and tap
    the broadcast channel so the run's events can be scored."""
    ad.screenshot_block = lambda force=False: (
        {"type": "text", "text": "(headless eval: no screen)"}
        if force else None)

    def _no_frame(*a, **k):
        raise RuntimeError("no screen in this environment")
    ad.capture_frame = _no_frame
    ad.capture_baseline = lambda: {}

    async def _no_cleanup(*a, **k):
        return None
    ad.cleanup_after_run = _no_cleanup

    async def _no_spend():
        return None
    ad.litellm_spend = _no_spend

    async def _ask_user(question: str) -> str:
        events.append({"type": "question", "text": question})
        return ("(eval: no human available — proceed with your best "
                "judgment, flag assumptions in the deliverable)")
    ad.ask_user = _ask_user

    real_execute = ad.execute_tool

    async def execute_tool(name, args, agent=None):
        if name in HEADLESS_TOOLS:
            return HEADLESS_NOTE, False
        return await real_execute(name, args, agent=agent)
    ad.execute_tool = execute_tool

    # LLM call counters — the message context is compacted on long runs,
    # so "assistant messages" undercounts turns; count requests instead.
    real_llm = ad.llm_request
    real_forcing = ad.llm_request_forcing
    calls = {"total": 0, "main": 0}

    async def llm_request(*a, **k):
        calls["total"] += 1
        return await real_llm(*a, **k)

    async def llm_request_forcing(*a, **k):
        calls["main"] += 1
        return await real_forcing(*a, **k)
    ad.llm_request = llm_request
    ad.llm_request_forcing = llm_request_forcing
    ad.eval_calls = calls

    real_broadcast = ad.broadcast

    async def broadcast(msg: dict):
        keep = {k: v for k, v in msg.items() if k != "data"}
        keep["t"] = time.time()
        events.append(keep)
        if verbose and msg.get("type") in ("action", "agent_msg", "done",
                                           "error", "plan", "subagent",
                                           "question"):
            who = f"[{msg['agent']}] " if msg.get("agent") else ""
            body = msg.get("tool") or msg.get("state") or ""
            text = (json.dumps(msg.get("args"))[:160] if msg.get("args")
                    else str(msg.get("text") or msg.get("result")
                             or msg.get("task") or "")[:160])
            print(f"    {who}{msg['type']:<10} {body:<14} {text}",
                  flush=True)
        await real_broadcast(msg)
    ad.broadcast = broadcast


# ── one run ────────────────────────────────────────────────────────────────

async def run_task(ad, task: dict, mode: str, args, events: list,
                   home: Path) -> dict:
    ad.state = ad.AgentState()
    ad.state.model = args.model
    ad.state.litellm_key = "sk-eval"
    ad.AGENT_ORCHESTRATE = mode == "orchestrate"
    ad.MAX_STEPS = args.max_steps
    ad.SUBAGENT_MAX_STEPS = args.worker_steps
    for d in ("scratch", "uploads", "Desktop", "Downloads"):
        (home / d).mkdir(parents=True, exist_ok=True)
    os.environ["HOME"] = str(home)
    ad.HOME_DIR = home
    ad.SCRATCH_DIR = home / "scratch"
    ad.UPLOAD_DIR = home / "uploads"
    events.clear()
    ad.eval_calls.update(total=0, main=0)
    cid = ad.conv_create(args.model, task["id"])["id"]
    ad.state.conversation_id = cid
    t0 = time.time()
    loop_task = asyncio.create_task(ad.agent_loop(cid, task["text"]))
    ad.state.task = loop_task
    timed_out = False
    try:
        await asyncio.wait_for(asyncio.shield(loop_task), args.timeout)
    except asyncio.TimeoutError:
        timed_out = True
        ad.state.stop = True
        try:
            await asyncio.wait_for(loop_task, 60)
        except (asyncio.TimeoutError, Exception):
            loop_task.cancel()
    except Exception as e:  # a crash inside the loop is a result too
        events.append({"type": "error", "text": f"harness: {e}"})
    wall = time.time() - t0
    messages = ad.conv_load_context(cid) or []
    m = metrics(ad, messages, events, task, mode, home)
    m.update(task=task["id"], mode=mode, model=args.model, wall_s=round(wall),
             timed_out=timed_out, tokens_in=ad.state.tokens_in,
             tokens_out=ad.state.tokens_out,
             steps=ad.eval_calls["main"], llm_calls=ad.eval_calls["total"],
             compactions=sum(1 for e in events if e.get("type") == "compact"),
             sent_paths=sorted(ad.state.sent_files))
    m["checks"] = run_checks(ad, task, m, home, events)
    m["score"] = (sum(1 for c in m["checks"] if c["ok"]), len(m["checks"]))
    return {"metrics": m, "events": events[:], "messages": messages}


# ── metrics ────────────────────────────────────────────────────────────────

def _text(content) -> str:
    if isinstance(content, str):
        return content
    if isinstance(content, list):
        return " ".join(str(b.get("text", "")) for b in content
                        if isinstance(b, dict) and b.get("type") == "text")
    return ""


def metrics(ad, messages: list, events: list, task: dict, mode: str,
            home: Path) -> dict:
    tool_counts: dict[str, int] = {}
    pending: dict[str, str] = {}
    results: list[tuple[str, str]] = []
    for msg in messages:
        if msg.get("role") == "assistant":
            for tc in msg.get("tool_calls") or []:
                name = (tc.get("function") or {}).get("name", "?")
                pending[tc.get("id")] = name
                tool_counts[name] = tool_counts.get(name, 0) + 1
        elif msg.get("role") == "tool":
            results.append((pending.get(msg.get("tool_call_id"), "?"),
                            _text(msg.get("content"))))
    steps = sum(1 for m in messages if m.get("role") == "assistant")
    res_text = "\n".join(r for _, r in results)
    subs = [e for e in events if e.get("type") == "subagent"]
    spawned = {e["name"] for e in subs}
    finished = {e["name"]: e for e in subs if e.get("state") != "running"}
    footers = [str(e.get("result", "")) for e in finished.values()]
    done_ev = [e for e in events if e.get("type") == "done"]
    done_text = done_ev[-1]["text"] if done_ev else ""
    todos = list(ad.state.todos)
    kinds: dict[str, int] = {}
    for t in todos:
        kinds[t.get("kind", "?")] = kinds.get(t.get("kind", "?"), 0) + 1
    sent = [{"name": e.get("name"), "size": e.get("size")}
            for e in events if e.get("type") == "file"]
    deliverables = []
    for p in sorted(home.iterdir()):
        if p.is_file() and not p.name.startswith("."):
            deliverables.append({"name": p.name,
                                 "digest": ad.file_digest(p)})
    return {
        "completed": bool(done_ev),
        "steps": steps,
        "tool_counts": tool_counts,
        "unknown_tool_calls": sum(1 for _, r in results
                                  if r.startswith("unknown tool:")),
        "headless_tool_calls": sum(1 for _, r in results
                                   if r.startswith(HEADLESS_NOTE[:20])),
        "blocked_calls": sum(1 for _, r in results
                             if r.startswith("BLOCKED:")),
        "idle_nudges": sum(1 for m in messages if m.get("role") == "user"
                           and isinstance(m.get("content"), str)
                           and m["content"].startswith(
                               ("Continue with tool calls",
                                "Plain text doesn't reach"))),
        "plan_posted": any(e.get("type") == "plan" for e in events),
        "plan_steps": len(todos),
        "plan_kinds": kinds,
        "plan_done_ratio": (round(sum(1 for t in todos
                                      if t.get("status") == "done")
                                  / len(todos), 2) if todos else None),
        "lint_bounces": res_text.count("plan NOT accepted"),
        "gate_holds": res_text.count("is NOT done — no"),
        "workers_spawned": len(spawned),
        "workers_done": sum(1 for e in finished.values()
                            if e.get("state") == "done"),
        "workers_failed": sum(1 for e in finished.values()
                              if e.get("state") in ("error", "stopped")),
        "worker_steps": sum(int(e.get("steps") or 0)
                            for e in finished.values()),
        "worker_unverified_footers": sum(
            1 for f in footers if "no page was fetched" in f
            or "no tools were used" in f),
        "verify_rejects": res_text.count("VERIFICATION FAILED"),
        "verify_caveat": "Checker flagged" in done_text,
        "errors": [e.get("text", "")[:200] for e in events
                   if e.get("type") == "error"],
        "files_sent": sent,
        "deliverables": deliverables,
        "done_text": done_text[:3000],
    }


# ── task checks ────────────────────────────────────────────────────────────

def _file_text(ad, p: Path) -> str:
    """Everything readable in a deliverable, for substring checks."""
    suf = p.suffix.lower()
    try:
        if suf in (".xlsx", ".xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
            return "\n".join(" ".join(str(c) for c in r if c is not None)
                             for ws in wb.worksheets
                             for r in ws.iter_rows(values_only=True))
        if suf == ".docx":
            import docx
            return "\n".join(x.text for x in docx.Document(str(p)).paragraphs)
        return p.read_text(errors="replace")
    except Exception:
        return ""


def _filled_rows(ad, p: Path) -> int:
    suf = p.suffix.lower()
    try:
        if suf in (".xlsx", ".xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
            return sum(1 for ws in wb.worksheets
                       for r in ws.iter_rows(values_only=True)
                       if any(c not in (None, "") for c in r))
        return sum(1 for ln in p.read_text(errors="replace").splitlines()
                   if ln.strip())
    except Exception:
        return 0


def run_checks(ad, task: dict, m: dict, home: Path, events: list) -> list:
    exp = task.get("expect") or {}
    checks: list[dict] = []

    def add(name: str, ok: bool, note: str = ""):
        checks.append({"check": name, "ok": bool(ok), "note": note})

    add("completed", m["completed"] and not m["timed_out"],
        "no task_complete" if not m["completed"] else "")
    sent = [Path(p) for p in m.get("sent_paths", []) if Path(p).is_file()]
    for spec in exp.get("files") or []:
        pat = spec["glob"]
        if spec.get("sent", True):
            cands = [p for p in sent if fnmatch.fnmatch(p.name, pat)]
        else:
            cands = [p for p in home.rglob(pat) if p.is_file()]
        if not cands:
            add(f"file {pat}" + (" sent" if spec.get("sent", True) else ""),
                False, "no matching file" + (" was sent" if spec.get(
                    "sent", True) else ""))
            continue
        best = max(cands, key=lambda p: _filled_rows(ad, p))
        rows = _filled_rows(ad, best)
        if "min_rows" in spec:
            add(f"{best.name} rows ≥ {spec['min_rows']}",
                rows >= spec["min_rows"], f"{rows} filled rows")
        text = _file_text(ad, best).lower()
        if spec.get("contains_any"):
            hit = [n for n in spec["contains_any"] if n.lower() in text]
            add(f"{best.name} mentions {spec['contains_any']}", bool(hit),
                f"found {hit}" if hit else "none found")
        if spec.get("contains_url"):
            n = len(re.findall(r"https?://", text))
            add(f"{best.name} has source URLs", n > 0, f"{n} URLs")
    dt = (m["done_text"] or "").lower()
    if exp.get("summary_any"):
        hit = [n for n in exp["summary_any"] if n.lower() in dt]
        add(f"summary mentions {exp['summary_any']}", bool(hit),
            f"found {hit}" if hit else "none found")
    if exp.get("summary_url"):
        add("summary cites a URL", bool(re.search(r"https?://", dt)))
    if exp.get("summary_digits"):
        add("summary has a figure", bool(re.search(r"\d", dt)))
    if "min_workers" in exp and m["mode"] == "orchestrate":
        add(f"workers ≥ {exp['min_workers']}",
            m["workers_spawned"] >= exp["min_workers"],
            f"{m['workers_spawned']} spawned")
    if "max_workers" in exp:
        add(f"workers ≤ {exp['max_workers']}",
            m["workers_spawned"] <= exp["max_workers"],
            f"{m['workers_spawned']} spawned")
    if exp.get("plan"):
        add("checklist posted", m["plan_steps"] > 0,
            f"{m['plan_steps']} steps, kinds {m['plan_kinds']}"
            + ("" if m["plan_posted"] else " (no summary card)"))
        add("checklist all done", m["plan_done_ratio"] == 1.0,
            f"{m['plan_done_ratio']}")
    add("no unverified-claim caveat", not m["verify_caveat"])
    add("no hallucinated tools", m["unknown_tool_calls"] == 0,
        f"{m['unknown_tool_calls']} unknown")
    if "max_steps" in exp:
        add(f"steps ≤ {exp['max_steps']}", m["steps"] <= exp["max_steps"],
            f"{m['steps']} steps")
    return checks


# ── driver ─────────────────────────────────────────────────────────────────

def fmt_row(m: dict) -> str:
    ok, n = m["score"]
    return (f"{m['task']:<18} {m['mode']:<11} {ok}/{n:<4} "
            f"{'done' if m['completed'] else 'FAIL':<5} "
            f"st={m['steps']:<3} calls={m['llm_calls']:<3} "
            f"wk={m['workers_spawned']}/{m['workers_done']:<3} "
            f"lint={m['lint_bounces']} hold={m['gate_holds']} "
            f"vfy={m['verify_rejects']} cmp={m['compactions']} "
            f"tok={m['tokens_in'] // 1000}k/{m['tokens_out'] // 1000}k "
            f"{m['wall_s']}s")


async def main() -> int:
    args = parse_args()
    tasks = json.loads(Path(args.tasks).read_text())
    if args.only:
        tasks = [t for t in tasks if t["id"] in set(args.only)]
        missing = set(args.only) - {t["id"] for t in tasks}
        if missing:
            print(f"unknown task ids: {sorted(missing)}", file=sys.stderr)
            return 2
    modes = ["orchestrate", "classic"] if args.mode == "both" else [args.mode]
    ad, eval_root = import_daemon(args)
    events: list = []
    patch_headless(ad, events, args.verbose)
    stamp = time.strftime("%Y%m%d-%H%M%S")
    out = RESULTS / stamp
    out.mkdir(parents=True)
    print(f"model {args.model} @ {args.llm}  →  {out}")
    rows = []
    for task in tasks:
        for mode in modes:
            for i in range(args.repeat):
                tag = f"{task['id']}-{mode}-{i + 1}"
                home = eval_root / "runs" / tag
                home.mkdir(parents=True)
                print(f"\n▶ {tag}: {task['text'][:90]}…", flush=True)
                res = await run_task(ad, task, mode, args, events, home)
                m = res["metrics"]
                rows.append(m)
                (out / f"{tag}.json").write_text(
                    json.dumps(res, indent=1, default=str))
                keep = out / tag
                keep.mkdir(exist_ok=True)
                for p in home.iterdir():
                    if p.is_file() and not p.name.startswith("."):
                        shutil.copy(p, keep / p.name)
                sc = home / "scratch"
                if sc.is_dir():
                    shutil.copytree(sc, keep / "scratch", dirs_exist_ok=True)
                for sp in m.get("sent_paths", []):
                    p = Path(sp)
                    if p.is_file() and keep not in p.parents:
                        (keep / "sent").mkdir(exist_ok=True)
                        shutil.copy(p, keep / "sent" / p.name)
                print("  " + fmt_row(m))
                for c in m["checks"]:
                    print(f"    {'✓' if c['ok'] else '✗'} {c['check']}"
                          + (f" — {c['note']}" if c["note"] else ""))
                if m["errors"]:
                    print(f"    errors: {m['errors'][:3]}")
                if not args.keep:
                    shutil.rmtree(home, ignore_errors=True)
    (out / "summary.json").write_text(json.dumps(rows, indent=1, default=str))
    print("\n" + "=" * 110)
    print(f"{'task':<18} {'mode':<11} score  ok    steps   calls     "
          f"workers  lint hold vfy cmp tokens        wall")
    for m in rows:
        print(fmt_row(m))
    by_mode: dict[str, list] = {}
    for m in rows:
        by_mode.setdefault(m["mode"], []).append(m)
    print("-" * 110)
    for mode, ms in by_mode.items():
        ok = sum(m["score"][0] for m in ms)
        n = sum(m["score"][1] for m in ms)
        print(f"{mode:<11} checks {ok}/{n}  completed "
              f"{sum(m['completed'] for m in ms)}/{len(ms)}  "
              f"avg steps {sum(m['steps'] for m in ms) / len(ms):.0f}  "
              f"avg tokens_in {sum(m['tokens_in'] for m in ms) // len(ms) // 1000}k  "
              f"avg wall {sum(m['wall_s'] for m in ms) // len(ms)}s")
    if not args.keep:
        shutil.rmtree(eval_root, ignore_errors=True)
    else:
        print(f"run homes kept under {eval_root}")
    return 0


if __name__ == "__main__":
    sys.exit(asyncio.run(main()))
