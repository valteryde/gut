"""Mechanical unit checks for the daemon's orchestration machinery — no
model, no network. Run: evals/.venv/bin/python evals/test_units.py"""
import asyncio
import os
import sys
import tempfile
from pathlib import Path
from unittest.mock import MagicMock

os.environ.setdefault("HOME", tempfile.mkdtemp())
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "backend"))
sys.modules["pyautogui"] = MagicMock()
import agent_daemon as ad  # noqa: E402

FAILS = []


def check(name, cond, extra=""):
    print(("PASS" if cond else "FAIL"), name, f"— {extra}" if extra else "")
    if not cond:
        FAILS.append(name)


def todos(items):
    return asyncio.run(ad.update_todos(items))


# ── step-kind inference ─────────────────────────────────────────────────
check("kind: research default",
      ad.infer_step_kind("Søg efter danske priser") == "research")
check("kind: build wins over research",
      ad.infer_step_kind("Lav et Excel-ark med priserne") == "build")
check("kind: deliver", ad.infer_step_kind("Send filen til brugeren") == "deliver")
check("kind: decide", ad.infer_step_kind("Vælg den billigste butik") == "decide")
check("kind: desktop", ad.infer_step_kind("Click the OK button") == "desktop")

# ── enumeration detection ───────────────────────────────────────────────
t = ad.enumerated_targets(
    "Søg efter danske priser på pulled pork, coleslaw og hjemmelavede "
    "ovnkartofler")
check("enum: 3 Danish items", len(t) >= 3, str(t))
check("enum: single target ok",
      ad.enumerated_targets("Find price of carrots in Denmark") == [])
check("enum: 3 English",
      len(ad.enumerated_targets(
          "Compare prices for Netflix, Disney+ and HBO Max")) >= 3)

# ── checklist lint ──────────────────────────────────────────────────────
ad.state = ad.AgentState()
out = todos([{"content": "Find prices for pork, coleslaw and potatoes",
              "status": "pending", "kind": "research"}])
check("lint: bundled research step rejected", "NOT accepted" in out,
      out[:70])

# ── evidence gate ───────────────────────────────────────────────────────
ad.state = ad.AgentState()
ad.state.todos = [{"content": "Research carrot prices",
                   "status": "in_progress", "kind": "research"}]
ad.state.step_marks["research carrot prices"] = dict(ad.state.evidence)
out = todos([{"content": "Research carrot prices", "status": "done",
              "kind": "research"}])
check("gate: done without evidence bounces",
      "NOT done" in out and ad.state.todos[0]["status"] == "in_progress")

ad.state = ad.AgentState()
ad.state.todos = [{"content": "Research carrot prices",
                   "status": "in_progress", "kind": "research"}]
ad.state.step_marks["research carrot prices"] = dict(ad.state.evidence)
ad.note_evidence("research")   # what a delivered worker report records
todos([{"content": "Research carrot prices", "status": "done",
        "kind": "research"}])
check("gate: done with report evidence passes",
      ad.state.todos[0]["status"] == "done")

ad.state = ad.AgentState()
ad.state.todos = [{"content": "Pick cheapest shop", "status": "in_progress",
                   "kind": "decide"}]
ad.state.step_marks["pick cheapest shop"] = dict(ad.state.evidence)
todos([{"content": "Pick cheapest shop", "status": "done", "kind": "decide"}])
check("gate: decide needs no evidence",
      ad.state.todos[0]["status"] == "done")

# gate exhaustion: after GATE_MAX_BOUNCES the list is accepted as posted
ad.state = ad.AgentState()
ad.state.gate_bounces = ad.GATE_MAX_BOUNCES
ad.state.todos = [{"content": "Research carrot prices",
                   "status": "in_progress", "kind": "research"}]
ad.state.step_marks["research carrot prices"] = dict(ad.state.evidence)
todos([{"content": "Research carrot prices", "status": "done",
        "kind": "research"}])
check("gate: exhausted bounces accepts as posted",
      ad.state.todos[0]["status"] == "done")

# ── worker findings gate ────────────────────────────────────────────────
d = Path(tempfile.mkdtemp())
check("findings: missing dir is empty", ad._findings_empty(d))
(d / "findings.json").write_text("[]")
check("findings: [] is empty", ad._findings_empty(d))
(d / "findings.json").write_text(
    '[{"value": "6,25", "unit": "kr", "source_url": "https://x.dk"}]')
check("findings: records count", not ad._findings_empty(d))

check("claims: price+unit", bool(ad._REPORT_CLAIM_RE.search(
    "6,25 kr/stk from Nemlig")))
check("claims: url", bool(ad._REPORT_CLAIM_RE.search("see https://x.dk")))
check("claims: honest UNVERIFIED prose is not a claim",
      not ad._REPORT_CLAIM_RE.search("No verifiable data. UNVERIFIED."))
check("claims: restating the goal is not a claim",
      not ad._REPORT_CLAIM_RE.search("budget for 220 personer failed"))

# ── file digests ────────────────────────────────────────────────────────
import openpyxl  # noqa: E402

p = Path(tempfile.mkdtemp()) / "t.xlsx"
wb = openpyxl.Workbook()
ws = wb.active
ws.append(["item", "price"])
ws.append(["pork", 42])
wb.save(p)
d = ad.file_digest(p)
check("digest: xlsx shows rows+cells", "rows" in d and "cells" in d, d[:80])
check("digest: empty xlsx shows 0 rows",
      "0" in ad.file_digest(
          (lambda q: (openpyxl.Workbook().save(q), q)[1])(
              Path(tempfile.mkdtemp()) / "e.xlsx")))

# ── run_command signature normalization ─────────────────────────────────
s = ad.StallDetector.signature
a = s("run_command", {"command": 'python3 -c "import openpyxl; wb=openpyxl.Workbook()"'})
b = s("run_command", {"command": "python3 -c  'import  openpyxl; wb = openpyxl.Workbook()'"})
check("stall: whitespace/quote variants are one call", a == b)
check("stall: different scripts differ",
      a != s("run_command", {"command": 'python3 -c "print(1)"'}))
check("stall: searches keep distinct signatures",
      s("web_search", {"queries": ["x"]}) != s("web_search",
                                              {"queries": ["y"]}))

# ── orchestrator tool visibility ─────────────────────────────────────────
ad.AGENT_ORCHESTRATE = True
hidden = {n for n, c in ad._TOOL_VISIBLE.items() if not c()} \
    if hasattr(ad, "_TOOL_VISIBLE") else set()
# fallback: check via the tools_for_run path used at request time
try:
    shown = {t["function"]["name"] for t in ad.tools_for_run(set())}
    check("orchestrate: web_search hidden", "web_search" not in shown)
    check("orchestrate: spawn_agent present", "spawn_agent" in shown)
except AttributeError:
    check("orchestrate: tool hiding", False, "tools_for_run not found")

# ── remember store ──────────────────────────────────────────────────────
ad.state = ad.AgentState()
out = ad.update_notes({"note": "API base is https://api.example.dk/v2"})
check("remember: stores a note",
      "remembered" in out
      and ad.state.notes == ["API base is https://api.example.dk/v2"],
      out[:80])
out = ad.update_notes({"note": "api base is https://api.example.dk/v2"})
check("remember: dedupes case-insensitively",
      "already noted" in out and len(ad.state.notes) == 1)
out = ad.update_notes({"note": "x" * (ad.NOTE_MAX_CHARS + 1)})
check("remember: oversized note rejected",
      "too long" in out and len(ad.state.notes) == 1)
out = ad.update_notes({})
check("remember: empty call lists state",
      "nothing to do" in out and "api.example.dk" in out)
out = ad.update_notes({"forget": "api.example"})
check("remember: forget removes",
      "forgot 1" in out and not ad.state.notes)

# cap: the list refuses note NOTES_MAX+1
ad.state = ad.AgentState()
ad.state.notes = [f"fact {i}" for i in range(ad.NOTES_MAX)]
out = ad.update_notes({"note": "one more"})
check("remember: full list rejects", "full" in out
      and len(ad.state.notes) == ad.NOTES_MAX)

# persistence: notes round-trip through <cid>.notes.json
ad.CONV_DIR = Path(tempfile.mkdtemp())
ad.state = ad.AgentState()
ad.state.conversation_id = "testconv123"
ad.update_notes({"note": "remember the milk"})
check("remember: persisted to disk",
      ad.conv_load_notes("testconv123") == ["remember the milk"])
check("remember: missing conv loads empty",
      ad.conv_load_notes("nope404") == [])

# ── build-step data-flow nudge ──────────────────────────────────────────
# A build step going in_progress on researched data gets pointed at files:
# worker scratch files when they exist, else "write the data file first".
ad.state = ad.AgentState()
ad.SCRATCH_DIR = Path(tempfile.mkdtemp(dir=ad.HOME_DIR))
ad.state.evidence["research"] = 1
out = todos([{"content": "Research carrot prices", "status": "done",
              "kind": "research"},
             {"content": "Build the comparison spreadsheet",
              "status": "in_progress", "kind": "build"}])
check("build nudge: no files → write data file",
      "Never type figures from memory" in out, out[-120:])

ad.state = ad.AgentState()
ad.state.evidence["research"] = 1
ad.SCRATCH_DIR = Path(tempfile.mkdtemp(dir=ad.HOME_DIR))
wd = ad.SCRATCH_DIR / "price-worker"
wd.mkdir(parents=True, exist_ok=True)
(wd / "findings.json").write_text('[{"value": 42, "unit": "kr"}]')
out = todos([{"content": "Build the comparison spreadsheet",
              "status": "in_progress", "kind": "build"}])
check("build nudge: worker files listed",
      "price-worker/findings.json" in out, out[-160:])

# no research evidence and no files → no nudge
ad.state = ad.AgentState()
ad.SCRATCH_DIR = Path(tempfile.mkdtemp(dir=ad.HOME_DIR))
out = todos([{"content": "Build the comparison spreadsheet",
              "status": "in_progress", "kind": "build"}])
check("build nudge: quiet without research",
      "figures from memory" not in out and "workers' files" not in out)

print(f"\n{len(FAILS)} failures" + (": " + ", ".join(FAILS) if FAILS else ""))
sys.exit(1 if FAILS else 0)
