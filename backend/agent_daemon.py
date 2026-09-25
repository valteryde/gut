"""Gut agent daemon.

FastAPI service inside the desktop container that:
  * exposes a WebSocket chat/control channel to the client UI,
  * drives the local X desktop (PyAutoGUI / scrot / xdotool) from LLM tool calls,
  * routes all model traffic through LiteLLM (BYOK) and reports live cost.
"""

from __future__ import annotations

import asyncio
import base64
import hashlib
from collections import deque
import hmac
import io
import json
import mimetypes
import os
import re
import shlex
import shutil
import signal
import socket
import ssl
import subprocess
import tempfile
import time
from contextlib import asynccontextmanager
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import (parse_qs, urljoin, urlparse, urlsplit,
                          urlunsplit)
from uuid import uuid4

import httpx
import pyautogui
from fastapi import (Body, FastAPI, HTTPException, Request, WebSocket,
                     WebSocketDisconnect)
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
from PIL import Image

try:
    from websockets.asyncio.client import connect as ws_connect
except ImportError:  # websockets<13
    from websockets import connect as ws_connect

try:
    import tcpmux
except ImportError:  # running as a package (uvicorn backend.agent_daemon)
    tcpmux = None

# ── Device config store ──────────────────────────────────────────────────
# Non-secret settings live in config.json on the data volume, managed from
# the app via POST /api/config — .env is only for bootstrap secrets (master
# key, db password, device password). File values override env defaults, the
# same way pushed provider keys win over env keys. Only whitelisted names
# are honoured, and env vars listed here keep working as boot defaults.
CONFIG_FILE = (Path(os.environ.get("GUT_DATA_DIR") or Path.home() / ".gut")
               / "config.json")
# env name -> module global it maps to (same name when absent). Keys in
# CONFIG_RESTART also persist but only take effect at the next boot —
# they feed start.sh (Xvfb size, DPI, ports), not this process.
CONFIG_GLOBALS = {
    "AGENT_MAX_STEPS": "MAX_STEPS",
    "AGENT_IDLE_REPLY_LIMIT": "IDLE_REPLY_LIMIT",
    "AGENT_COMPACT_RATIO": "COMPACT_RATIO",
    "AGENT_CONTEXT_LIMIT": "COMPACT_CONTEXT_LIMIT",
    "AGENT_COMPACT_KEEP": "COMPACT_KEEP",
    "AGENT_COMPACT_INPUT_CHARS": "COMPACT_INPUT_MAX_CHARS",
    "AGENT_CYCLE_WINDOW": "CYCLE_WINDOW",
    "AGENT_RESEARCH_NUDGE": "RESEARCH_NUDGE_STREAK",
    "AGENT_WRAP_UP_STEPS": "WRAP_UP_STEPS",
    "AGENT_VERIFY_MAX_REJECTS": "VERIFY_MAX_REJECTS",
    "AGENT_VERIFY_MIN_CALLS": "VERIFY_MIN_CALLS",
    "AGENT_VERIFY_INPUT_CHARS": "VERIFY_INPUT_CHARS",
    "AGENT_VERIFY_RESULT_CHARS": "VERIFY_RESULT_CHARS",
    "SCREENSHOT_AUTO_PIXELS": "SHOT_AUTO_PIXELS",
    "GUT_UNO_PORT": "UNO_PORT",
}
CONFIG_KEYS = frozenset(CONFIG_GLOBALS) | frozenset({
    "DEFAULT_MODEL", "ESCALATION_MODEL", "ESCALATION_RESCUES",
    "SUBAGENT_MODEL", "SUBAGENT_MAX_STEPS", "SUBAGENT_MAX_CONCURRENT",
    "SUBAGENT_MAX_TOTAL", "AGENT_ORCHESTRATE", "AGENT_GATE_BOUNCES",
    "COMPACT_MODEL", "JANITOR_MODEL", "SCROLL_MAX_CLICKS",
    "MODEL_CONTEXT_LIMITS", "AGENT_VERIFY", "VERIFY_MODEL",
    "AGENT_MAX_USD", "TOOL_RESULT_HISTORY", "TOOL_RESULT_STUB_CHARS",
    "SCREENSHOT_MAX_EDGE", "SCREENSHOT_MAX_PIXELS", "SCREENSHOT_HISTORY",
    "TODO_REMIND_STEPS", "TODO_NUDGE_STEPS", "TODO_MAX_ITEMS",
    "ACTION_SETTLE_SECS", "INTER_ACTION_DELAY",
    "CLICK_A11Y", "CLICK_SNAP_PX", "NORMALIZED_COORD_MODELS",
    "LLM_MAX_RETRIES", "ASK_USER_TIMEOUT", "COMMAND_TIMEOUT",
    "SEND_FILE_MAX_BYTES", "ATTACH_TOTAL_MAX_BYTES",
    "GUT_CLEANUP", "JANITOR_MAX_STEPS", "OPENSERP_URL", "OPENSERP_ENGINES",
    "OPENSERP_MAX_CONCURRENT", "SEARCH_LANG", "SEARCH_REGION",
    # Boot-time settings — persisted here, applied by start.sh next boot.
    "RESOLUTION", "UI_SCALE", "DEVICE_NAME", "WALLPAPER_HUE", "CDP_PORT",
})
CONFIG_RESTART = frozenset({
    "RESOLUTION", "UI_SCALE", "DEVICE_NAME", "WALLPAPER_HUE", "CDP_PORT",
    "GUT_UNO_PORT",
})


def load_config_file() -> dict:
    try:
        data = json.loads(CONFIG_FILE.read_text())
        return {str(k): str(v) for k, v in data.items()
                if str(k) in CONFIG_KEYS and str(v) != ""}
    except (OSError, json.JSONDecodeError, AttributeError):
        return {}


def store_config_file(cfg: dict) -> None:
    CONFIG_FILE.parent.mkdir(parents=True, exist_ok=True)
    tmp = CONFIG_FILE.with_suffix(".tmp")
    tmp.write_text(json.dumps(cfg))
    tmp.chmod(0o600)
    tmp.replace(CONFIG_FILE)


# Values the env held before config.json was layered on, so a key removed
# from the file falls back to the deploy default instead of going blank.
_ENV_ORIG = {k: os.environ.get(k) for k in CONFIG_KEYS}
for _k, _v in load_config_file().items():
    os.environ[_k] = _v


def _coerce_like(old, raw: str):
    """Convert a config string to the type of the existing global."""
    if isinstance(old, bool):
        return str(raw).strip().lower() not in ("off", "0", "false", "no")
    if isinstance(old, int):
        return int(float(raw))
    if isinstance(old, float):
        return float(raw)
    if isinstance(old, list):
        return [p for p in str(raw).lower().split(",") if p]
    return str(raw)


LITELLM_URL = os.environ.get("LITELLM_URL", "http://localhost:4000").rstrip("/")
LITELLM_MASTER_KEY = os.environ.get("LITELLM_MASTER_KEY", "")
DEFAULT_MODEL = os.environ.get("DEFAULT_MODEL", "claude-sonnet-4-5")
RESOLUTION = os.environ.get("RESOLUTION", "1920x1080")
MAX_STEPS = int(os.environ.get("AGENT_MAX_STEPS", "150"))
# Consecutive text-only replies (no tool calls) before the run is ended with
# the model's last reply as the wrap-up. Weaker models answer in prose instead
# of calling task_complete; without a cap the "continue" nudge loops to
# MAX_STEPS paying full input cost every turn.
IDLE_REPLY_LIMIT = int(os.environ.get("AGENT_IDLE_REPLY_LIMIT", "3"))
ASK_USER_TIMEOUT = int(os.environ.get("ASK_USER_TIMEOUT", "600"))
COMMAND_TIMEOUT = int(os.environ.get("COMMAND_TIMEOUT", "60"))
# Background helper agents (spawn_agent): headless — web/text tools and a
# display-less shell, own step cap and optional cheaper model.
# SUBAGENT_MODEL empty = same model as the main agent.
SUBAGENT_MAX_STEPS = int(os.environ.get("SUBAGENT_MAX_STEPS", "40"))
SUBAGENT_MAX_CONCURRENT = int(os.environ.get("SUBAGENT_MAX_CONCURRENT", "4"))
SUBAGENT_MODEL = os.environ.get("SUBAGENT_MODEL", "")
# Helpers one run may spawn in total — bounds an orchestrator that keeps
# re-spawning a worker for a step that never yields evidence.
SUBAGENT_MAX_TOTAL = int(os.environ.get("SUBAGENT_MAX_TOTAL", "12"))
# Orchestrator mode: the main agent plans, delegates and assembles — it
# has no web_search/fetch_url of its own, so every research step is a
# focused worker with one goal. One agent juggling the plan, the
# research and the deliverable in a single context is where small models
# lose the thread; a worker with one goal and five tools has nothing to
# confuse. off = the classic single agent that may research itself.
AGENT_ORCHESTRATE = os.environ.get("AGENT_ORCHESTRATE", "on").lower() not in (
    "off", "0", "false", "no")
# Plan-lint / evidence-gate bounces one run tolerates before the checklist
# is accepted as posted — a gate a weak model can't satisfy must not wedge
# the run (same shape as AGENT_VERIFY_MAX_REJECTS).
GATE_MAX_BOUNCES = int(os.environ.get("AGENT_GATE_BOUNCES", "4"))
# Stronger model a struggling run escalates to — pair a cheap DEFAULT_MODEL
# with a premium ESCALATION_MODEL so easy work stays cheap and only
# demonstrably-stuck runs pay premium rates. Empty = never escalate.
ESCALATION_MODEL = os.environ.get("ESCALATION_MODEL", "")
# Unstick rescues the run gets on the starting model before switching.
ESCALATION_RESCUES = int(os.environ.get("ESCALATION_RESCUES", "2"))
# Stall detection (see StallDetector). The exact-repeat counter only sees
# the last 10 calls, so a loop longer than that never trips it: a run that
# cycles through the same 16 dead ends repeats each one once per lap. The
# cycle detector fires when AGENT_CYCLE_WINDOW acting calls in a row were
# all seen earlier in the run and disables the whole cycle at once.
CYCLE_WINDOW = int(os.environ.get("AGENT_CYCLE_WINDOW", "12"))
# Steps before the cap at which the run is told to stop exploring and
# deliver what it has — a run cut off mid-research delivers nothing at all.
WRAP_UP_STEPS = int(os.environ.get("AGENT_WRAP_UP_STEPS", "10"))
# Wheel clicks per scroll call. Models pass pixel counts at times ("500"),
# which jumps to the page end and back without ever showing the middle.
SCROLL_MAX_CLICKS = int(os.environ.get("SCROLL_MAX_CLICKS", "20"))
# Per-run dollar ceiling — bounds the worst case of a runaway loop. The run
# stops with an error once its spend passes this; 0 = no cap.
AGENT_MAX_USD = float(os.environ.get("AGENT_MAX_USD", "0"))
# Tool outputs older than the newest TOOL_RESULT_HISTORY results shrink to
# TOOL_RESULT_STUB_CHARS chars — stale DOM dumps and page text are the
# biggest payloads in history and every request re-bills them.
TOOL_RESULT_HISTORY = int(os.environ.get("TOOL_RESULT_HISTORY", "8"))
TOOL_RESULT_STUB_CHARS = int(os.environ.get("TOOL_RESULT_STUB_CHARS", "300"))
# Auxiliary loops on cheaper models: COMPACT_MODEL summarizes history
# (text-only — any cheap chat model works), JANITOR_MODEL runs the post-run
# cleanup pass (must be vision-capable). Empty = use the run's model.
COMPACT_MODEL = os.environ.get("COMPACT_MODEL", "")
JANITOR_MODEL = os.environ.get("JANITOR_MODEL", "")
# Wrap-up verification: before task_complete is accepted, a checker call
# audits the summary against a ledger of the run's actual tool calls —
# facts, URLs and "I sent the file" claims no tool result backs get bounced
# back as a critique instead of reaching the user. VERIFY_MODEL empty =
# COMPACT_MODEL or the run's model; AGENT_VERIFY=off disables the check.
AGENT_VERIFY = os.environ.get("AGENT_VERIFY", "on").lower() not in (
    "off", "0", "false", "no")
VERIFY_MODEL = os.environ.get("VERIFY_MODEL", "")
# Rejections one run tolerates before the wrap-up is accepted anyway with
# the caveats attached to it — a checker that can never be satisfied must
# not loop the run forever.
VERIFY_MAX_REJECTS = int(os.environ.get("AGENT_VERIFY_MAX_REJECTS", "2"))
# Runs with fewer completed tool calls skip the audit — a two-step task has
# nothing worth cross-checking and the call would be pure latency.
VERIFY_MIN_CALLS = int(os.environ.get("AGENT_VERIFY_MIN_CALLS", "4"))
# Ledger budget: head + tail of the call list are kept, the middle drops —
# the claims under review gather at both ends (early research, late writes).
VERIFY_INPUT_CHARS = int(os.environ.get("AGENT_VERIFY_INPUT_CHARS", "120000"))
VERIFY_RESULT_CHARS = int(os.environ.get("AGENT_VERIFY_RESULT_CHARS", "400"))
# Long-horizon support. When a request's prompt_tokens exceed
# AGENT_COMPACT_RATIO of the model's usable context window (max_input_tokens
# from LiteLLM's /model/info, capped by MODEL_CONTEXT_LIMITS;
# AGENT_CONTEXT_LIMIT is the fallback when it reports none — deliberately
# conservative: over-compacting wastes one call, under-compacting kills the
# run), history is summarized into a handoff note and the loop continues on
# summary + todos + the last AGENT_COMPACT_KEEP messages.
COMPACT_RATIO = float(os.environ.get("AGENT_COMPACT_RATIO", "0.75"))
COMPACT_CONTEXT_LIMIT = int(os.environ.get("AGENT_CONTEXT_LIMIT", "128000"))
COMPACT_KEEP = int(os.environ.get("AGENT_COMPACT_KEEP", "6"))
# Ceiling on the context a model may use, per model — comma-separated
# name=tokens entries ("ollama/qwen2.5vl:7b=32k", k/m suffixes ok); "*"
# caps every model with no entry of its own. A cap only shrinks the
# window model_context_limit resolves — the detected size wins when it's
# smaller — so the agent compacts early enough to stay under it. Empty =
# every model may use its full window.
MODEL_CONTEXT_LIMITS = os.environ.get("MODEL_CONTEXT_LIMITS", "")
# Steps without a plan call before the checklist is nudged back
# into view on long tasks. TODO_NUDGE_STEPS covers the empty case — a run
# that deep with no checklist usually means the task only looked small.
TODO_REMIND_STEPS = int(os.environ.get("TODO_REMIND_STEPS", "20"))
TODO_NUDGE_STEPS = int(os.environ.get("TODO_NUDGE_STEPS", "10"))
TODO_MAX_ITEMS = int(os.environ.get("TODO_MAX_ITEMS", "30"))
# Consecutive turns of exactly one research call (web_search/fetch_url)
# before the "parallelize or delegate" nudge rides the next tool result —
# serial research is how a five-minute job turns into 50 billed turns.
RESEARCH_NUDGE_STREAK = int(os.environ.get("AGENT_RESEARCH_NUDGE", "5"))
SCREENSHOT_MAX_EDGE = int(os.environ.get("SCREENSHOT_MAX_EDGE", "1568"))
SCREENSHOT_MAX_PIXELS = int(os.environ.get("SCREENSHOT_MAX_PIXELS", "1000000"))
SCREENSHOT_HISTORY = int(os.environ.get("SCREENSHOT_HISTORY", "3"))
# Cap auto-attached frames lower than SCREENSHOT_MAX_PIXELS — routine "did
# it change?" looks don't need full detail; the screenshot tool still
# returns full-res on demand. 0 = same cap as manual screenshots.
SHOT_AUTO_PIXELS = int(os.environ.get("SCREENSHOT_AUTO_PIXELS", "0"))
# Screen actions take effect before the UI finishes repainting — without a
# pause the post-action screenshot can capture the pre-action frame, which
# reads as "the click did nothing" and invites a blind re-click into
# whatever has since appeared. ACTION_SETTLE_SECS is how long we poll for a
# changed frame before declaring the screen unchanged; INTER_ACTION_DELAY
# is the minimum gap between consecutive screen-mutating calls in one batch
# (they were planned against the same frame).
ACTION_SETTLE_SECS = float(os.environ.get("ACTION_SETTLE_SECS", "1.5"))
INTER_ACTION_DELAY = float(os.environ.get("INTER_ACTION_DELAY", "0.35"))
# Pixel-click assist via AT-SPI: every aimed point gets looked up in the
# a11y tree so the tool result can say what was actually hit, and a point
# that lands on dead space within CLICK_SNAP_PX of an actionable element
# snaps to that element's center. CLICK_A11Y=off disables both (clicks then
# go exactly where aimed with no annotation).
CLICK_A11Y = os.environ.get("CLICK_A11Y", "on").lower() not in (
    "off", "0", "false", "no")
CLICK_SNAP_PX = int(os.environ.get("CLICK_SNAP_PX", "24"))
# VLMs that emit coordinates on a normalized 0-1000 grid rather than
# screenshot pixels — Qwen3-VL standardised on this (its ViT rescales
# inputs internally, so its pixel space is unknowable from our side) and
# Gemini 2.5 does the same. Comma-separated substrings matched against the
# model id; set empty to disable.
NORMALIZED_COORD_MODELS = [
    p for p in os.environ.get(
        "NORMALIZED_COORD_MODELS", "qwen3-vl,gemini").lower().split(",")
    if p]
LLM_MAX_RETRIES = int(os.environ.get("LLM_MAX_RETRIES", "4"))
# Files/images sent to the user travel the chat WebSocket as base64 — uvicorn's
# default ws max message is 16 MB, so cap payloads comfortably under it.
SEND_FILE_MAX_BYTES = int(os.environ.get("SEND_FILE_MAX_BYTES", str(9 * 1024 * 1024)))
CDP_PORT = int(os.environ.get("CDP_PORT", "9222"))
CDP_HTTP = f"http://localhost:{CDP_PORT}"
# Structured-control helpers run under the SYSTEM python (/usr/bin/python3)
# — pyatspi and uno arrive as apt packages (python3-pyatspi / python3-uno),
# not in the daemon venv. UNO_PORT is the LibreOffice listener the gut-office
# wrapper injects into every soffice/libreoffice launch.
SYS_PY = "/usr/bin/python3"
GUT_DIR = Path(__file__).resolve().parent
ATSPI_HELPER = GUT_DIR / "atspi.py"
UNO_HELPER = GUT_DIR / "uno_eval.py"
UNO_PORT = os.environ.get("GUT_UNO_PORT", "2002")
# OpenSERP endpoint for web_search — the compose stacks run one on the
# internal network, gut-bot runs it as a systemd unit on localhost.
# Empty = fall back to DuckDuckGo's HTML endpoint.
OPENSERP_URL = os.environ.get("OPENSERP_URL", "").rstrip("/")
# Concurrent /mega/search requests daemon-wide — openserp launches a Chrome
# per query, so unbounded parallelism from parallel workers is a self-DoS.
OPENSERP_MAX_CONCURRENT = int(os.environ.get("OPENSERP_MAX_CONCURRENT", "3"))
_openserp_sem: asyncio.Semaphore | None = None
_openserp_sem_n = 0


def _search_sem() -> asyncio.Semaphore:
    """The search semaphore for the current OPENSERP_MAX_CONCURRENT —
    rebuilt lazily so a runtime config change takes effect on the next
    acquire (holders of the old one just finish)."""
    global _openserp_sem, _openserp_sem_n
    n = max(1, OPENSERP_MAX_CONCURRENT)
    if _openserp_sem is None or _openserp_sem_n != n:
        _openserp_sem, _openserp_sem_n = asyncio.Semaphore(n), n
    return _openserp_sem
# Locale passed to openserp (lang/region) and DuckDuckGo (kl) — unset =
# whatever the container IP implies, which is usually wrong for non-English
# or location-specific queries. The model can override per call.
SEARCH_LANG = os.environ.get("SEARCH_LANG", "")
SEARCH_REGION = os.environ.get("SEARCH_REGION", "")
# Engines openserp's /mega/search fans out to, comma-separated. Empty =
# chosen per query language (see _openserp_engines): Baidu answers every
# query in Chinese whatever `lang` says and Yandex leans Russian, so for a
# Danish price query they only push the real hits down the list.
OPENSERP_ENGINES = os.environ.get("OPENSERP_ENGINES", "")
SHOT_PATH = Path("/tmp/gut_screen.png")
HOME_DIR = Path(os.environ.get("HOME") or Path.home())
KEY_FILE = Path(os.environ.get("GUT_KEY_FILE") or HOME_DIR / ".gut_litellm_key")
DEVICE_NAME = os.environ.get("DEVICE_NAME") or socket.gethostname()
# Shared secret the client must present (Bearer header or ?token=) on every
# API call and the chat socket. Empty = open (acceptable on localhost-only
# installs). The gut-bot package and the Electron local stack both generate
# one at install time and reuse it as the VNC password.
GUT_API_TOKEN = os.environ.get("GUT_API_TOKEN", "")


def _build_version() -> str:
    v = os.environ.get("GUT_VERSION")
    if v:
        return v
    try:
        return Path("/opt/gut/VERSION").read_text().strip() or "dev"
    except OSError:
        return "dev"


GUT_VERSION = _build_version()


def _install_kind() -> str:
    # docker = the gut-desktop image (GUT_VERSION is a baked-in ENV); deb =
    # the gut-bot package (/opt/gut payload, /etc/gut env); else dev.
    if os.environ.get("GUT_VERSION") or Path("/.dockerenv").exists():
        return "docker"
    if (Path("/etc/gut/gut-bot.env").exists()
            or Path("/opt/gut/VERSION").exists()):
        return "deb"
    return "dev"


INSTALL_KIND = _install_kind()
# Conversation transcripts + model context persist on the desktop-home
# volume (or ~/.gut outside docker) so they survive container rebuilds.
GUT_DATA_DIR = Path(os.environ.get("GUT_DATA_DIR") or Path.home() / ".gut")
CONV_DIR = GUT_DATA_DIR / "conversations"
# TLS: start.sh generates a self-signed cert into the data dir; a thin
# TLS->TCP proxy in this process terminates GUT_TLS_PORT and forwards to the
# plain uvicorn socket, so a second app instance (and a second lifespan) is
# never needed. Clients don't need a CA — they pin the cert fingerprint
# after the password-authenticated /api/hello handshake. GUT_NOVNC_TLS_PORT
# is only advertised to clients; the TLS websockify is spawned by start.sh.
TLS_CERT = Path(os.environ.get("GUT_TLS_CERT") or GUT_DATA_DIR / "tls" / "cert.pem")
TLS_KEY = Path(os.environ.get("GUT_TLS_KEY") or GUT_DATA_DIR / "tls" / "key.pem")
TLS_PORT = int(os.environ.get("GUT_TLS_PORT", "8443"))
NOVNC_TLS_PORT = int(os.environ.get("GUT_NOVNC_TLS_PORT", "6443"))
# Where the plain uvicorn socket lives — the TLS proxy forwards to it.
# start.sh passes the same var to uvicorn's --port so overrides stay in
# sync (AGENT_PORT in .env only moves the host-side compose mapping).
HTTP_PORT = int(os.environ.get("GUT_HTTP_PORT", "8000"))
# When start.sh parks uvicorn on a loopback-only port (GUT_UVICORN_PORT),
# this process multiplexes the public ports itself: TLS ClientHellos get
# routed to the :8443/:6443 terminators, everything else to the plain
# backends — so pinned TLS works wherever the plain ports already reach,
# no extra firewall holes or port forwards needed.
UVICORN_PORT = int(os.environ.get("GUT_UVICORN_PORT", "0"))
BACKEND_PORT = UVICORN_PORT or HTTP_PORT
MUX = bool(UVICORN_PORT)
NOVNC_PORT = int(os.environ.get("GUT_NOVNC_PORT", "6080"))
NOVNC_PLAIN_PORT = int(os.environ.get("GUT_NOVNC_PLAIN_PORT", "6081"))
TLS_FP = None  # sha256 of the cert the proxy actually presents
TLS_ERR = None  # why TLS is down, when it is — surfaced via /api/hello
# Files the user attaches in the composer land on the desktop itself, where
# the agent's file/shell tools can read them (paths are relative to home).
UPLOAD_DIR = Path(os.environ.get("GUT_UPLOAD_DIR") or HOME_DIR / "uploads")
MAX_ATTACHMENTS = 8
# Per-run cleanup: a deterministic sweep (close tabs/windows/processes the
# run created, wipe ~/scratch) followed by a short janitor LLM pass for
# residue the sweep can't see. GUT_CLEANUP=off disables both.
GUT_CLEANUP = os.environ.get("GUT_CLEANUP", "on").lower() not in (
    "off", "0", "false", "no")
JANITOR_MAX_STEPS = int(os.environ.get("JANITOR_MAX_STEPS", "10"))
SCRATCH_DIR = HOME_DIR / "scratch"
# Baseline of what was running when a task started, persisted so a daemon
# restart mid-run still lets the next boot sweep that run's leftovers.
RUN_STATE_FILE = GUT_DATA_DIR / "run_state.json"
# Base64 inflates ~33%; keep the ws message comfortably under uvicorn's
# 16 MB frame cap.
ATTACH_TOTAL_MAX_BYTES = int(
    os.environ.get("ATTACH_TOTAL_MAX_BYTES", str(12 * 1024 * 1024)))
# Self-update status file (deb installs) — written by the detached updater
# script, read by GET /api/update.
UPDATE_STATUS_FILE = GUT_DATA_DIR / "update.json"
GITHUB_REPO = os.environ.get("GUT_RELEASE_REPO", "valteryde/gut")

pyautogui.FAILSAFE = False
pyautogui.PAUSE = 0.05

SYSTEM_PROMPT = """You are Gut, an autonomous operator of a Linux desktop (XFCE4, {res}).
You perceive the screen through screenshots and act with mouse/keyboard tools.
{role}
Environment:
{web}- Google Chrome is installed with DevTools on localhost:{cdp}. Use the
  browser_* tools only when {textalt} can't do the job — pages needing
  JS, logins/sessions, forms, or visual checks (DOM refs, not pixels):
  open_url to get somewhere, browser_text to read the
  page's text, browser_dom to list interactive elements as #refs, then
  browser_click / browser_type by ref; browser_eval runs arbitrary JS.
  open_url reports the HTTP status and title: a 404 means the
  URL was wrong — {on404}, don't try variations of it. Read
  pages with browser_text, not by scrolling through screenshots.
  Fall back to pixel tools for anything outside the page.
- LibreOffice Writer, Calc and Impress are installed
  (`libreoffice --writer/--calc/--impress`). office_eval runs Python-UNO
  against the LIVE document — insert content, format, save, export as PDF —
  no GUI clicking. Globals: `doc` (current document, None if none open),
  `desktop`, `load(path)`, `file_url`, `uno`; set `result` to return a
  value. For producing documents from scratch, prefer writing the file
  directly (python-docx, openpyxl, pandoc, or `soffice --headless
  --convert-to pdf out.docx` — the soffice wrapper gives batch runs their
  own profile, so this is safe while the GUI is open) and only open the
  GUI to eyeball the result. Build the whole file from one script so the
  numbers have a single source —
      python3 - <<'EOF'
      import openpyxl
      wb = openpyxl.Workbook(); ws = wb.active
      ws.append(["Item", "Qty", "Price"]); ws.append(["Pork", 44, 88])
      wb.save("budget.xlsx")
      EOF
  then verify the artifact before send_file: reload it, ls the path, or
  print a check total — never deliver a file you haven't confirmed exists
  and opens.
- desktop_tree is the native-app equivalent of browser_dom: it dumps the
  focused window's accessibility tree as numbered #refs with role, name,
  actions and bounds. Act on them with desktop_act (press/activate/select
  by ref), desktop_type (set text directly) or desktop_click (pixel-click
  the ref's bounds). Always prefer these over guessing pixel coordinates;
  re-run desktop_tree after the UI changes — refs go stale.
- run_command gives you a bash shell (cwd {home}, DISPLAY already set) —
  and Python 3 with pip: openpyxl, python-docx and pillow are
  preinstalled, `pip install` covers the rest. This is your main tool for
  math, parsing, data and producing files — compute and generate in code,
  never by hand: a script is one source of truth where hand-typed numbers
  and file contents drift. An ImportError means install the package or
  pick an installed one — never a reason to abandon code for a GUI or
  echo workaround. Launch GUI apps in the background so the command
  returns, e.g. `google-chrome &`.
- {home}/scratch is wiped when the task ends — use it for temp and
  intermediate files. Keep anything needed later elsewhere in {home}.
- {coords}
- Older screenshots are dropped from context — only recent frames are kept.
  The text log of your actions stays; call screenshot for a fresh look.

Talking to the user — act like a teammate, not a live feed:
- Chat messages reach the user only via the send_* tools, ask_user and
  task_complete — never rely on anything else to reach them.
- The text you write alongside tool calls flashes in the user's status
  line as you work (and lands in a verbose log they can open): a short
  "comparing the quotes…" keeps them oriented. A few words now and then —
  it's an ephemeral status, not a message.
- send_message: a short chat update. Use sparingly — a milestone on a long
  task, a blocker, a finding worth flagging. Silence is fine while work is
  straightforward; do not narrate steps.
- send_file: deliver an artifact (report, spreadsheet, export, download) as a
  chat attachment. Paths are relative to {home}. The user can't browse this
  filesystem — before task_complete, send_file every file they'll want,
  including attached files you edited; a path in the summary is not delivery.
- send_image: show the user an image — a file, or a fresh screenshot when no
  path is given.
- Files the user attaches to a message are saved under {home}/uploads/ — the
  message text lists the exact paths; open them with run_command or the file
  tools. Attached images are also shown to you inline.
- ask_user: pause for input only when blocked on something only a human can
  provide — a decision, credential, 2FA or CAPTCHA. Never for information
  visible on screen, and never just because you're stuck — brainstorm more
  approaches instead.
- task_complete: ends the task and sends `summary` as your wrap-up message.
  The summary IS the answer: when the user asked a question, it holds the
  values with units and the source URLs verbatim — "sources are
  provided" or "see the report" is not an answer, the user sees nothing
  else. Then what was done, where files are, and what to check.

Planning — match the effort to the task:
- Quick one-off actions: just do them — no plan, no checklist.
- Anything with several distinct phases or likely more than ~15 steps
  (research-then-write, multi-site collection, install-and-configure) —
  and any task that gathers data and then produces a file always counts:
  call plan once before you start acting, with `summary` (a concise plan
  — it posts to the user as a card, no approval needed, keep working)
  and `steps` (your checklist).
  Pass the full `steps` list every call, keep exactly one item
  in_progress, mark steps done as you go — task_complete bounces a
  stale list back, so it ends the run truthful.
- One step = one target, one verifiable outcome. Give each step a `kind`:
  research (look ONE thing up — it becomes one worker), build (make or
  compute a file with run_command), desktop (act on the screen), deliver
  (send_file), decide (a judgment needing no tool). "Find prices for A, B
  and C" is three research steps, never one — a bundled research step is
  rejected and you must re-post it split. A step that takes 15 tool calls
  was really several steps, so split it when you notice.
- A step only becomes done on evidence of its kind since it started — a
  worker's report for research, a command that ran for build, a send_file
  for deliver. Marking a step done with no such evidence puts it straight
  back in_progress: do the work, then mark it.
- Update the list the moment you start work that belongs to a different
  step — a stale checklist triggers forced reposts that spend real turns.
  When in doubt, post the plan and checklist — they cost little and the
  user watches both live. If the scope changes, call plan again with a
  new summary and list.
- On very long runs your older context gets compacted into a handoff
  summary — the checklist always survives it. Anything else worth keeping
  (paths, URLs, decisions, findings) belongs in the todo text or in files
  under {home}.

Accuracy — never fabricate:
- Facts that end up in a deliverable (prices, dates, names, statistics,
  URLs) must come from a tool result in this run — {factsrc}, a file
  you read, command output. Never fill gaps from memory or invent
  plausible-looking values.
- When the user asks for sources, a source is a specific page that was
  actually opened ({srcsrc}) — never a bare homepage, and never
  a URL constructed to look right.
- If real attempts can't verify a value, mark it as an estimate in the
  deliverable and tell the user which parts are unverified. A flagged
  estimate beats a confident invention — a wrong "fact" delivered as truth
  is the worst possible outcome.
{leads}
Guidelines:
- A fresh screenshot is attached automatically after each turn's actions; only
  call screenshot when nothing changed or you need an extra look.
- Batch predictable sequences into one response — emit several tool calls at
  once (e.g. click field → type → press enter{batch}). Split only when the
  next step depends on a result.
- Prefer desktop_* refs, office_eval, keyboard shortcuts and run_command
  over pixel hunting — raw coordinates are the fallback, not the default.
- In file chooser dialogs press ctrl+l to open the location bar, type the
  absolute path and hit enter — never navigate the places list by mouse.
- {online}
- Think in English — your reasoning and tool arguments stay English for
  quality — but face the user in their language: chat messages, questions
  and deliverables match the language they write in, and they may switch
  languages between requests. {locale}
- If an action changes nothing after two tries, stop and brainstorm at least
  5 different approaches (keyboard navigation, menus, run_command, the
  browser_* tools, a different app entirely) and try the most promising
  untried one. Never keep repeating the same click, and never give up —
  there is almost always another way forward.
- Click results tell you what element actually sits under the point — read
  them. "isn't clickable there" or "nothing at that point" means your aim
  was off or the target isn't a widget; re-aim or switch to desktop_tree
  #refs instead of clicking the same spot again.
- type_text and key go to whatever window has focus — if keystrokes aren't
  landing, click the target field first (or use browser_type on web pages).
- Never click tel:/mailto: links; they're blocked and just pop a dead-end OS
  dialog. Read phone numbers and addresses with browser_text instead.
- If an OS dialog does appear ('Open xdg-open?', permission prompts),
  press Escape to dismiss it rather than pixel-clicking buttons. File
  pickers you opened yourself are fine — use ctrl+l (see above) or
  desktop_tree to work them.
- If a login, 2FA, CAPTCHA or genuinely ambiguous decision blocks you, call
  ask_user — the human can click into the live screen to help, then resume you.
- When your task ends the system closes the apps, windows and browser tabs
  you opened and stops leftover processes. Logins and cookies persist across
  tasks — never log out or wipe browser data as "cleanup".

{delegate}"""

# The operating model, per mode. Orchestrator: plan → one worker per
# research target → assemble from their files → deliver. Classic: the
# single agent that may research itself, with delegation as an option.
ROLE_ORCHESTRATOR = """
You work as an orchestrator: you plan, delegate research to focused
workers, assemble their findings into the deliverable, and talk to the
user. You have no web_search or fetch_url of your own — every lookup is a
worker with one goal. You keep the screen, the shell and the user.
"""
ROLE_CLASSIC = ""

WEB_ORCHESTRATOR = """- Web research is done by workers (spawn_agent), one goal each — see
  Workers below. run_command + curl still works for APIs and downloads
  whose URL a worker already reported.
"""
WEB_CLASSIC = """- To find or read web content, start with the text tools — fast and cheap,
  no browser or screenshots needed: web_search finds pages, fetch_url reads
  a page's text and links (run_command + curl works for APIs and downloads).
"""

LEADS_CLASSIC = """- Search snippets are leads, not sources: fetch_url the result page before
  putting its claims or its URL into the deliverable.
"""
LEADS_ORCHESTRATOR = """- A worker's report ends with an evidence footer the daemon wrote —
  pages fetched, files written. Figures in a report whose footer shows no
  page fetched are unverified: send the worker back or flag them.
"""

DELEGATE_ORCHESTRATOR = """Workers — spawn_agent runs a focused helper in the background:
- One worker, one goal: "the current price per kg of pulled pork in
  Denmark", not "prices for the whole menu". Pass the goal as `task`, in
  English, with the unit, locale and time frame you need — the worker sees
  only that text, never your conversation. The daemon wraps it in a fixed
  contract: verify through tools, write findings to
  {home}/scratch/<name>/, report values + sources + an UNVERIFIED list.
- Spawn every independent research step at once (up to {subcap} run
  concurrently), then collect_agent — its report arrives as a message
  either way. Keep the desktop work and the user yourself; a worker has
  no screen, no browser and cannot ask anyone anything.
- Assemble, don't transcribe: build the deliverable with one script that
  reads the workers' files (JSON/CSV under {home}/scratch/) so every
  number in it has a single source. Retyping figures from reports into a
  file is how numbers drift. Then verify the artifact — reload it, print
  a check total — before send_file.
- When the deliverable carries sources, a source is the URL the worker
  fetched — copy it verbatim from its file/report into the artifact.
  A shop or site name ("Nemlig", "tilbudsugen.dk" as text) is not a
  source; a row whose source cell has no URL is unsourced.
- A worker that reports nothing usable gets a narrower goal, not a
  retry of the same one; {subtotal} spawns per run is the cap.
"""
DELEGATE_CLASSIC = """Delegating — spawn_agent runs a helper agent in the background:
- Give it self-contained headless subtasks: web research, reading or writing
  files, crunching data with run_command. It has no screen, no browser and
  no way to reach the user — anything needing eyes, clicks or logins is yours.
- Its final report arrives as a message mid-run; collect_agent(name) blocks
  until it (or any helper) reports. Share artifacts through files under {home}.
- Delegate aggressively. The moment a task has two or more independent
  research targets, sources to cross-check, or file/data jobs that need no
  screen, spawn one helper per target and keep the desktop work yourself —
  "research these 5 companies" means five helpers, not five of your own
  steps. A list of lookups (prices for N items, facts about N targets) is
  the canonical case: one helper per group, or one web_search `queries`
  batch — never a turn per item. Doing everything serially yourself is the
  slow path: if a subtask can run headless while you act, it should be a
  helper. At most {subcap} helpers run at once.
"""


def system_prompt(res: str, coords: str) -> str:
    """The main agent's system prompt for the active mode."""
    orch = AGENT_ORCHESTRATE
    return SYSTEM_PROMPT.format(
        res=res, cdp=CDP_PORT, coords=coords, home=HOME_DIR,
        role=ROLE_ORCHESTRATOR if orch else ROLE_CLASSIC,
        web=WEB_ORCHESTRATOR if orch else WEB_CLASSIC,
        textalt="a worker" if orch else "the text tools",
        on404=("have a worker find the right page"
               if orch else "go back to web_search"),
        factsrc=("a worker's report and the file it wrote"
                 if orch else "a fetched page"),
        srcsrc=("a URL a worker fetched and listed in its report"
                if orch else "with fetch_url or browser_text"),
        leads=LEADS_ORCHESTRATOR if orch else LEADS_CLASSIC,
        batch=("; several spawn_agent calls" if orch else
               "). Independent web_search and fetch_url calls in one reply "
               "run in parallel, and web_search's `queries` list runs "
               "several lookups in a single call — N lookups is one call, "
               "not N turns"),
        online=("For anything online, a worker first; browser_* only for "
                "what genuinely needs a browser (JS, auth, interaction)."
                if orch else
                "For anything online, web_search/fetch_url first; browser_* "
                "only when they fail or the page genuinely needs a browser "
                "(JS, auth, interaction)."),
        locale=("Tell each worker the locale its goal targets (a request "
                "in Danish asking for Danish prices → Denmark, DKK, Danish "
                "shops); never assume the server's locale."
                if orch else
                "Set web_search's lang/region to the locale each query "
                "targets (a request in Danish asking for Danish prices → "
                "DA/DK); never assume the server's locale."),
        delegate=(DELEGATE_ORCHESTRATOR if orch else DELEGATE_CLASSIC).format(
            home=HOME_DIR, subcap=SUBAGENT_MAX_CONCURRENT,
            subtotal=SUBAGENT_MAX_TOTAL))


SUBAGENT_PROMPT = """You are '{name}', a focused worker spawned by Gut on a Linux desktop.
You have exactly one goal — the GOAL below — and the orchestrator that
spawned you only ever sees your final report.

Tools:
- web_search + fetch_url: find and read web content (text only).
  web_search's `queries` list runs several searches in one call; set
  lang/region to the locale the goal targets (Danish prices → DA/DK).
- run_command: bash, cwd {home}. There is no display — GUI apps and anything
  needing a screen fail; stay headless (curl, scripts, files, packages).
- send_file: deliver a file to the user (paths relative to {home}) — only
  when the goal says to.
- task_complete: finish; `summary` becomes your report.

Method — thorough on one thing beats shallow on many:
1. Search (2-3 queries at once), pick the 2-4 most promising result pages.
2. fetch_url each of them — a snippet is a lead, not a source. Read the
   actual figure, unit and date off the page.
3. Cross-check: two sources agreeing is a finding; one source is a lead
   you flag as single-source.
4. Write findings to {workdir}/findings.json (machine-readable: value,
   unit, currency, source_url, date, note) plus anything else the goal
   asks for. Create the directory first (mkdir -p). Every figure and URL
   in your report must also be in that file — the file is the record,
   the report only summarizes it; task_complete checks this.
5. task_complete with a report of 5-15 lines: the values with units and
   currency, the source URLs you actually fetched, and an UNVERIFIED list
   of anything you could not confirm. Never pad, never speculate.

Rules:
- Stay on the goal. Anything outside it — even if interesting — is not
  your job; note it in one line at most.
- Every number, name and URL in your report must come from a tool result
  in this run. If real attempts fail, say UNVERIFIED and what you tried —
  a flagged gap beats a plausible guess.
- You cannot see the screen, drive the browser, or ask anyone anything —
  a login wall or CAPTCHA is a blocker you report, not something to guess
  around.
- Reason and report in English whatever the goal's language — the
  orchestrator faces the user.
- A reply with no tool calls also ends your run, with the reply as the
  report — but prefer task_complete so the intent is clear.
"""

# The user-turn wrapper spawn_agent puts around the orchestrator's goal —
# the contract is the daemon's, so a terse or sloppy goal still yields a
# report in the shape the gate and the orchestrator expect.
SUBAGENT_TASK = """GOAL: {goal}

Deliver: values with units/currency, the URLs you fetched, an UNVERIFIED
list. Write {workdir}/findings.json before you finish (mkdir -p first).
Only this goal — nothing else."""

COORD_PROMPT_PIXEL = ("Tool coordinates refer to pixels in the screenshot "
                      "image you received.")
COORD_PROMPT_NORM = ("Tool coordinates use a normalized 0-1000 grid over the "
                     "screenshot — [500, 500] is the center of the screen, "
                     "[0, 0] the top-left corner.")
# Gemini's spatial convention is [y, x] on the 0-1000 grid — telling it to
# emit [x, y] makes it fight its training, so ask for [y, x] explicitly.
COORD_PROMPT_NORM_YX = ("Tool coordinates use a normalized 0-1000 grid over "
                        "the screenshot in [y, x] order — [500, 500] is the "
                        "center of the screen, [0, 0] the top-left corner.")


# ── Step kinds ─────────────────────────────────────────────────────────────
# Every checklist step has a kind, and the kind decides what counts as
# evidence that it was actually done (see update_todos). The model may set
# it; when it doesn't, a lexicon guess fills it in — the gate must never
# depend on a weak model remembering a field.
STEP_KINDS = ("research", "build", "desktop", "deliver", "decide")
_KIND_LEXICON = (
    ("decide", r"\b(decide|decision|choose|pick|assume|beslut|vælg|antag)"),
    ("deliver", r"\b(send|deliver|attach|share|hand ?over|send_file|"
                r"sende?|aflever|lever|vedhæft)\b"),
    ("desktop", r"\b(open|click|log ?in|login|sign ?in|type|navigate|"
                r"install|configure|screenshot|browser|app|window|form|"
                r"upload|download|åbn|klik|installer)\b"),
    ("build", r"\b(creat|build|make|generat|writ|compil|assembl|produc|"
              r"export|sav|calculat|comput|format|excel|spreadsheet|"
              r"sheet|xlsx|docx|csv|pdf|report|script|budget|"
              r"lav|opret|skriv|gem|udregn|beregn|ark|dokument|rapport)"),
    ("research", r"\b(search|find|look ?up|research|compar|check|price|"
                 r"prices|cost|source|verify|collect|gather|identif|"
                 r"list|survey|søg|undersøg|sammenlign|tjek|pris|"
                 r"priser|kilde|indsaml)"),
)


def infer_step_kind(content: str) -> str:
    """Guess a step's kind from its wording — decide/deliver/desktop/build
    verbs win over research ones, since "find prices and build a sheet"
    is a build step that happens to mention prices; desktop beats build
    because a desktop step accepts build evidence (run_command) while the
    reverse would false-bounce a clicked-through step. Falls back to
    research: a vague step is more likely a lookup than a screen action."""
    text = content.lower()
    for kind, pat in _KIND_LEXICON:
        if re.search(pat, text):
            return kind
    return "research"


# Enumerations inside one research step — "priser på A, B og C" — are the
# canonical bundled step: one worker asked for three things does one well
# and guesses two. Separators: commas, semicolons, slashes and the
# conjunctions of the languages the agent meets most.
_ENUM_SPLIT_RE = re.compile(
    r"\s*(?:,|;|/|&|\bog\b|\beller\b|\band\b|\bor\b|\bund\b|\boder\b)\s*",
    re.I)


def enumerated_targets(content: str) -> list[str]:
    """The items a step enumerates, when it bundles three or more —
    otherwise []. Splits on list separators and keeps items that look like
    nouns (short, no verbs of their own)."""
    tail = re.split(r"\b(?:på|for|of|om|about|af|til|to|from|fra)\b",
                    content, maxsplit=1, flags=re.I)
    body = tail[-1] if len(tail) > 1 else content
    items = [i.strip(" .:-") for i in _ENUM_SPLIT_RE.split(body)]
    items = [i for i in items if 1 <= len(i.split()) <= 5]
    return items if len(items) >= 3 else []


class AgentState:
    def __init__(self) -> None:
        self.model = DEFAULT_MODEL
        self.litellm_key: str = LITELLM_MASTER_KEY
        self.clients: set[WebSocket] = set()
        self.task: asyncio.Task | None = None
        self.paused = False
        self.stop = False
        self.pending_answer: asyncio.Future | None = None
        self.coord_scale = 1.0  # screenshot px -> real desktop px
        # Dimensions of the last image sent to the model (after downscale).
        try:
            w, h = RESOLUTION.lower().split("x")
            self.shot_size = (int(w), int(h))
        except ValueError:
            self.shot_size = (1024, 768)
        self.frame_hash: str | None = None  # last frame sent to the model
        self.browser_ws: str | None = None  # CDP ws url of the active page
        # Consecutive open_url calls that hit an error
        # status — the model constructing URLs from memory instead of using
        # ones a tool returned.
        self.url_fail_streak = 0
        # Tabs open before the run started — the user's; open_url never
        # navigates those away (the cleanup sweep leaves them alone too).
        self.baseline_tabs: set[str] = set()
        # Last desktop_tree ref map: ref -> a11y path ("app|i,j,...").
        # Stale after any UI change — desktop_tree refreshes it.
        self.a11y_map: dict[int, str] = {}
        # Contextual disclosure (tools_for_run): a browser_dom/desktop_tree
        # this run has produced refs to act on; tools_shown is the monotonic
        # set of CONTEXT_TOOLS already declared to the model.
        self.dom_seen = False
        self.tree_seen = False
        self.tools_shown: set[str] = set()
        self.conversation_id: str | None = None  # conversation of the current/last run
        self.session_usd = 0.0
        self.tokens_in = 0
        self.tokens_out = 0
        self.models_synced = False  # provider keys pushed into LiteLLM
        self.escalated = False  # run already switched to ESCALATION_MODEL
        # Background helpers: name -> {task, desc, conv, status, result,
        # model, usd, steps, delivered}. Finished entries queue on
        # subagent_inbox for injection into their conversation's context.
        self.subagents: dict[str, dict] = {}
        self.subagent_inbox = deque()
        # User messages sent while a run is active: {conv, text, files,
        # mode, seq}. "steer" entries inject into the running
        # conversation's context at the next step; "queue" entries are
        # picked up when the run ends — same-conversation ones even before
        # the end-of-run cleanup sweep.
        self.user_msgs = deque()
        # The running conversation's `plan` checklist (persisted at
        # <cid>.todos.json) and the long-horizon bookkeeping for reminders
        # and compaction.
        self.todos: list[dict] = []
        self.plan_shared = False
        self.steps_since_todo = 0
        # Consecutive turns that were exactly one research call — batched
        # turns and any other tool reset it. Trips RESEARCH_NUDGE_STREAK.
        self.solo_research = 0
        # Completion gate: task_complete bounces once while checklist items
        # are unfinished (todo_nudge_done); the reconcile flag pins the
        # bounced call's successor to plan so the list gets a real final
        # update, not just a note it can ignore.
        self.todo_nudge_done = False
        self.todo_reconcile = False
        # Evidence gate (update_todos): counters of evidence-bearing events
        # this run by step kind — a worker report landing (research), a
        # command or office_eval that ran (build), a send_file (deliver), a
        # screen action (desktop). step_marks snapshots the counters when a
        # step goes in_progress; marking it done needs a higher count of
        # its kind since — proof that *something* of the right sort
        # happened, not a model's say-so. gate_bounces counts rejected plan
        # updates (lint + gate) toward GATE_MAX_BOUNCES; spawned counts
        # helpers toward SUBAGENT_MAX_TOTAL.
        self.evidence: dict[str, int] = dict.fromkeys(STEP_KINDS, 0)
        self.step_marks: dict[str, dict[str, int]] = {}
        self.plan_mark: dict[str, int] = dict.fromkeys(STEP_KINDS, 0)
        self.gate_bounces = 0
        self.spawned = 0
        self.steps_since_compact = 99
        self.ctx_limit: dict[str, int] = {}  # model -> max_input_tokens
        # Context-window occupancy of the running conversation: the last
        # request's billed prompt_tokens and a proportional split
        # (system/shots/tools/chat) normalized to it — broadcast on `cost`.
        self.ctx_used = 0
        self.ctx_parts: dict = {}
        self.no_tool_choice: set[str] = set()  # models that 400 on tool_choice
        self.no_tool_required: set[str] = set()  # …and on the "required" form
        # Delivery bookkeeping for the unsent-output nudge at task_complete:
        # run_start marks the current task's beginning; delivered maps an
        # attachment's resolved path -> mtime when the user sent it (cumulative
        # across runs); sent_files maps path -> mtime at send_file time, so a
        # post-send edit counts as unsent again.
        self.run_start = 0.0
        self.delivered: dict[str, float] = {}
        self.sent_files: dict[str, float] = {}
        self.output_nudge_done = False
        # Wrap-up verification: rejections spent this run, and the last
        # critique — appended to the done text when the reject cap forces
        # acceptance, cleared when a summary passes clean.
        self.verify_rejects = 0
        self.verify_caveat: str | None = None

    @property
    def running(self) -> bool:
        return self.task is not None and not self.task.done()

    @property
    def phase(self) -> str:
        if not self.running:
            return "idle"
        if self.pending_answer is not None and not self.pending_answer.done():
            return "waiting_user"
        return "paused" if self.paused else "running"


state = AgentState()


# ── Conversations ────────────────────────────────────────────────────────────
# A conversation lives on one device and carries its own model context, so a
# follow-up message continues where the agent left off. On disk:
#   <id>.meta.json     — {id, title, model, created_at, updated_at, events}
#   <id>.jsonl         — append-only transcript events, each {seq, ts, ...event}
#   <id>.context.json  — the LLM message array, rewritten after each step

_CID_RE = re.compile(r"[0-9a-zA-Z_-]{1,64}")


def _conv_paths(cid: str) -> tuple[Path, Path, Path]:
    if not _CID_RE.fullmatch(cid or ""):
        raise ValueError(f"bad conversation id: {cid!r}")
    return (CONV_DIR / f"{cid}.meta.json",
            CONV_DIR / f"{cid}.jsonl",
            CONV_DIR / f"{cid}.context.json")


def _read_meta(cid: str) -> dict | None:
    meta_path, _, _ = _conv_paths(cid)
    try:
        return json.loads(meta_path.read_text())
    except (OSError, json.JSONDecodeError):
        return None


def _write_meta(meta: dict) -> None:
    meta_path, _, _ = _conv_paths(meta["id"])
    tmp = meta_path.with_suffix(".tmp")
    tmp.write_text(json.dumps(meta))
    tmp.replace(meta_path)


def conv_create(model: str, title: str = "") -> dict:
    CONV_DIR.mkdir(parents=True, exist_ok=True)
    cid = uuid4().hex[:12]
    now = time.time()
    meta = {"id": cid, "title": title, "model": model,
            "created_at": now, "updated_at": now, "events": 0,
            "cost_usd": 0.0, "tokens_in": 0, "tokens_out": 0}
    _write_meta(meta)
    _conv_paths(cid)[1].touch()
    return meta


def conv_list() -> list[dict]:
    metas = []
    try:
        paths = list(CONV_DIR.glob("*.meta.json"))
    except OSError:
        return []
    for p in paths:
        try:
            metas.append(json.loads(p.read_text()))
        except (OSError, json.JSONDecodeError):
            continue
    metas.sort(key=lambda m: m.get("updated_at") or 0, reverse=True)
    return metas


def conv_get(cid: str) -> dict | None:
    meta = _read_meta(cid)
    if meta is None:
        return None
    _, events_path, _ = _conv_paths(cid)
    events = []
    try:
        for line in events_path.read_text().splitlines():
            if line.strip():
                events.append(json.loads(line))
    except (OSError, json.JSONDecodeError):
        pass
    return {"meta": meta, "events": events, "todos": conv_load_todos(cid)}


def conv_delete(cid: str) -> bool:
    if _read_meta(cid) is None:
        return False
    for p in _conv_paths(cid):
        p.unlink(missing_ok=True)
    try:
        _todo_path(cid).unlink(missing_ok=True)
    except ValueError:
        pass
    return True


def conv_append_event(cid: str, event: dict) -> int | None:
    """Append a transcript event; returns its seq (1-based per conversation).

    The first user message also becomes the conversation title.
    """
    meta = _read_meta(cid)
    if meta is None:
        return None
    seq = int(meta.get("events") or 0) + 1
    meta["events"] = seq
    meta["updated_at"] = time.time()
    if not meta.get("title") and event.get("type") == "user":
        lines = str(event.get("text") or "").strip().splitlines()
        files = event.get("files") or []
        fallback = files[0]["name"] if files else "Conversation"
        meta["title"] = (lines[0][:80] if lines else "") or fallback
    rec = {"seq": seq, "ts": meta["updated_at"]}
    rec.update({k: v for k, v in event.items()
                if k not in ("seq", "conversation_id")})
    _, events_path, _ = _conv_paths(cid)
    try:
        with events_path.open("a") as f:
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")
        _write_meta(meta)
    except OSError:
        return None
    return seq


def conv_set_model(cid: str, model: str) -> None:
    meta = _read_meta(cid)
    if meta is not None:
        meta["model"] = model
        _write_meta(meta)


def conv_add_usage(cid: str, usd: float, tokens_in: int, tokens_out: int,
                   ctx_tokens: int | None = None) -> None:
    """Accumulate spend into the conversation's meta file.

    `ctx_tokens` marks the request as carrying the conversation's own
    context — its prompt size is stored as the conv's last known context
    occupancy. Auxiliary calls (subagents, janitor, compact, verify) pass
    nothing so their prompts don't overwrite it.
    """
    try:
        meta = _read_meta(cid)
    except ValueError:
        return
    if meta is None:
        return
    meta["cost_usd"] = round(float(meta.get("cost_usd") or 0) + usd, 6)
    meta["tokens_in"] = int(meta.get("tokens_in") or 0) + tokens_in
    meta["tokens_out"] = int(meta.get("tokens_out") or 0) + tokens_out
    if ctx_tokens is not None:
        meta["ctx_tokens"] = ctx_tokens
    try:
        _write_meta(meta)
    except OSError:
        pass


def conv_load_context(cid: str) -> list | None:
    _, _, ctx_path = _conv_paths(cid)
    try:
        messages = json.loads(ctx_path.read_text())
    except (OSError, json.JSONDecodeError):
        return None
    if not isinstance(messages, list) or not messages:
        return None
    migrate_legacy_tools(messages)
    return messages


def conv_save_context(cid: str, messages: list) -> None:
    _, _, ctx_path = _conv_paths(cid)
    try:
        tmp = ctx_path.with_suffix(".tmp")
        tmp.write_text(json.dumps(messages))
        tmp.replace(ctx_path)
    except OSError as e:
        print(f"[gut] context save failed for {cid}: {e}")


def _todo_path(cid: str) -> Path:
    if not _CID_RE.fullmatch(cid or ""):
        raise ValueError(f"bad conversation id: {cid!r}")
    return CONV_DIR / f"{cid}.todos.json"


def conv_load_todos(cid: str) -> list[dict]:
    try:
        items = json.loads(_todo_path(cid).read_text())
    except (OSError, json.JSONDecodeError, ValueError):
        return []
    if not isinstance(items, list):
        return []
    return [i for i in items
            if isinstance(i, dict) and str(i.get("content") or "").strip()]


def conv_save_todos(cid: str, items: list[dict]) -> None:
    try:
        p = _todo_path(cid)
        tmp = p.with_suffix(".tmp")
        tmp.write_text(json.dumps(items))
        tmp.replace(p)
    except (OSError, ValueError) as e:
        print(f"[gut] todos save failed for {cid}: {e}")


def conv_ctx_estimate(cid: str, meta: dict) -> dict:
    """Context-window occupancy for the conversation meter.

    The running conversation reports the last billed prompt_tokens with its
    live split; an idle one gets a stored-context estimate normalized to the
    last real prompt size saved in its meta — `estimated` marks when that
    anchor is missing and the raw chars/4 guess is all we have.
    """
    if state.running and state.conversation_id == cid and state.ctx_used:
        out = {"tokens": state.ctx_used, "parts": state.ctx_parts,
               "estimated": False}
        limit = state.ctx_limit.get(state.model)
        if limit:
            out["limit"] = limit
            out["compact_at"] = int(limit * COMPACT_RATIO)
        return out
    messages = conv_load_context(cid)
    if not messages:
        return {"tokens": 0, "parts": {}, "estimated": True}
    raw = ctx_breakdown(messages)
    total = sum(raw.values())
    anchor = meta.get("ctx_tokens")
    if anchor:
        scale = anchor / max(1, total)
        return {"tokens": anchor, "estimated": False,
                "parts": {k: round(v * scale) for k, v in raw.items()}}
    return {"tokens": round(total), "estimated": True,
            "parts": {k: round(v) for k, v in raw.items()}}


def sanitize_context(messages: list) -> None:
    """Repair a stored context so the API accepts it on resume.

    A run stopped mid-turn can leave an assistant tool_call without its tool
    result (Anthropic rejects those) or an orphan tool result. Fill missing
    results with a placeholder and drop orphans, in place.
    """
    i = 0
    while i < len(messages):
        msg = messages[i]
        calls = (msg.get("tool_calls") or []) if msg.get("role") == "assistant" else []
        if not calls:
            i += 1
            continue
        want = {tc.get("id") for tc in calls}
        j = i + 1
        while j < len(messages) and messages[j].get("role") == "tool":
            if messages[j].get("tool_call_id") in want:
                want.discard(messages[j].get("tool_call_id"))
                j += 1
            else:
                del messages[j]  # orphan result for a call we no longer have
        for k, tid in enumerate(sorted(t for t in want if t)):
            messages.insert(j + k, {
                "role": "tool", "tool_call_id": tid,
                "content": "(interrupted before a result was produced)"})
        i = j + len(want)


# ── LiteLLM ─────────────────────────────────────────────────────────────────

async def provision_key() -> str:
    """Reuse or mint a 'gut-agent' virtual key so spend is attributed to it.

    The key is persisted on the desktop-home volume so lifetime spend survives
    container restarts. Falls back to the master key if LiteLLM is unreachable.
    """
    if not LITELLM_MASTER_KEY:
        return ""
    for attempt in range(30):
        try:
            async with httpx.AsyncClient(timeout=15) as c:
                if KEY_FILE.exists():
                    saved = KEY_FILE.read_text().strip()
                    if saved:
                        r = await c.get(
                            f"{LITELLM_URL}/key/info",
                            params={"key": saved},
                            headers={"Authorization": f"Bearer {LITELLM_MASTER_KEY}"},
                        )
                        if r.status_code == 200:
                            print("[gut] reusing persisted virtual key")
                            return saved
                r = await c.post(
                    f"{LITELLM_URL}/key/generate",
                    headers={"Authorization": f"Bearer {LITELLM_MASTER_KEY}"},
                    json={"key_alias": "gut-agent"},
                )
                r.raise_for_status()
                key = r.json()["key"]
                try:
                    KEY_FILE.write_text(key)
                except OSError:
                    pass
                return key
        except Exception as e:  # LiteLLM may still be migrating on first boot
            print(f"[gut] key generate attempt {attempt + 1} failed: {e}")
            await asyncio.sleep(5)
    print("[gut] falling back to master key")
    return LITELLM_MASTER_KEY


async def litellm_spend() -> float | None:
    """Lifetime spend for the agent key, from LiteLLM's spend tracking API."""
    try:
        async with httpx.AsyncClient(timeout=15) as c:
            r = await c.get(
                f"{LITELLM_URL}/key/info",
                params={"key": state.litellm_key},
                headers={"Authorization": f"Bearer {LITELLM_MASTER_KEY}"},
            )
            r.raise_for_status()
            info = r.json().get("info") or {}
            return float(info.get("spend") or 0)
    except Exception:
        return None


def track_cost(resp: httpx.Response) -> tuple[float, int, int]:
    """Add a response's spend to the session totals; returns the delta
    (usd, tokens_in, tokens_out) so the caller can attribute it."""
    usd = tin = tout = 0
    try:
        usd = float(resp.headers.get("x-litellm-response-cost") or 0)
        state.session_usd += usd
    except ValueError:
        pass
    try:
        usage = resp.json().get("usage") or {}
        tin = int(usage.get("prompt_tokens") or 0)
        tout = int(usage.get("completion_tokens") or 0)
        state.tokens_in += tin
        state.tokens_out += tout
    except Exception:
        pass
    return usd, tin, tout


async def push_cost() -> None:
    msg = {
        "type": "cost",
        "session_usd": round(state.session_usd, 6),
        "lifetime_usd": await litellm_spend(),
        "tokens_in": state.tokens_in,
        "tokens_out": state.tokens_out,
        "model": state.model,
    }
    # The running conversation's own spend, for clients showing per-chat usage.
    if state.running and state.conversation_id:
        try:
            meta = _read_meta(state.conversation_id) or {}
        except ValueError:
            meta = {}
        msg["conversation_id"] = state.conversation_id
        msg["conversation_usd"] = meta.get("cost_usd") or 0
    # Context-window occupancy of the running conversation — the billed
    # prompt size plus its category split; the window and the compaction
    # threshold let the client render a gauge. Auxiliary push_cost callers
    # (subagents, janitor) re-broadcast these unchanged — they describe the
    # same conversation_id.
    if state.ctx_used and state.running and state.conversation_id:
        msg["ctx_tokens"] = state.ctx_used
        msg["ctx_parts"] = state.ctx_parts
        limit = state.ctx_limit.get(state.model)
        if limit:
            msg["ctx_limit"] = limit
            msg["ctx_compact_at"] = int(limit * COMPACT_RATIO)
    await broadcast(msg)


# ── Provider keys (BYOK) ─────────────────────────────────────────────────────
# Model provider keys reach the daemon two ways:
#   * env vars on this process — compose's shared .env, or the deb's
#     /etc/gut/gut-bot.env (source "env"), and
#   * pushed from the app via POST /api/keys — persisted in
#     PROVIDER_KEYS_FILE on the data volume (source "pushed").
# Pushed keys win over env ones for the same provider. For every effective
# key the daemon registers that provider's models in LiteLLM's DB-backed
# model store (store_model_in_db) under deployment ids prefixed "gut-", so
# keys take effect with no service restarts and the shipped config can stay
# keyless. A static config.yaml entry under the same model_name suppresses
# the managed deployment — hand-written config always wins.
PROVIDER_MODELS = {
    "ANTHROPIC_API_KEY": [
        ("claude-sonnet-4-5", "anthropic/claude-sonnet-4-5-20250929"),
        ("claude-haiku-4-5", "anthropic/claude-haiku-4-5-20251001"),
    ],
    "OPENAI_API_KEY": [
        ("gpt-5", "openai/gpt-5"),
        ("gpt-5-mini", "openai/gpt-5-mini"),
        ("gpt-4.1-mini", "openai/gpt-4.1-mini"),
        ("gpt-4o", "openai/gpt-4o"),
    ],
    "GEMINI_API_KEY": [
        ("gemini-2.5-pro", "gemini/gemini-2.5-pro"),
        ("gemini-2.5-flash", "gemini/gemini-2.5-flash"),
        ("gemini-2.5-flash-lite", "gemini/gemini-2.5-flash-lite"),
    ],
    "DEEPSEEK_API_KEY": [
        ("deepseek-chat", "deepseek/deepseek-chat"),
    ],
    "OPENROUTER_API_KEY": [
        ("openrouter/claude-sonnet-4.5",
         "openrouter/anthropic/claude-sonnet-4.5"),
        ("openrouter/gpt-5", "openrouter/openai/gpt-5"),
        ("openrouter/gpt-4.1-mini", "openrouter/openai/gpt-4.1-mini"),
        ("openrouter/gemini-2.5-pro", "openrouter/google/gemini-2.5-pro"),
        ("openrouter/gemini-2.5-flash", "openrouter/google/gemini-2.5-flash"),
        ("openrouter/gemini-2.5-flash-lite",
         "openrouter/google/gemini-2.5-flash-lite"),
        ("openrouter/qwen3-vl", "openrouter/qwen/qwen3-vl-235b-a22b-instruct"),
        ("openrouter/qwen3-vl-30b",
         "openrouter/qwen/qwen3-vl-30b-a3b-instruct"),
        # Text-only — cheap SUBAGENT_MODEL, not for driving the desktop.
        ("openrouter/deepseek-v3.2", "openrouter/deepseek/deepseek-v3.2"),
    ],
}
# The server providers' "keys" are base URLs (plus an optional API key for
# OpenAI-compatible endpoints) and their model lists are whatever those
# servers expose — discovered live in reconcile_models instead of static
# entries above.
PROVIDER_KEYS = tuple(PROVIDER_MODELS) + (
    "OLLAMA_API_BASE", "OPENAI_COMPAT_BASE", "OPENAI_COMPAT_API_KEY",
    # Keyed SERP APIs for web_search — not model providers, but they ride
    # the same pushed-key channel (provider_keys.json applies live, env is
    # the fallback) since scraping folds the moment a host IP gets flagged.
    "TAVILY_API_KEY", "BRAVE_API_KEY", "SERPER_API_KEY")
PROVIDER_KEYS_FILE = GUT_DATA_DIR / "provider_keys.json"
MANAGED_PREFIX = "gut-"


def load_provider_keys() -> dict:
    try:
        data = json.loads(PROVIDER_KEYS_FILE.read_text())
        return {k: str(v).strip() for k, v in data.items()
                if k in PROVIDER_KEYS and str(v).strip()}
    except (OSError, json.JSONDecodeError, AttributeError):
        return {}


def store_provider_keys(keys: dict) -> None:
    PROVIDER_KEYS_FILE.parent.mkdir(parents=True, exist_ok=True)
    tmp = PROVIDER_KEYS_FILE.with_suffix(".tmp")
    tmp.write_text(json.dumps(keys))
    tmp.chmod(0o600)
    tmp.replace(PROVIDER_KEYS_FILE)


def effective_provider_keys() -> dict:
    """{env key: (value, "env"|"pushed")} — pushed keys override env."""
    eff = {k: (os.environ[k].strip(), "env") for k in PROVIDER_KEYS
           if os.environ.get(k, "").strip()}
    for k, v in load_provider_keys().items():
        eff[k] = (v, "pushed")
    return eff


# Fallback when the server can't report modality — for Ollama the reliable
# signal is 'clip' in details.families (the vision encoder it bundles into
# multimodal models); an OpenAI-compatible /models has no such field.
VISION_NAME_RE = re.compile(
    r"llava|moondream|minicpm-v|qwen[\d.]*-?vl|vision|gemma3|"
    r"mistral-small|granite|bakllava", re.I)


def normalize_server_base(raw: str, default_port: int = 0,
                          default_path: str = "",
                          keep_path: bool = False) -> str:
    """'192.168.1.5' → 'http://192.168.1.5[:default_port][default_path]'.

    An explicit scheme/port wins; 'https://' with no port keeps 443 (a
    reverse proxy), plain http gets default_port. '' when there's no usable
    host."""
    b = (raw or "").strip().rstrip("/")
    if not b:
        return ""
    if "://" not in b:
        b = f"http://{b}"
    try:
        u = urlsplit(b)
        host = u.hostname or ""
        if (u.scheme not in ("http", "https") or u.username or u.password
                or not re.fullmatch(
                    r"[A-Za-z0-9.\-_]+|[0-9a-fA-F:]*:[0-9a-fA-F:]+", host)):
            return ""
        # rstrip(':') cleans inputs like 'host:' that parse with no port.
        netloc = u.netloc.rstrip(":")
        if u.port is None and default_port and u.scheme == "http":
            netloc = f"{netloc}:{default_port}"
        u = u._replace(netloc=netloc)
        path = (u.path.rstrip("/") or default_path) if keep_path else ""
        return urlunsplit((u.scheme, u.netloc, path, "", ""))
    except ValueError:
        return ""


def normalize_ollama_base(raw: str) -> str:
    """'192.168.1.5' → 'http://192.168.1.5:11434' — path always dropped, a
    pasted one would corrupt api_base joins."""
    return normalize_server_base(raw, default_port=11434)


def normalize_compat_base(raw: str) -> str:
    """'192.168.1.5:8000' → 'http://192.168.1.5:8000/v1' — OpenAI-compatible
    servers mount at /v1; an explicit path (/openai/v1, …) wins."""
    return normalize_server_base(raw, default_path="/v1", keep_path=True)


def ollama_base() -> str:
    """The configured Ollama server URL, normalized — pushed wins over env."""
    eff = effective_provider_keys().get("OLLAMA_API_BASE")
    return normalize_ollama_base(eff[0]) if eff else ""


def compat_base() -> str:
    """Configured OpenAI-compatible server URL (…/v1), normalized."""
    eff = effective_provider_keys().get("OPENAI_COMPAT_BASE")
    return normalize_compat_base(eff[0]) if eff else ""


def compat_key() -> str:
    """Optional key for the compat server — most LAN servers need none."""
    eff = effective_provider_keys().get("OPENAI_COMPAT_API_KEY")
    return eff[0] if eff else ""


async def ollama_tags(base: str) -> list[dict]:
    """GET {base}/api/tags → [{name, vision}] — the server's pulled models."""
    async with httpx.AsyncClient(timeout=8) as c:
        r = await c.get(f"{base}/api/tags")
        r.raise_for_status()
        out = []
        for m in r.json().get("models", []):
            name = m.get("name") or m.get("model")
            if not name:
                continue
            families = (m.get("details") or {}).get("families") or []
            out.append({"name": name,
                        "vision": "clip" in families
                                  or bool(VISION_NAME_RE.search(name))})
        return sorted(out, key=lambda m: m["name"])


async def compat_models(base: str, key: str = "") -> list[dict]:
    """GET {base}/models → [{name, vision}] — the OpenAI model listing."""
    headers = {"Authorization": f"Bearer {key}"} if key else {}
    async with httpx.AsyncClient(timeout=8) as c:
        r = await c.get(f"{base}/models", headers=headers)
        r.raise_for_status()
        return sorted(
            ({"name": m["id"], "vision": bool(VISION_NAME_RE.search(m["id"]))}
             for m in r.json().get("data", []) if m.get("id")),
            key=lambda m: m["name"])


def _probe_err(e: Exception, server: str, auth_hint: bool = False) -> str:
    """A failed server probe as a friendly one-liner — httpx's str() is
    empty on timeouts. `server` is a noun phrase like 'an Ollama server';
    auth_hint makes 401/403 say 'key rejected' instead of 'wrong server'."""
    if isinstance(e, httpx.HTTPStatusError):
        c = e.response.status_code
        if auth_hint and c in (401, 403):
            return f"answered HTTP {c} — key rejected"
        return f"answered HTTP {c} — is this {server}?"
    if isinstance(e, httpx.TimeoutException):
        return "timed out"
    if isinstance(e, httpx.ConnectError):
        return "no answer — is it running there?"
    return str(e) or type(e).__name__


# Provider values that are server addresses, not secrets — normalized on
# store, and the hint shown when the input can't be parsed into a host.
SERVER_BASE_KEYS = {
    "OLLAMA_API_BASE": (normalize_ollama_base,
                        "e.g. 192.168.1.5 or http://192.168.1.5:11434"),
    "OPENAI_COMPAT_BASE": (normalize_compat_base,
                           "e.g. 192.168.1.5:8000 or http://host:1234/v1"),
}


async def litellm_admin(method: str, path: str, **kw) -> httpx.Response:
    async with httpx.AsyncClient(timeout=15) as c:
        r = await c.request(
            method, f"{LITELLM_URL}{path}",
            headers={"Authorization": f"Bearer {LITELLM_MASTER_KEY}"}, **kw)
        r.raise_for_status()
        return r


def _managed_id(m: dict) -> str:
    """Deployment id when `m` is one of ours (DB-stored), else ''."""
    info = m.get("model_info") or {}
    mid = str(info.get("id") or "")
    if mid.startswith(MANAGED_PREFIX) or info.get("managed_by") == "gut":
        return mid
    return ""


def _deployment_id(model_name: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "-", model_name).strip("-").lower()
    return f"{MANAGED_PREFIX}{slug}"


async def litellm_deployments() -> list[dict]:
    r = await litellm_admin("GET", "/model/info")
    data = r.json()
    items = data.get("data", []) if isinstance(data, dict) else data
    return [m for m in items if isinstance(m, dict)]


_reconcile_lock = asyncio.Lock()


async def reconcile_models() -> None:
    """Sync LiteLLM's deployments with the effective provider keys.

    Deletes every managed (gut-*) deployment then re-adds the desired set —
    cheap, idempotent, and self-healing when keys change or the LiteLLM DB
    was reset. A model_name already served by a static (unmanaged) config
    entry is left to it.
    """
    # model_name -> (litellm_params, extra model_info)
    desired: dict[str, tuple[dict, dict]] = {}
    for env_key, models in PROVIDER_MODELS.items():
        eff = effective_provider_keys().get(env_key)
        if not eff:
            continue
        for name, litellm_model in models:
            desired[name] = ({"model": litellm_model, "api_key": eff[0],
                              "max_tokens": 8192}, {})
    # Ollama's stored value is the server's base URL; its models are whatever
    # that server has pulled. A dead server means no ollama models this
    # pass — never a failed sync for the other providers.
    base = ollama_base()
    if base:
        try:
            for m in await ollama_tags(base):
                desired[f"ollama/{m['name']}"] = (
                    {"model": f"ollama/{m['name']}", "api_base": base,
                     "max_tokens": 8192},
                    {"supports_vision": m["vision"]})
        except Exception as e:
            print(f"[gut] ollama server {base} unreachable: "
                  f"{str(e) or type(e).__name__}")
    # Same for a generic OpenAI-compatible server (vLLM, LM Studio,
    # llama.cpp…): models from GET {base}/models, routed through LiteLLM's
    # openai provider with api_base. No configured key → a dummy Bearer,
    # which these servers ignore but LiteLLM's client requires.
    cbase, ckey = compat_base(), compat_key()
    if cbase:
        try:
            for m in await compat_models(cbase, ckey):
                desired[f"compat/{m['name']}"] = (
                    {"model": f"openai/{m['name']}", "api_base": cbase,
                     "api_key": ckey or "gut-local", "max_tokens": 8192},
                    {"supports_vision": m["vision"]})
        except Exception as e:
            print(f"[gut] openai-compat server {cbase} unreachable: "
                  f"{str(e) or type(e).__name__}")
    async with _reconcile_lock:
        current = await litellm_deployments()
        static_names = {m.get("model_name") for m in current
                        if m.get("model_name") and not _managed_id(m)}
        for m in current:
            mid = _managed_id(m)
            if mid:
                await litellm_admin("POST", "/model/delete",
                                    json={"id": mid})
        added = []
        for name, (params, info) in desired.items():
            if name in static_names:
                continue
            await litellm_admin("POST", "/model/new", json={
                "model_name": name,
                "litellm_params": params,
                "model_info": {"id": _deployment_id(name),
                               "managed_by": "gut", **info},
            })
            added.append(name)
        # If the configured default model doesn't exist (no key for it),
        # fall back to the first model that does — a task sent before the
        # user picks a model shouldn't hit a missing deployment.
        available = static_names | set(added)
        if available and state.model not in available:
            state.model = added[0] if added else sorted(available)[0]
            await push_status()


async def model_sync_loop() -> None:
    """Retry reconcile until LiteLLM accepts it, then stay quiet. Key pushes
    reset state.models_synced to trigger another pass."""
    while True:
        if not state.models_synced:
            try:
                await reconcile_models()
                state.models_synced = True
                print("[gut] litellm models synced")
            except Exception as e:
                print(f"[gut] model sync failed (retrying in 20s): {e}")
        await asyncio.sleep(20)


# ── WebSocket plumbing ───────────────────────────────────────────────────────

# Event types that make up the persisted transcript — status/cost/hello are
# ephemeral channel noise and are not recorded.
TRANSCRIPT_TYPES = frozenset({
    "user", "agent_msg", "done", "file", "image",
    "thought", "action", "action_result", "question", "error",
    "subagent", "cleanup", "plan", "compact"})


def record_event(msg: dict) -> None:
    """Persist a transcript event into the running conversation, then tag the
    outgoing message with (conversation_id, seq) so clients can dedupe live
    events against fetched history."""
    if msg.get("type") not in TRANSCRIPT_TYPES or "seq" in msg:
        return
    cid = state.conversation_id
    if not cid or not state.running:
        return
    seq = conv_append_event(cid, msg)
    if seq is not None:
        msg["conversation_id"] = cid
        msg["seq"] = seq


async def broadcast(msg: dict) -> None:
    record_event(msg)
    data = json.dumps(msg)
    dead = []
    for ws in list(state.clients):
        try:
            await ws.send_text(data)
        except Exception:
            dead.append(ws)
    for ws in dead:
        state.clients.discard(ws)


async def broadcast_conv(msg: dict, conv_id: str | None) -> None:
    """Record a transcript event into a specific conversation, then send.
    Helper events belong to the conversation that spawned the helper — not
    whatever run (if any) is active when they fire — so they bypass
    record_event's state.conversation_id tagging."""
    if (conv_id and msg.get("type") in TRANSCRIPT_TYPES
            and "seq" not in msg):
        seq = conv_append_event(conv_id, msg)
        if seq is not None:
            msg["conversation_id"] = conv_id
            msg["seq"] = seq
    await broadcast(msg)


async def push_status() -> None:
    await broadcast({"type": "status", "state": state.phase,
                     "model": state.model,
                     "run_started":
                         state.run_start if state.running else None,
                     "conversation_id":
                         state.conversation_id if state.running else None})


# ── Desktop control ──────────────────────────────────────────────────────────

def capture_frame(region: tuple[int, int, int, int] | None = None,
                  max_pixels: int | None = None) -> tuple[str, str]:
    """Return (base64 JPEG, content hash) of the current screen. `region` is
    a real-pixel (left, upper, right, lower) crop for zoomed detail reads —
    it may be upscaled for legibility and does NOT move the coordinate
    bookkeeping (clicks keep mapping through the last full frame).
    `max_pixels` overrides the pixel cap for cheaper low-res auto frames."""
    subprocess.run(["scrot", "-o", str(SHOT_PATH)], check=True)
    img = Image.open(SHOT_PATH).convert("RGB")
    if region:
        box = (max(0, region[0]), max(0, region[1]),
               min(img.width, region[2]), min(img.height, region[3]))
        if box[2] <= box[0] or box[3] <= box[1]:
            raise ValueError("region is outside the screen")
        img = img.crop(box)
    # Vision APIs silently rescale images past their limits (Anthropic: >1568px
    # long edge or >~1.15MP) and the model then emits coordinates in that
    # rescaled space — clicks land off-target. Stay under both limits. Full
    # frames are never enlarged; crops exist to be read, so upscaling a small
    # region is the point (bounded by the same caps).
    pixels = max_pixels or SCREENSHOT_MAX_PIXELS
    scale = min(
        SCREENSHOT_MAX_EDGE / max(img.size),
        (pixels / (img.width * img.height)) ** 0.5,
    )
    scale = min(scale, 8.0) if region else min(scale, 1.0)
    if abs(scale - 1.0) > 1e-3:
        img = img.resize((max(1, round(img.width * scale)),
                          max(1, round(img.height * scale))))
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=80)
    if not region:
        state.coord_scale = 1.0 / scale
        state.shot_size = img.size
    data = buf.getvalue()
    return base64.b64encode(data).decode(), hashlib.sha256(data).hexdigest()


def screenshot_block(force: bool = False) -> dict | None:
    """Fresh screenshot block, or None if the frame is byte-identical to the
    last one sent (dedup — no point re-billing the model for a static screen).

    Non-forced (auto-attached) frames honor SCREENSHOT_AUTO_PIXELS — a
    cheaper, lower-res look for routine change checks; force=True keeps the
    full cap for explicit screenshot calls and run starts."""
    b64, h = capture_frame(
        max_pixels=None if force else (SHOT_AUTO_PIXELS or None))
    if not force and h == state.frame_hash:
        return None
    state.frame_hash = h
    return {
        "type": "image_url",
        "image_url": {"url": f"data:image/jpeg;base64,{b64}"},
    }


def coords_normalized() -> bool:
    """True when the active model emits 0-1000 normalized coordinates."""
    m = state.model.lower()
    return any(p in m for p in NORMALIZED_COORD_MODELS)


def coord_prompt_for(model: str) -> str:
    """The coordinate-convention blurb baked into the system prompt for a
    given model."""
    m = model.lower()
    if not any(p in m for p in NORMALIZED_COORD_MODELS):
        return COORD_PROMPT_PIXEL
    return COORD_PROMPT_NORM_YX if "gemini" in m else COORD_PROMPT_NORM


def swap_coord_prompt(messages: list, old_model: str) -> None:
    """Rewrite the coordinate convention in the stored system prompt after a
    mid-run model switch — a pixel model escalated from a normalized-grid
    model (or vice versa) would otherwise aim clicks in the wrong space."""
    want, old = coord_prompt_for(state.model), coord_prompt_for(old_model)
    if want == old or not messages:
        return
    content = messages[0].get("content")
    if isinstance(content, str):
        messages[0]["content"] = content.replace(old, want)
    elif isinstance(content, list):
        for b in content:
            if isinstance(b.get("text"), str) and old in b["text"]:
                b["text"] = b["text"].replace(old, want)
                break


def to_real_xy(coord) -> tuple[int, int]:
    x, y = float(coord[0]), float(coord[1])
    if coords_normalized():
        # All Gemini variants emit [y, x] on the 0-1000 grid (Google's
        # spatial convention), not just the computer-use model.
        if "gemini" in state.model.lower():
            x, y = y, x
        w, h = state.shot_size
        x, y = x * w / 1000, y * h / 1000
    s = state.coord_scale
    return round(x * s), round(y * s)


def _pos(args) -> tuple[int, int, str]:
    """Resolve tool coords to desktop px. The label echoes the model's own
    screenshot coords — showing the rescaled desktop number reads as a
    'coordinate mismatch' and models start compensating for it."""
    x, y = to_real_xy(args["coordinate"])
    sx, sy = args["coordinate"][0], args["coordinate"][1]
    return x, y, f"{sx},{sy}"


_KEYMAP = {
    "ctrl": "ctrl", "control": "ctrl", "cmd": "win", "command": "win",
    "super": "win", "meta": "win", "option": "alt", "return": "enter",
    "escape": "esc", "del": "delete", "pgup": "pageup", "pgdn": "pagedown",
    " ": "space",
}


def press_keys(spec: str) -> str:
    if "+" in spec:
        keys = [_KEYMAP.get(k.strip().lower(), k.strip().lower()) for k in spec.split("+")]
        pyautogui.hotkey(*keys)
        return f"pressed {'+'.join(keys)}"
    keys = [_KEYMAP.get(k.lower(), k.lower()) for k in spec.split()]
    pyautogui.press(keys)
    return f"pressed {' '.join(keys)}"


def _session_combo(spec: str) -> bool:
    """ctrl+alt combos that hijack or end the whole session — logout
    (delete), xkill (escape), X zap (backspace), VT switch (F-keys)."""
    sep = "+" if "+" in spec else None
    keys = {_KEYMAP.get(k.strip().lower(), k.strip().lower())
            for k in spec.split(sep)}
    if not {"ctrl", "alt"} <= keys:
        return False
    rest = keys - {"ctrl", "alt", "shift", "win"}
    return bool(rest & {"delete", "esc", "backspace"}) or \
        any(k.startswith("f") and k[1:].isdigit() for k in rest)


def type_text(text: str) -> str:
    try:
        text.encode("ascii")
        short_ascii = len(text) <= 40
    except UnicodeEncodeError:
        short_ascii = False
    if not short_ascii:
        # Long or non-ASCII text goes through the clipboard: far faster than
        # synthetic keystrokes and immune to keyboard-layout quirks.
        try:
            subprocess.run(["xclip", "-selection", "clipboard"],
                           input=text.encode(), check=True)
            pyautogui.hotkey("ctrl", "v")
            return f"pasted {len(text)} chars via clipboard"
        except (OSError, subprocess.SubprocessError):
            pass  # fall through to synthetic typing
    try:
        text.encode("ascii")
        pyautogui.write(text, interval=0.01)
    except (UnicodeEncodeError, pyautogui.PyAutoGUIException):
        # xdotool handles unicode better than pyautogui's Xlib path
        subprocess.run(["xdotool", "type", "--delay", "12", "--", text], check=False)
    return f"typed {len(text)} chars"


_ESCAPED_NL_RE = re.compile(r"\\n")
_INLINE_SCRIPT_RE = re.compile(r"python3?\s+-c\s|<<\s*'?\"?\w+")


def _repair_script_newlines(command: str) -> tuple[str, bool]:
    """Small models emit multi-line scripts as one line with literal
    backslash-n between statements (the JSON escape leaked into the
    string). Bash hands Python the two characters `\\` `n`, which is a
    SyntaxError the model then retries forty times with a different typo
    each. When a python -c / heredoc command has no real newline and two
    or more escaped ones, they were meant as line breaks — turn them into
    real ones and say so. The same repair the search tools make for
    mojibake in queries."""
    if "\n" in command or len(_ESCAPED_NL_RE.findall(command)) < 2 \
            or not _INLINE_SCRIPT_RE.search(command):
        return command, False
    return _ESCAPED_NL_RE.sub("\n", command), True


def run_command(command: str, headless: bool = False,
                helper: str = "") -> str:
    command, repaired = _repair_script_newlines(command)
    out = _run_command(command, headless, helper)
    if repaired:
        out += ("\n[note: your command arrived as one line with literal "
                "\\n sequences — they were converted to real newlines "
                "before running. Emit real line breaks (a heredoc: "
                "python3 - <<'EOF' … EOF) so this isn't needed.]")
    return out


def _run_command(command: str, headless: bool = False,
                 helper: str = "") -> str:
    # Redirect via a real file, not pipes: a backgrounded child (`foo &`)
    # inherits stdout/stderr, and communicate() would block on pipe EOF until
    # that child exits — a false "still running" timeout for every GUI launch.
    env = None
    if headless:
        # Subagent shell: no display, so GUI launches fail fast instead of
        # silently hijacking the screen the main agent is using. GUT_HELPER
        # tags the whole process tree so the end-of-run sweep doesn't kill
        # a still-working helper's processes (see _helper_claimed).
        env = {k: v for k, v in os.environ.items()
               if k not in ("DISPLAY", "XAUTHORITY", "WAYLAND_DISPLAY")}
        if helper:
            env["GUT_HELPER"] = helper
    fd, out_path = tempfile.mkstemp(prefix="gut-cmd-", suffix=".out")
    try:
        with os.fdopen(fd, "w") as f:
            p = subprocess.run(
                ["bash", "-lc", command],
                stdout=f, stderr=subprocess.STDOUT,
                timeout=COMMAND_TIMEOUT, cwd=str(HOME_DIR), env=env,
            )
        out = Path(out_path).read_text(errors="replace").strip()
        return (out or f"(exit {p.returncode}, no output)")[:4000]
    except subprocess.TimeoutExpired:
        return "command still running after timeout — it may have been left in the foreground"
    finally:
        Path(out_path).unlink(missing_ok=True)


def list_windows() -> str:
    p = subprocess.run(["wmctrl", "-l"], capture_output=True, text=True)
    return p.stdout.strip() or "(no windows)"


def focus_window(match: str) -> str:
    p = subprocess.run(["wmctrl", "-a", match], capture_output=True, text=True)
    if p.returncode != 0:
        return f"no window matching '{match}' — try list_windows"
    return f"focused window matching '{match}'"


def _unfocused_warning() -> str | None:
    """type_text/key deliver keystrokes to whatever window holds X focus.
    When that's the desktop or panel — or nothing — the input goes nowhere,
    so warn instead of letting the model assume it landed."""
    try:
        p = subprocess.run(["xdotool", "getactivewindow", "getwindowname"],
                           capture_output=True, text=True, timeout=5)
        win = p.stdout.strip() if p.returncode == 0 else ""
    except (OSError, subprocess.SubprocessError):
        win = ""
    if not win:
        return ("NOTE: no application window has focus — keystrokes may have "
                "gone nowhere. focus_window or click a field first.")
    if win.lower() in ("desktop", "xfce4-panel"):
        return (f"NOTE: focus is on '{win}', not an app — keystrokes may "
                "have gone nowhere. focus_window or click a field first.")
    return None


# ── Structured desktop control (AT-SPI + UNO) ────────────────────────────────
# Same idea as the CDP bridge below, for native apps: read the accessibility
# tree and act on named elements instead of hunting pixels; drive the live
# LibreOffice document over its UNO socket. Both helpers run under the
# system python and need the session bus — borrowed via _session_env (the
# daemon itself starts outside dbus-run-session).

def _sys_py(helper: Path, *argv: str, stdin: str = "",
            timeout: int = 20) -> str:
    if not helper.exists():
        return f"helper missing: {helper} (image/package out of date?)"
    try:
        p = subprocess.run([SYS_PY, str(helper), *argv], input=stdin,
                           capture_output=True, text=True, timeout=timeout,
                           env=_session_env())
    except subprocess.TimeoutExpired:
        return f"{helper.name} timed out — the UI may be unresponsive"
    except OSError as e:
        return f"cannot run {helper.name}: {e}"
    out = (p.stdout or "").strip()
    err = (p.stderr or "").strip()
    if out:
        if p.returncode and err:
            out += f"\n(stderr: {err[-400:]})"
        return out
    return err or f"(exit {p.returncode}, no output)"


def desktop_tree(app: str = "") -> str:
    out = _sys_py(ATSPI_HELPER, "tree", app)
    body, sep, raw = out.partition("\n@@map ")
    if not sep:
        state.a11y_map = {}
        return out
    try:
        state.a11y_map = {int(k): v for k, v in json.loads(raw).items()}
    except (ValueError, AttributeError):
        state.a11y_map = {}
    return body


def _a11y_path(ref) -> str | None:
    try:
        return state.a11y_map.get(int(ref))
    except (TypeError, ValueError):
        return None


def desktop_act(ref, action: str = "") -> str:
    path = _a11y_path(ref)
    if not path:
        return f"unknown ref {ref} — run desktop_tree first"
    return _sys_py(ATSPI_HELPER, "act", path, action)


def desktop_click(ref) -> str:
    path = _a11y_path(ref)
    if not path:
        return f"unknown ref {ref} — run desktop_tree first"
    out = _sys_py(ATSPI_HELPER, "bounds", path)
    try:
        x, y, w, h, cx, cy = (int(v) for v in out.split()[:6])
    except (ValueError, IndexError):
        return f"no bounds for ref {ref}: {out}"
    if not (w and h):
        return f"ref {ref} has no on-screen bounds — try desktop_act"
    pyautogui.click(cx, cy)
    return f"clicked ref {ref} at {cx},{cy}"


# Roles where the exact pixel matters (cursor placement, drawing, web
# content) or that are plain containers — clicking them is either
# deliberate dead-space clicking or a near miss we may snap away from.
# Snap is only allowed when the hit role is a NEUTRAL container (or
# nothing was hit at all), so a click inside a document/text/canvas never
# gets hijacked to a nearby widget.
_SNAP_NEUTRAL_ROLES = {
    "", "panel", "filler", "viewport", "layered pane", "scroll pane",
    "frame", "window", "dialog", "section", "grouping", "unknown",
    "application", "root pane", "tool bar", "menu bar", "status bar",
    "menu", "separator", "split pane", "desktop frame",
}


def _a11y_at(x: int, y: int) -> dict | None:
    """Element report at real screen point (x, y); None when CLICK_A11Y is
    off. A failed lookup returns {"error": ...} so the click result can say
    grounding is down instead of silently omitting the annotation — clicks
    still work either way."""
    if not CLICK_A11Y:
        return None
    out = _sys_py(ATSPI_HELPER, "at", str(x), str(y), timeout=15)
    try:
        return json.loads(out)
    except ValueError:
        return {"error": out.strip()[:120] or "no output"}


def _a11y_desc(n: dict) -> str:
    name = (n.get("name") or "").strip()
    role = n.get("role") or "element"
    return f'{role} "{name[:40]}"' if name else role


def _snap_target(hit: dict | None, near: dict | None) -> dict | None:
    """The nearby element to snap a missed click to — only for near misses
    onto dead space, never for clicks inside content where position matters
    and never onto giant container-level widgets."""
    if not near or near.get("dist", 1e9) > CLICK_SNAP_PX:
        return None
    b = near.get("bounds") or [0, 0, 0, 0]
    if not (b[2] and b[3]) or max(b[2], b[3]) > 600:
        return None
    if hit is not None \
            and (hit.get("role") or "").lower() not in _SNAP_NEUTRAL_ROLES:
        return None
    return near


def _click_at(args, button: str = "left", double: bool = False) -> str:
    """Pixel click with a11y grounding: reports what sits under the aim
    point, and snaps to a nearby actionable element when the aim landed on
    dead space just beside it."""
    x, y, pos = _pos(args)
    tx, ty, note = x, y, ""
    info = _a11y_at(x, y)
    if info and info.get("error"):
        note = f" — click grounding unavailable ({info['error']})"
    elif info is not None:
        hit, near = info.get("hit"), info.get("near")
        if hit and (hit.get("actions") or hit.get("editable")):
            note = f" — on {_a11y_desc(hit)}"
        elif (snap := _snap_target(hit, near)) is not None:
            bx, by, bw, bh = snap["bounds"]
            tx, ty = bx + bw // 2, by + bh // 2
            aim = _a11y_desc(hit) if hit else "empty space"
            note = f" — aimed at {aim}; snapped to {_a11y_desc(snap)}"
        elif hit:
            note = (f" — {_a11y_desc(hit)} isn't clickable there; if that "
                    "was a miss, desktop_tree gives actionable #refs")
        else:
            note = " — nothing at that point"
    if double:
        pyautogui.doubleClick(tx, ty)
        return f"double-clicked {pos}{note}"
    pyautogui.click(tx, ty, button=button)
    verb = {"left": "left-clicked", "right": "right-clicked",
            "middle": "middle-clicked"}[button]
    return f"{verb} {pos}{note}"


def desktop_type(ref, text: str) -> str:
    path = _a11y_path(ref)
    if not path:
        return f"unknown ref {ref} — run desktop_tree first"
    out = _sys_py(ATSPI_HELPER, "settext", path, text)
    if not out.startswith("error:"):
        return out
    # Not an editable node — fall back to focus + real keystrokes.
    foc = _sys_py(ATSPI_HELPER, "focus", path)
    if foc.startswith("error:"):
        return f"{out}; focus fallback also failed: {foc}"
    return type_text(text)


def office_eval(code: str) -> str:
    # Cold-starting LibreOffice can take tens of seconds on a weak VPS.
    return _sys_py(UNO_HELPER, stdin=code, timeout=90)


# ── Chrome DevTools Protocol ──────────────────────────────────────────────────
# Chrome launches with --remote-debugging-port (see the wrapper in Dockerfile),
# so the agent can drive the page via DOM refs instead of pixel hunting.

class CDPError(RuntimeError):
    pass


async def cdp_up() -> bool:
    try:
        async with httpx.AsyncClient(timeout=3) as c:
            r = await c.get(f"{CDP_HTTP}/json/version")
            return r.status_code == 200
    except httpx.HTTPError:
        return False


# Schemes this desktop can't handle — clicking them just spawns a modal
# 'Open xdg-open?' dialog that blocks the page until someone dismisses it.
BLOCKED_SCHEMES = ("tel", "mailto", "sms", "callto", "sip", "skype",
                   "facetime", "zoommtg", "msteams")


def block_external_schemes() -> None:
    """Keep tel:/mailto:/etc links from spawning 'Open xdg-open?' dialogs.

    Chrome consults protocol_handler.excluded_schemes in the profile's
    Preferences and silently ignores clicks on listed schemes — exactly what
    we want on a desktop with nothing to hand them to.
    """
    prefs = Path.home() / ".gut-chrome" / "Default" / "Preferences"
    try:
        data = json.loads(prefs.read_text()) if prefs.exists() else {}
        schemes = data.setdefault("protocol_handler", {}) \
                      .setdefault("excluded_schemes", {})
        for s in BLOCKED_SCHEMES:
            schemes[s] = True
        prefs.parent.mkdir(parents=True, exist_ok=True)
        tmp = prefs.with_name("Preferences.tmp")
        tmp.write_text(json.dumps(data))
        tmp.replace(prefs)
    except (OSError, json.JSONDecodeError):
        pass


def accept_chrome_eula() -> None:
    """Write the 'EULA Accepted' sentinel into every Chrome user-data dir.

    Chrome/Chromium 151+ shows a modal 'Additional Terms of Service' on
    first run (eula_required defaults on for Linux) and skips it when the
    sentinel exists — covering launches that miss --no-first-run, e.g. a
    bare chromium via /etc/chromium.d or the real binary re-exec'd without
    the wrapper's --user-data-dir.
    """
    for d in (".gut-chrome", ".config/google-chrome", ".config/chromium"):
        try:
            p = Path.home() / d
            p.mkdir(parents=True, exist_ok=True)
            (p / "EULA Accepted").touch(exist_ok=True)
        except OSError:
            pass


_last_browser_heal = 0.0


def _kick_browser_provision() -> None:
    """Deb installs: hand the missing/broken browser to ensure-browser.sh so
    a failed launch self-heals instead of waiting for a service restart.
    The script is idempotent and self-throttling, so kicks are cheap."""
    global _last_browser_heal
    if not Path("/opt/gut/ensure-browser.sh").exists():
        return
    now = time.monotonic()
    if now - _last_browser_heal < 120:
        return
    _last_browser_heal = now
    try:
        subprocess.Popen(["sudo", "-n", "/opt/gut/ensure-browser.sh"])
    except OSError:
        pass


async def ensure_browser(url: str | None = None) -> str:
    """Guarantee Chrome is running with CDP; returns a status line."""
    if await cdp_up():
        return "chrome already running"
    running = subprocess.run(["pgrep", "-f", "chrome|chromium"],
                             capture_output=True).returncode == 0
    if running:
        # Chrome is up but without CDP (e.g. started before the wrapper flags
        # existed) — relaunch so the browser_* tools can work.
        subprocess.run(["pkill", "-f", "chrome|chromium"], check=False)
        await asyncio.sleep(1.5)
    block_external_schemes()  # Chrome must be stopped while we edit prefs
    accept_chrome_eula()
    target = f" {shlex.quote(url)}" if url else ""
    subprocess.Popen(["bash", "-lc",
                      f"nohup google-chrome{target} >/dev/null 2>&1 &"])
    for _ in range(50):
        if await cdp_up():
            return "chrome started"
        await asyncio.sleep(0.5)
    _kick_browser_provision()
    raise CDPError(
        f"chrome did not expose CDP on :{CDP_PORT} — provisioning a "
        "working browser in the background; retry in a minute")


async def cdp_send(ws, method: str, params: dict | None = None) -> dict:
    cdp_send.next_id += 1
    await ws.send(json.dumps({"id": cdp_send.next_id, "method": method,
                              "params": params or {}}))
    while True:  # skip CDP events until our response arrives
        msg = json.loads(await asyncio.wait_for(ws.recv(), 30))
        if msg.get("id") == cdp_send.next_id:
            if "error" in msg:
                raise CDPError(f"{method}: {msg['error'].get('message')}")
            return msg.get("result") or {}


cdp_send.next_id = 0


async def cdp_page_ws_url(new_tab_url: str | None = None) -> str:
    """WS URL of the page target to drive. `new_tab_url` opens a fresh tab."""
    async with httpx.AsyncClient(timeout=5) as c:
        if new_tab_url:
            ver = (await c.get(f"{CDP_HTTP}/json/version")).json()
            async with ws_connect(ver["webSocketDebuggerUrl"]) as bws:
                tid = (await cdp_send(bws, "Target.createTarget",
                                      {"url": new_tab_url}))["targetId"]
                try:  # make the tab visible so DOM and screenshot agree
                    await cdp_send(bws, "Target.activateTarget",
                                   {"targetId": tid})
                except CDPError:
                    pass
            state.browser_ws = f"ws://localhost:{CDP_PORT}/devtools/page/{tid}"
            return state.browser_ws
        targets = (await c.get(f"{CDP_HTTP}/json/list")).json()
    pages = [t for t in targets if t.get("type") == "page"]
    if not pages:
        raise CDPError("chrome has no open pages — call open_url first")
    urls = {t["webSocketDebuggerUrl"] for t in pages}
    if state.browser_ws in urls:
        return state.browser_ws
    state.browser_ws = pages[0]["webSocketDebuggerUrl"]
    return state.browser_ws


async def cdp_eval(ws, expression: str):
    res = await cdp_send(ws, "Runtime.evaluate", {
        "expression": expression, "returnByValue": True, "awaitPromise": True})
    if res.get("exceptionDetails"):
        raise CDPError("js exception: "
                       + json.dumps(res["exceptionDetails"])[:300])
    result = res.get("result") or {}
    return result.get("value") if result.get("value") is not None \
        else result.get("description")


# Numbered-ref DOM dump: every interactive element visible in the viewport gets
# a #ref stored in window.__gutRefs for browser_click / browser_type.
_DOM_JS = r"""
(() => {
  const sel = 'a,button,input,textarea,select,option,summary,label,' +
              '[role="button"],[role="link"],[role="checkbox"],[role="tab"],' +
              '[contenteditable="true"],[onclick]';
  const vw = innerWidth, vh = innerHeight;
  window.__gutRefs = {};
  const out = [];
  let i = 0;
  for (const el of document.querySelectorAll(sel)) {
    const r = el.getBoundingClientRect();
    if (r.width < 2 || r.height < 2) continue;
    if (r.bottom < 0 || r.right < 0 || r.top > vh || r.left > vw) continue;
    const cs = getComputedStyle(el);
    if (cs.display === 'none' || cs.visibility === 'hidden') continue;
    i++;
    window.__gutRefs[i] = el;
    const txt = (el.innerText || el.value || el.placeholder ||
                 el.getAttribute('aria-label') || el.title || el.name || '')
      .trim().replace(/\s+/g, ' ').slice(0, 80);
    const cx = Math.round(r.left + r.width / 2);
    const cy = Math.round(r.top + r.height / 2);
    out.push('#' + i + ' <' + el.tagName.toLowerCase() + '>' +
             (txt ? ' "' + txt + '"' : '') + ' @(' + cx + ',' + cy + ')');
    if (i >= 200) break;
  }
  return location.href + '\n' + document.title + '\n' + out.join('\n');
})()
"""

_REF_CENTER_JS = r"""
(() => {
  const el = (window.__gutRefs || {})[%d];
  if (!el) return null;
  el.scrollIntoView({block: 'center', inline: 'center'});
  const r = el.getBoundingClientRect();
  return {x: r.left + r.width / 2, y: r.top + r.height / 2,
          href: el.href || el.getAttribute('href') || ''};
})()
"""

# Page text, for reading content (contact info, listings, articles) without
# burning screenshots on scroll-and-squint.
_TEXT_JS = r"""
(() => {
  const t = ((document.body && document.body.innerText) || '')
    .replace(/\n{3,}/g, '\n\n').trim();
  return location.href + '\n' + document.title + '\n\n' + t;
})()
"""


async def _wait_for_load(ws, timeout: float = 8.0) -> None:
    """Best-effort settle wait after a navigation (or a click that may
    trigger one). Runtime.evaluate throws while the execution context is
    being swapped mid-load — we treat that as 'still loading'."""
    await asyncio.sleep(0.3)  # give a triggered navigation a chance to start
    deadline = asyncio.get_running_loop().time() + timeout
    while asyncio.get_running_loop().time() < deadline:
        try:
            if await cdp_eval(ws, "document.readyState") == "complete":
                return
        except (CDPError, asyncio.TimeoutError):
            pass
        await asyncio.sleep(0.25)


# Outcome of the navigation that just settled: the document's HTTP status
# (Chrome 109+ exposes it on the navigation timing entry; 0 = unknown), the
# title and the final URL.
_PAGE_INFO_JS = r"""
(() => {
  const nav = performance.getEntriesByType('navigation')[0];
  return {status: (nav && nav.responseStatus) || 0,
          title: (document.title || '').trim().replace(/\s+/g, ' ').slice(0, 120),
          url: location.href};
})()
"""

_URL_GUESS_NOTE = (
    "That is {n} URLs in a row that don't resolve to a real page. Stop "
    "constructing URLs from memory — find real ones with web_search and "
    "open only URLs a tool result gave you (search results, links from "
    "fetch_url or browser_dom).")


async def _navigate(ws, url: str) -> str:
    """Page.navigate, settle, and say what came back — an error status and
    the title — so a 404 or an error page lives on in text. The frame that
    showed it is pruned from context a few steps later; the tool result is
    what the model still sees when it considers opening the same URL again.
    Repeated misses add a nudge back to search."""
    await cdp_send(ws, "Page.enable")
    nav = await cdp_send(ws, "Page.navigate", {"url": url})
    err = nav.get("errorText")
    status, title, final = 0, "", url
    if not err:
        await _wait_for_load(ws)
        try:
            info = await cdp_eval(ws, _PAGE_INFO_JS) or {}
        except CDPError:
            info = {}
        status = int(info.get("status") or 0)
        title = str(info.get("title") or "")
        final = str(info.get("url") or url)
    if err:
        out = f"FAILED to load {url}: {err}"
    else:
        out = f"opened {url}"
        if final.rstrip("/") != url.rstrip("/"):
            out += f" → {final}"
        if status >= 400:
            out += f" — HTTP {status}"
        if title:
            out += f' — title "{title}"'
    if err or status >= 400:
        state.url_fail_streak += 1
        if state.url_fail_streak >= 2:
            out += "\n" + _URL_GUESS_NOTE.format(n=state.url_fail_streak)
    else:
        state.url_fail_streak = 0
    return out


async def browser_dom() -> str:
    await ensure_browser()
    ws_url = await cdp_page_ws_url()
    async with ws_connect(ws_url, max_size=20 * 1024 * 1024) as ws:
        out = await cdp_eval(ws, _DOM_JS)
    return str(out or "(no interactive elements)")[:12000]


async def browser_text() -> str:
    await ensure_browser()
    ws_url = await cdp_page_ws_url()
    async with ws_connect(ws_url, max_size=20 * 1024 * 1024) as ws:
        out = await cdp_eval(ws, _TEXT_JS)
    return str(out or "(empty page)")[:12000]


async def browser_click(ref: int) -> str:
    await ensure_browser()
    ws_url = await cdp_page_ws_url()
    async with ws_connect(ws_url, max_size=20 * 1024 * 1024) as ws:
        pt = await cdp_eval(ws, _REF_CENTER_JS % ref)
        if not pt:
            return (f"ref #{ref} not found — run browser_dom again "
                    "(the page has changed)")
        scheme = str(pt.get("href") or "").split(":", 1)[0].lower()
        if scheme in BLOCKED_SCHEMES:
            return (f"ref #{ref} is a {scheme}: link — no handler exists on "
                    "this desktop, it would just pop a dead-end dialog. "
                    "Read its target with browser_text instead.")
        await cdp_send(ws, "Input.dispatchMouseEvent",
                       {"type": "mouseMoved", "x": pt["x"], "y": pt["y"]})
        for evt in ("mousePressed", "mouseReleased"):
            await cdp_send(ws, "Input.dispatchMouseEvent", {
                "type": evt, "x": pt["x"], "y": pt["y"],
                "button": "left", "clickCount": 1})
        await _wait_for_load(ws, timeout=4.0)  # in case the click navigated
    return f"clicked #{ref}"


async def browser_type(ref: int, text: str) -> str:
    await ensure_browser()
    ws_url = await cdp_page_ws_url()
    async with ws_connect(ws_url, max_size=20 * 1024 * 1024) as ws:
        ok = await cdp_eval(ws, (
            "(() => { const el = (window.__gutRefs || {})[%d];"
            " if (!el) return false; el.scrollIntoView({block: 'center'});"
            " el.focus(); try { el.select(); } catch (e) {}"
            " return true; })()") % ref)
        if not ok:
            return (f"ref #{ref} not found — run browser_dom again "
                    "(the page has changed)")
        await cdp_send(ws, "Input.insertText", {"text": text})
    return f"typed {len(text)} chars into #{ref}"


async def browser_eval(expression: str) -> str:
    await ensure_browser()
    ws_url = await cdp_page_ws_url()
    async with ws_connect(ws_url, max_size=20 * 1024 * 1024) as ws:
        val = await cdp_eval(ws, expression)
    return str(val)[:4000]


async def _agent_tab() -> tuple[str, str] | None:
    """(ws url, target id) of a tab this run may navigate away: the active
    agent tab if it still exists, else any page that isn't one of the user's
    pre-run tabs. None when there is nothing to reuse."""
    try:
        async with httpx.AsyncClient(timeout=5) as c:
            targets = (await c.get(f"{CDP_HTTP}/json/list")).json()
    except (httpx.HTTPError, ValueError):
        return None
    ours = [(t["webSocketDebuggerUrl"], t["id"]) for t in targets
            if t.get("type") == "page" and t.get("id") not in state.baseline_tabs]
    for ws_url, tid in ours:
        if ws_url == state.browser_ws:
            return ws_url, tid
    return ours[0] if ours else None


async def open_url(url: str, new_tab: bool = False) -> str:
    """Navigate the agent's tab to `url`; a fresh tab only on request or
    when there is none to reuse (Chrome just launched, or every open tab is
    the user's). A new tab per call used to leave dozens of pages open by
    the end of a long run — heavy on a small VPS and slow to screenshot."""
    launched = await ensure_browser() != "chrome already running"
    if launched:
        state.browser_ws = None
    tab = None
    if not new_tab:
        for _ in range(12 if launched else 1):
            tab = await _agent_tab()  # the first tab trails CDP by a beat
            if tab or not launched:
                break
            await asyncio.sleep(0.25)
    if tab is None:
        ws_url = await cdp_page_ws_url(new_tab_url="about:blank")
    else:
        ws_url, tid = tab
        state.browser_ws = ws_url
        try:  # foreground it so screenshots and the DOM agree
            async with httpx.AsyncClient(timeout=5) as c:
                ver = (await c.get(f"{CDP_HTTP}/json/version")).json()
            async with ws_connect(ver["webSocketDebuggerUrl"]) as bws:
                await cdp_send(bws, "Target.activateTarget", {"targetId": tid})
        except (httpx.HTTPError, CDPError, OSError, KeyError):
            pass
    async with ws_connect(ws_url, max_size=20 * 1024 * 1024) as ws:
        out = await _navigate(ws, url)
    return f"chrome started; {out}" if launched else out


# ── Web search & fetch ─────────────────────────────────────────────────────
# Text-only alternatives to driving Chrome: the agent searches via OpenSERP
# (OPENSERP_URL — the compose stacks run one on the internal network; unset or
# unreachable = DuckDuckGo's HTML endpoint) and reads pages over plain HTTP.
# Deliberately absent from SCREEN_TOOLS: these turns carry no screenshot.

WEB_UA = ("Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
          "(KHTML, like Gecko) Chrome/126.0.0.0 Safari/537.36")

# Content of _SKIP_TAGS is dropped entirely (boilerplate); _BLOCK_TAGS only
# force line breaks so text doesn't run together.
_SKIP_TAGS = {"script", "style", "noscript", "template", "svg", "head",
              "nav", "header", "footer", "aside", "form"}
_BLOCK_TAGS = {
    "p", "div", "br", "hr", "li", "ul", "ol", "tr", "td", "th", "table",
    "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "main",
    "blockquote", "pre", "fieldset", "figure", "figcaption",
}


class _PageText(HTMLParser):
    """Minimal HTML -> text+links extractor (stdlib, no deps)."""

    def __init__(self, base: str) -> None:
        super().__init__(convert_charrefs=True)
        self.base = base
        self.skip = 0
        self.parts: list[str] = []
        self.links: list[tuple[str, str]] = []
        self._a_href: str | None = None
        self._a_text: list[str] = []
        self.title = ""
        self._in_title = False

    def _break(self) -> None:
        if self.parts and not self.parts[-1].endswith("\n"):
            self.parts.append("\n")

    def handle_starttag(self, tag, attrs) -> None:
        if tag == "title":
            self._in_title = True
        if tag in _SKIP_TAGS:
            self.skip += 1
            return
        if self.skip:
            return
        if tag in _BLOCK_TAGS:
            self._break()
        if tag == "a":
            href = dict(attrs).get("href")
            self._a_href = urljoin(self.base, href) if href else None
            self._a_text = []

    def handle_endtag(self, tag) -> None:
        if tag == "title":
            self._in_title = False
        if tag in _SKIP_TAGS:
            self.skip = max(0, self.skip - 1)
            return
        if self.skip:
            return
        if tag in _BLOCK_TAGS:
            self._break()
        if tag == "a":
            txt = " ".join("".join(self._a_text).split())
            if self._a_href and txt and len(txt) <= 140:
                self.links.append((txt, self._a_href))
            self._a_href = None
            self._a_text = []

    def handle_data(self, data) -> None:
        if self._in_title:
            self.title += data
            return
        if self.skip:
            return
        self.parts.append(data)
        if self._a_href is not None:
            self._a_text.append(data)

    def result(self) -> tuple[str, str, list]:
        lines = [" ".join(l.split()) for l in "".join(self.parts).split("\n")]
        text = "\n".join(l for l in lines if l)
        links = [(t, u) for t, u in self.links
                 if urlparse(u).scheme in ("http", "https")]
        return text, " ".join(self.title.split()), links


class _DDGResults(HTMLParser):
    """Parser for html.duckduckgo.com result pages (the OPENSERP_URL fallback)."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.results: list[dict] = []
        self._cls = ""
        self._href: str | None = None
        self._buf: list[str] = []

    def handle_starttag(self, tag, attrs) -> None:
        cls = dict(attrs).get("class", "").split()
        if "result__a" in cls:
            self._cls, self._href, self._buf = "a", dict(attrs).get("href"), []
        elif "result__snippet" in cls:
            self._cls, self._buf = "snippet", []

    def handle_data(self, data) -> None:
        if self._cls:
            self._buf.append(data)

    def handle_endtag(self, tag) -> None:
        # Both result__a and result__snippet are anchors — inner tags like
        # <b> must not finalize the element early (they'd truncate text).
        if tag != "a" or not self._cls:
            return
        text = " ".join("".join(self._buf).split())
        if self._cls == "a":
            url = self._href or ""
            if "uddg=" in url:  # unwrap the /l/?uddg= redirect
                url = parse_qs(urlparse(url).query).get("uddg", [url])[0]
            if urlparse(url).scheme in ("http", "https"):
                self.results.append({"title": text, "url": url, "snippet": ""})
        elif self._cls == "snippet" and self.results:
            self.results[-1]["snippet"] = text
        self._cls = ""


class _SearchBlocked(Exception):
    """The engine answered with an anti-bot challenge instead of results —
    reparsing can't help, so the caller should stop the retry loop."""


async def _ddg_search(query: str, lang: str = "", region: str = "") -> list[dict]:
    data = {"q": query}
    if lang and region:  # kl = "country-language", e.g. dk-da
        data["kl"] = f"{region.lower()}-{lang.lower()}"
    async with httpx.AsyncClient(timeout=15, follow_redirects=True,
                                 headers={"User-Agent": WEB_UA}) as c:
        r = await c.post("https://html.duckduckgo.com/html/", data=data)
        r.raise_for_status()
    p = _DDGResults()
    p.feed(r.text)
    # DDG answers flagged IPs with an image CAPTCHA ("anomaly-modal") that
    # parses cleanly to zero results — distinguish it from an honest empty.
    if not p.results and ("anomaly-modal" in r.text
                          or "challenge-form" in r.text):
        raise _SearchBlocked
    return p.results


async def _tavily_search(query: str, n: int, key: str) -> list[dict]:
    async with httpx.AsyncClient(timeout=20) as c:
        r = await c.post("https://api.tavily.com/search",
                         json={"api_key": key, "query": query,
                               "max_results": n, "search_depth": "basic"})
        r.raise_for_status()
    return [{"title": str(it.get("title", "")), "url": str(it.get("url", "")),
             "snippet": str(it.get("content") or "")}
            for it in r.json().get("results", [])]


async def _brave_search(query: str, n: int,
                        lang: str, region: str, key: str) -> list[dict]:
    params: dict = {"q": query, "count": n}
    if region:
        params["country"] = region.lower()
    if lang:
        params["search_lang"] = lang.lower()
    async with httpx.AsyncClient(
            timeout=15, headers={"X-Subscription-Token": key,
                                 "Accept": "application/json"}) as c:
        r = await c.get("https://api.search.brave.com/res/v1/web/search",
                        params=params)
        r.raise_for_status()
    return [{"title": str(it.get("title", "")), "url": str(it.get("url", "")),
             "snippet": str(it.get("description") or "")}
            for it in (r.json().get("web") or {}).get("results", [])]


async def _serper_search(query: str, n: int,
                         lang: str, region: str, key: str) -> list[dict]:
    body: dict = {"q": query, "num": n}
    if region:
        body["gl"] = region.lower()
    if lang:
        body["hl"] = lang.lower()
    async with httpx.AsyncClient(timeout=15,
                                 headers={"X-API-KEY": key}) as c:
        r = await c.post("https://google.serper.dev/search", json=body)
        r.raise_for_status()
    return [{"title": str(it.get("title", "")),
             "url": str(it.get("link", "")),
             "snippet": str(it.get("snippet") or "")}
            for it in r.json().get("organic", [])]


async def _api_search(query: str, n: int,
                      lang: str, region: str) -> tuple[list, str]:
    """First configured keyed SERP API, or ([], "") when none is. Keys are
    env vars or pushed via /api/keys — effective_provider_keys merges both."""
    eff = effective_provider_keys()
    if eff.get("TAVILY_API_KEY"):
        return await _tavily_search(query, n, eff["TAVILY_API_KEY"][0]), \
            "tavily"
    if eff.get("BRAVE_API_KEY"):
        return await _brave_search(query, n, lang, region,
                                   eff["BRAVE_API_KEY"][0]), "brave"
    if eff.get("SERPER_API_KEY"):
        return await _serper_search(query, n, lang, region,
                                    eff["SERPER_API_KEY"][0]), "serper"
    return [], ""


def _openserp_engines(lang: str) -> str:
    """Engines for /mega/search: the western trio, plus the regional engine
    only for the languages it actually serves."""
    if OPENSERP_ENGINES:
        return OPENSERP_ENGINES
    engines = ["google", "bing", "duckduckgo"]
    lang = lang.upper()
    if lang in ("RU", "UK", "BE", "KK", "KY", "UZ", "TG", "HY", "AZ", "TT"):
        engines.append("yandex")
    if lang.startswith("ZH"):
        engines.append("baidu")
    return ",".join(engines)


def _clean_query(query: str) -> tuple[str, bool]:
    """Strip control characters from a search query, repairing the mojibake
    some models emit for non-ASCII — a NUL byte followed by two hex digits
    stands in for the high-byte char (\\x00f8 = ø). Returns (cleaned,
    mangled) so the caller can warn the model its query arrived broken."""
    fixed = re.sub(r"\x00([0-9a-fA-F]{2})",
                   lambda m: chr(int(m.group(1), 16)), query)
    fixed = "".join(c for c in fixed if ord(c) >= 32).strip()
    return fixed, fixed != query.strip()


async def web_search(query: str, max_results: int = 8,
                     lang: str = "", region: str = "") -> str:
    query, mangled = _clean_query(query)
    if not query:
        return ("empty or garbled query — if you sent non-ASCII it arrived "
                "as control characters; resend with plain UTF-8")
    n = max(1, min(int(max_results or 8), 15))
    lang = (lang or SEARCH_LANG).strip().upper()
    region = (region or SEARCH_REGION).strip().upper()
    results, backend = [], ""
    openserp_failed = False
    try:
        results, backend = await _api_search(query, n, lang, region)
    except Exception as e:
        print(f"[gut] api search failed: {type(e).__name__} {e}")
    if not results and OPENSERP_URL:
        # /mega/search fans out to the listed engines and merges+dedupes
        # — per-engine blocks (CAPTCHA, rate limits) don't sink the query.
        # A 502 means every engine failed at once (usually transient
        # rate-limiting), so retry once before falling back to DDG — the
        # caller's own requery with different terms is the better retry.
        # The semaphore bounds daemon-wide concurrency: openserp spawns a
        # Chrome per query, and a batch of workers each firing multi-query
        # searches at once is exactly what 502s a small instance.
        params: dict = {"text": query, "limit": n,
                        "engines": _openserp_engines(lang)}
        if lang:
            params["lang"] = lang
        if region:
            params["region"] = region
        for attempt in range(2):
            try:
                async with _search_sem():
                    async with httpx.AsyncClient(timeout=25) as c:
                        r = await c.get(f"{OPENSERP_URL}/mega/search",
                                        params=params)
                if r.status_code == 200:
                    results = [{"title": str(it.get("title", "")),
                                "url": str(it.get("url", "")),
                                "snippet": str(it.get("snippet") or "")}
                               for it in r.json().get("results", [])]
                    backend = "openserp"
                    break
                openserp_failed = True
                print(f"[gut] openserp search HTTP {r.status_code}: "
                      f"{r.text[:200]}")
            except Exception as e:
                openserp_failed = True
                # httpx timeout exceptions stringify to "" — log the class
                # so a hung openserp isn't indistinguishable from a 502.
                print(f"[gut] openserp search failed: "
                      f"{type(e).__name__} {e}")
            if attempt == 0:
                await asyncio.sleep(2)
    if not results:
        try:
            results, backend = await _ddg_search(query, lang, region), \
                "duckduckgo"
        except _SearchBlocked:
            print(f"[gut] duckduckgo anti-bot challenge — query '{query}'")
            out = ("search is blocked: duckduckgo served an anti-bot "
                   "challenge"
                   + (" and openserp is failing" if openserp_failed else "")
                   + " — rephrasing won't help; fetch known pages directly "
                     "with fetch_url (browser tools if you have them)")
            return _mangled_note(out, mangled)
        except Exception as e:
            out = f"search failed ({e}) — use the browser tools instead"
            return _mangled_note(out, mangled)
    if not results:
        out = ("no results — try rephrasing, a different lang/region, or "
               "fetch a known page directly with fetch_url")
    else:
        lines = [f"{i}. {r['title']}\n   {r['url']}"
                 + (f"\n   {r['snippet']}" if r["snippet"] else "")
                 for i, r in enumerate(results[:n], 1)]
        out = f"results for '{query}' via {backend}:\n" + "\n".join(lines)
    return _mangled_note(out, mangled)


def _mangled_note(out: str, mangled: bool) -> str:
    """When a query arrived with control characters, the search ran on the
    cleaned text — tell the model so it resends proper UTF-8 (æ/ø/å) instead
    of retrying blind when the results look wrong."""
    if not mangled:
        return out
    return out + ("\n\n[note: your query contained control characters — "
                  "mangled unicode, likely æ/ø/å. Emit the characters "
                  "directly in the query; it was cleaned before searching, "
                  "so resend it correctly if these results look off.]")


async def _openserp_extract(url: str, cap: int) -> str | None:
    """Render `url` through openserp's bundled Chrome and return the page as
    markdown — rescues JS-heavy pages (SPAs) that a plain GET reads as an
    empty shell. None when unavailable or extraction failed."""
    if not OPENSERP_URL:
        return None
    try:
        async with httpx.AsyncClient(timeout=45) as c:
            r = await c.get(f"{OPENSERP_URL}/extract",
                            params={"url": url, "mode": "auto",
                                    "format": "markdown"})
        if r.status_code != 200:
            return None
        body = r.text.strip()
        if body.startswith("{"):
            # JSON envelope ({page_content, metadata}) or an error payload.
            try:
                data = json.loads(body)
            except json.JSONDecodeError:
                data = None
            if isinstance(data, dict):
                if data.get("error") or \
                        (data.get("metadata") or {}).get("error"):
                    return None
                body = str(data.get("page_content") or
                           data.get("content") or "").strip()
        return body[:cap] or None
    except Exception as e:
        print(f"[gut] openserp extract failed for {url}: {e}")
        return None


async def fetch_url(url: str, max_chars: int = 6000) -> str:
    url = url.strip()
    if not re.match(r"https?://", url):
        return "url must start with http:// or https://"
    cap = max(500, min(int(max_chars or 6000), 16000))
    try:
        async with httpx.AsyncClient(timeout=20, follow_redirects=True,
                                     headers={"User-Agent": WEB_UA}) as c:
            async with c.stream("GET", url) as r:
                ctype = (r.headers.get("content-type", "")
                         .split(";")[0].strip().lower())
                enc = r.encoding or "utf-8"
                final_url = str(r.url)
                raw = b""
                async for chunk in r.aiter_bytes(65536):
                    raw += chunk
                    if len(raw) > 2_000_000:
                        break
    except Exception as e:
        rendered = await _openserp_extract(url, cap)
        if rendered:
            return f"{url}\n(rendered via openserp)\n\n{rendered}"
        return (f"fetch failed: {e} — if the page needs JS or a login, "
                "use the browser tools")
    if ctype in ("", "text/html", "application/xhtml+xml"):
        p = _PageText(final_url)
        p.feed(raw.decode(enc, errors="replace"))
        text, title, links = p.result()
        if len(text) < 400:
            # A near-empty body on a plain GET usually means a JS-rendered
            # SPA — try openserp's headless-Chrome extractor before giving
            # the model a useless shell.
            rendered = await _openserp_extract(final_url, cap)
            if rendered and len(rendered) > len(text):
                return (f"# {title or final_url}\n{final_url}\n"
                        f"(rendered via openserp)\n\n{rendered}")
        out = [f"# {title or final_url}", final_url, "", text[:cap]]
        if len(text) > cap:
            out.append(f"\n[truncated — {len(text)} chars total; refetch with "
                       "a higher max_chars or grab a section with curl]")
        seen: set[str] = set()
        link_lines = []
        for txt, href in links:
            if href in seen or len(link_lines) >= 25:
                continue
            seen.add(href)
            link_lines.append(f"  {txt} — {href}")
        if link_lines:
            out.append("\nlinks:")
            out.extend(link_lines)
        return "\n".join(out)
    if ctype.startswith("text/") or ctype in ("application/json",
                                             "application/xml"):
        return f"{final_url}\n\n" + raw[:cap].decode(enc, errors="replace")
    return (f"{final_url} is {ctype or 'unknown type'} ({len(raw)} bytes) — "
            "not readable as text; download it with run_command/curl or "
            "open it in the browser")


# Source URLs harvested from a research result ride its action_result event
# so the client's research pane can stack sources without re-parsing the
# 500-char preview it ships in `result`.
_RESULT_URL_RE = re.compile(r"https?://[^\s)\]'\"<>]+")


def _result_urls(tool: str, text: str) -> list[str]:
    if tool not in PARALLEL_TOOLS or not text:
        return []
    out, seen = [], set()
    for u in _RESULT_URL_RE.findall(text):
        u = u.rstrip(".,;:'\"")
        if u not in seen:
            seen.add(u)
            out.append(u)
            if len(out) >= 12:
                break
    return out


# Screen-mutating tools trigger one fresh screenshot per turn, attached to the
# last tool result (skipped when the frame is unchanged). See agent_loop.
SCREEN_TOOLS = {
    "click", "mouse_move", "scroll", "type_text", "key", "run_command",
    "wait", "browser_click", "browser_type", "open_url", "focus_window",
    "desktop_act", "desktop_click", "desktop_type", "office_eval",
}
# The subset whose whole point is a visible effect. A run_command or
# office_eval that leaves the screen as it was did its job in a file — an
# orchestrator assembling a deliverable runs six of those in a row, and
# counting them toward the "your actions have no effect" streak used to
# spend a rescue (escalation, or a check-in with the user) on a run that
# was working fine.
VISUAL_TOOLS = SCREEN_TOOLS - {"run_command", "office_eval", "wait"}

# Declared in the order the prompt wants them weighed: plan first, then the
# cheap text web tools, browser refs, the screen, native-app refs, pixel
# fallbacks, helpers, and the user-facing tools last. Models attend more to
# the head of a long list — this used to open with eight pixel tools and end
# with the planning ones.
TOOLS = [
    {"type": "function", "function": {
        "name": "plan",
        "description": "Post your plan and keep your checklist current — one "
                       "call does both. `summary` (markdown) is shown to the "
                       "user as a plan card: give it on the first call and "
                       "again only if the plan changes. `steps` replaces "
                       "your working checklist, pinned live above the chat "
                       "and kept across context compaction: pass the full "
                       "list every time with exactly one item in_progress "
                       "and mark steps done as you go. One step = one "
                       "target, one outcome: a research step names ONE "
                       "thing to look up (it becomes one worker), never a "
                       "list. A step only counts as done once evidence of "
                       "its kind exists — a worker report for research, a "
                       "command that ran for build, a send_file for "
                       "deliver — otherwise it is put back in_progress. "
                       "Multi-phase tasks: call this before you start "
                       "acting. task_complete bounces an unfinished "
                       "checklist — keep it truthful. Quick one-off "
                       "actions: skip it.",
        "parameters": {"type": "object", "properties": {
            "summary": {"type": "string",
                        "description": "concise plan for the user — first "
                                       "call, or when the plan changes"},
            "steps": {"type": "array", "items": {"type": "object",
                "properties": {
                    "content": {"type": "string"},
                    "status": {"type": "string",
                               "enum": ["pending", "in_progress", "done"]},
                    "kind": {"type": "string",
                             "enum": list(STEP_KINDS),
                             "description": "research = look something up "
                                            "(one worker); build = make or "
                                            "compute a file/result with "
                                            "run_command; desktop = act on "
                                            "the screen; deliver = "
                                            "send_file; decide = a judgment "
                                            "call needing no tool"}},
                "required": ["content", "status"]}}},
            "required": ["steps"]}}},
    {"type": "function", "function": {
        "name": "web_search",
        "description": "Search the web — returns numbered results with title, "
                       "URL and snippet. Fast and text-only (no browser, no "
                       "screenshots): always prefer this for finding pages or "
                       "answers online. `queries` runs several searches in "
                       "parallel in a single call — always batch independent "
                       "lookups this way instead of one call per turn.",
        "parameters": {"type": "object", "properties": {
            "query": {"type": "string",
                      "description": "one search query (or use `queries`)"},
            "queries": {"type": "array", "items": {"type": "string"},
                        "description": "up to 8 queries run in parallel in "
                                       "one call — preferred for independent "
                                       "lookups"},
            "max_results": {"type": "integer",
                            "description": "results to return, default 8"},
            "lang": {"type": "string",
                     "description": "ISO language code for the results, e.g. "
                                    "DA, EN, DE — set it to the language of "
                                    "the query (defaults to SEARCH_LANG)"},
            "region": {"type": "string",
                       "description": "ISO country code for local results, "
                                      "e.g. DK, US — set it when the query is "
                                      "location-specific (defaults to "
                                      "SEARCH_REGION)"}}}}},
    {"type": "function", "function": {
        "name": "fetch_url",
        "description": "Fetch a URL over HTTP and return the page's text plus "
                       "its links — no browser needed. Much cheaper than "
                       "browser_* for reading articles, docs, posts. "
                       "JS-heavy pages are automatically rendered through a "
                       "headless browser when one is configured; pages that "
                       "still need a login come back thin — use the browser "
                       "then.",
        "parameters": {"type": "object", "properties": {
            "url": {"type": "string"},
            "max_chars": {"type": "integer",
                          "description": "text cap, default 6000 (max 16000)"}},
            "required": ["url"]}}},
    {"type": "function", "function": {
        "name": "run_command",
        "description": f"Run a bash command in the desktop session (cwd {HOME_DIR}) — "
                       "your main tool for compute, files and data. Python 3 "
                       "is available with openpyxl, python-docx and pillow "
                       "preinstalled, plus pip for anything missing. Append "
                       "'&' when launching GUI apps so it returns immediately.",
        "parameters": {"type": "object", "properties": {
            "command": {"type": "string"}}, "required": ["command"]}}},
    {"type": "function", "function": {
        "name": "spawn_agent",
        "description": "Spawn a background helper agent on a self-contained "
                       "headless subtask — web research, file and data work. "
                       "It gets web_search, fetch_url, run_command and "
                       "send_file (no screen, no browser, no user contact) "
                       "and reports back as a message when done. You keep "
                       "working meanwhile; collect_agent waits for a report.",
        "parameters": {"type": "object", "properties": {
            "task": {"type": "string",
                     "description": "complete instructions — the helper sees "
                                    "only this, not your conversation"},
            "name": {"type": "string",
                     "description": "short label, e.g. 'visa-research'"},
            "model": {"type": "string",
                      "description": "override model (default: yours)"}},
            "required": ["task"]}}},
    {"type": "function", "function": {
        "name": "collect_agent",
        "description": "Wait for a spawned helper to finish and return its "
                       "report. With no name, waits for the next one done.",
        "parameters": {"type": "object", "properties": {
            "name": {"type": "string"}}}}},
    {"type": "function", "function": {
        "name": "open_url",
        "description": "Open a URL in Chrome (launches it if needed) in your "
                       "current tab, replacing the page there; returns the "
                       "HTTP status and title. Only open URLs a tool result "
                       "gave you — never construct them. Prefer this over "
                       "clicking the Chrome icon and typing in the address "
                       "bar.",
        "parameters": {"type": "object", "properties": {
            "url": {"type": "string"},
            "new_tab": {"type": "boolean",
                        "description": "open in a new tab and keep the "
                                       "current page (default false)"}},
            "required": ["url"]}}},
    {"type": "function", "function": {
        "name": "browser_text",
        "description": "Read the current page as text (URL, title, body "
                       "text). Much cheaper than reading screenshots — use "
                       "it to extract info, listings, contact details, etc.",
        "parameters": {"type": "object", "properties": {}}}},
    {"type": "function", "function": {
        "name": "browser_dom",
        "description": "List the current page's interactive elements (links, "
                       "buttons, inputs, ...) as numbered #refs with their "
                       "on-screen positions. Required before browser_click or "
                       "browser_type (which appear once you've called this); "
                       "re-run after navigation or DOM changes.",
        "parameters": {"type": "object", "properties": {}}}},
    {"type": "function", "function": {
        "name": "browser_click",
        "description": "Click page element #ref from the last browser_dom. "
                       "Far more reliable than pixel clicks on web pages.",
        "parameters": {"type": "object", "properties": {
            "ref": {"type": "integer",
                    "description": "element # from browser_dom"}},
            "required": ["ref"]}}},
    {"type": "function", "function": {
        "name": "browser_type",
        "description": "Focus element #ref, select any existing text, and type "
                       "into it. For dropdowns, browser_click the option instead.",
        "parameters": {"type": "object", "properties": {
            "ref": {"type": "integer"}, "text": {"type": "string"}},
            "required": ["ref", "text"]}}},
    {"type": "function", "function": {
        "name": "browser_eval",
        "description": "Evaluate a JS expression in the current page and return "
                       "the result — page state, form values, fetch(), etc.",
        "parameters": {"type": "object", "properties": {
            "expression": {"type": "string"}}, "required": ["expression"]}}},
    {"type": "function", "function": {
        "name": "screenshot",
        "description": "Capture the current screen. Pass `region` "
                       "[x, y, w, h] (in screenshot coordinates) to zoom "
                       "into a detail — the crop is enlarged for "
                       "readability. Positions in a crop are relative to "
                       "the crop: keep issuing clicks in normal "
                       "full-screen coordinates.",
        "parameters": {"type": "object", "properties": {
            "region": {"type": "array", "items": {"type": "integer"},
                       "minItems": 4, "maxItems": 4,
                       "description": "optional [x, y, w, h] crop"}}}}},
    {"type": "function", "function": {
        "name": "wait",
        "description": "Wait for the screen to change (page loads, app launches, "
                       "animations). Sleeps the given seconds, then returns a "
                       "fresh screenshot.",
        "parameters": {"type": "object", "properties": {
            "seconds": {"type": "number",
                        "description": "seconds to wait, 0.1–60"}},
            "required": ["seconds"]}}},
    {"type": "function", "function": {
        "name": "desktop_tree",
        "description": "Dump the focused window's accessibility tree — the "
                       "native-app equivalent of browser_dom: numbered #refs "
                       "with role, name, available actions and on-screen "
                       "bounds. Pass `app` (name substring) to target a "
                       "window that isn't focused. Required before "
                       "desktop_act/desktop_click/desktop_type (which appear "
                       "once you've called this); re-run after the UI "
                       "changes — refs go stale.",
        "parameters": {"type": "object", "properties": {
            "app": {"type": "string",
                    "description": "app name substring, e.g. 'libreoffice'"}}}}},
    {"type": "function", "function": {
        "name": "desktop_act",
        "description": "Perform an accessibility action on #ref from the last "
                       "desktop_tree — 'press' a button, 'activate' a menu "
                       "item, 'select' a row. Far more reliable than pixel "
                       "clicks. `action` may be omitted when the element has "
                       "only one.",
        "parameters": {"type": "object", "properties": {
            "ref": {"type": "integer"},
            "action": {"type": "string",
                       "description": "action name from the ref's [..] list"}},
            "required": ["ref"]}}},
    {"type": "function", "function": {
        "name": "desktop_type",
        "description": "Set the text of editable #ref directly (no "
                       "keystrokes); falls back to focusing the node and "
                       "typing when it isn't editable.",
        "parameters": {"type": "object", "properties": {
            "ref": {"type": "integer"}, "text": {"type": "string"}},
            "required": ["ref", "text"]}}},
    {"type": "function", "function": {
        "name": "desktop_click",
        "description": "Pixel-click the center of #ref's bounds from the last "
                       "desktop_tree — for elements with no useful action.",
        "parameters": {"type": "object", "properties": {
            "ref": {"type": "integer"}}, "required": ["ref"]}}},
    {"type": "function", "function": {
        "name": "office_eval",
        "description": "Run Python-UNO code against the LIVE LibreOffice "
                       "document — insert content, format, save, export as "
                       "PDF — without touching the GUI. Globals: `doc` (the "
                       "open document, None if none), `desktop`, `load(path)` "
                       "to open a file, `file_url`, `uno`; assign `result` "
                       "to return a value.",
        "parameters": {"type": "object", "properties": {
            "code": {"type": "string"}}, "required": ["code"]}}},
    {"type": "function", "function": {
        "name": "list_windows",
        "description": "List open windows (id, desktop, title) via wmctrl.",
        "parameters": {"type": "object", "properties": {}}}},
    {"type": "function", "function": {
        "name": "focus_window",
        "description": "Raise/focus the window whose title contains `match`.",
        "parameters": {"type": "object", "properties": {
            "match": {"type": "string"}}, "required": ["match"]}}},
    {"type": "function", "function": {
        "name": "click",
        "description": "Click at [x, y] — the pixel fallback; prefer "
                       "browser_click / desktop_act refs on pages and native "
                       "apps. `button` left (default), right or middle; "
                       "`count` 2 for a double-click.",
        "parameters": {"type": "object", "properties": {
            "coordinate": {"type": "array", "items": {"type": "integer"},
                           "minItems": 2, "maxItems": 2},
            "button": {"type": "string", "enum": ["left", "right", "middle"]},
            "count": {"type": "integer", "enum": [1, 2]}},
            "required": ["coordinate"]}}},
    {"type": "function", "function": {
        "name": "mouse_move", "description": "Move the pointer to [x, y] without clicking.",
        "parameters": {"type": "object", "properties": {
            "coordinate": {"type": "array", "items": {"type": "integer"},
                           "minItems": 2, "maxItems": 2}},
            "required": ["coordinate"]}}},
    {"type": "function", "function": {
        "name": "scroll", "description": "Scroll at [x, y] (or current position).",
        "parameters": {"type": "object", "properties": {
            "coordinate": {"type": "array", "items": {"type": "integer"},
                           "minItems": 2, "maxItems": 2},
            "direction": {"type": "string", "enum": ["up", "down", "left", "right"]},
            "amount": {"type": "integer",
                       "description": "mouse-wheel clicks, NOT pixels: 3 "
                                      f"default, ~10 is a screenful, max "
                                      f"{SCROLL_MAX_CLICKS}"}},
            "required": ["direction"]}}},
    {"type": "function", "function": {
        "name": "type_text", "description": "Type text at the current focus.",
        "parameters": {"type": "object", "properties": {
            "text": {"type": "string"}}, "required": ["text"]}}},
    {"type": "function", "function": {
        "name": "key",
        "description": "Press a key, combo, or sequence, e.g. 'enter', "
                       "'ctrl+c', 'alt+tab', 'f5', 'tab tab tab'.",
        "parameters": {"type": "object", "properties": {
            "keys": {"type": "string"}}, "required": ["keys"]}}},
    {"type": "function", "function": {
        "name": "send_message",
        "description": "Send a chat message to the user. Like a teammate: "
                       "milestones, blockers, things worth interrupting for — "
                       "not step-by-step narration.",
        "parameters": {"type": "object", "properties": {
            "text": {"type": "string"}}, "required": ["text"]}}},
    {"type": "function", "function": {
        "name": "send_file",
        "description": "Send a file from this machine to the user as a "
                       f"download in chat (reports, spreadsheets, exports, "
                       f"downloads). Paths are relative to {HOME_DIR}.",
        "parameters": {"type": "object", "properties": {
            "path": {"type": "string"},
            "note": {"type": "string",
                     "description": "short caption shown with the file"}},
            "required": ["path"]}}},
    {"type": "function", "function": {
        "name": "send_image",
        "description": "Send an image to the user, shown inline in chat. "
                       "With no path, sends a fresh screenshot of the desktop.",
        "parameters": {"type": "object", "properties": {
            "path": {"type": "string",
                     "description": "image file; omit for a live screenshot"},
            "caption": {"type": "string"}}}}},
    {"type": "function", "function": {
        "name": "ask_user",
        "description": "Pause and ask the human a question (logins, 2FA, CAPTCHAs, "
                       "ambiguous decisions). Waits for their reply.",
        "parameters": {"type": "object", "properties": {
            "question": {"type": "string"}}, "required": ["question"]}}},
    {"type": "function", "function": {
        "name": "task_complete",
        "description": "End the task. `summary` is sent to the user as your "
                       "wrap-up message and is all they see: put the actual "
                       "answer in it — values with units, source URLs "
                       "verbatim, file names — then what was done and what "
                       "to check. Never 'sources are provided'; paste them. "
                       "send_file any files the user needs first. Unfinished "
                       "checklist items or unsent files bounce the call — "
                       "reconcile and retry.",
        "parameters": {"type": "object", "properties": {
            "summary": {"type": "string"}}, "required": ["summary"]}}},
]

# Tools that only make sense once something else has happened — a ref tool
# with no refs to act on, office_eval with no document, collect_agent with
# no helper. Hidden until then: the list the model weighs is shorter, and a
# browser_click with a guessed ref before any browser_dom (seen in the wild)
# becomes impossible. Once shown, a tool stays shown for the rest of the run
# — a request that no longer declares a tool its own history calls is
# something providers may reject, and it keeps the cached prefix stable.
CONTEXT_TOOLS = {
    "browser_click": lambda: state.dom_seen,
    "browser_type": lambda: state.dom_seen,
    "desktop_act": lambda: state.tree_seen,
    "desktop_click": lambda: state.tree_seen,
    "desktop_type": lambda: state.tree_seen,
    "office_eval": lambda: OFFICE_AVAILABLE,
    "collect_agent": lambda: bool(state.subagents),
    # Orchestrator mode takes the text web tools away from the main agent
    # for good: research is a worker's job, and a tool the model can't
    # call is a rule it can't forget. (A resumed conversation whose
    # history already calls them keeps them — tools_for_run below.)
    "web_search": lambda: not AGENT_ORCHESTRATE,
    "fetch_url": lambda: not AGENT_ORCHESTRATE,
}


# office_eval drives LibreOffice through UNO and starts soffice itself when
# it isn't running (uno_eval.py) — its precondition is "installed", not
# "running". Gating on a live process hid the no-GUI document tool until
# the agent had already opened LibreOffice by hand.
OFFICE_AVAILABLE = bool(shutil.which("soffice")) or \
    os.path.exists("/usr/local/bin/soffice")


def tools_for_run(messages: list) -> list:
    """The main agent's tool list for this request — TOOLS minus the
    CONTEXT_TOOLS whose precondition hasn't been met yet in this run."""
    for name, ready in CONTEXT_TOOLS.items():
        if name not in state.tools_shown and ready():
            state.tools_shown.add(name)
    for m in messages:  # tools the stored history calls stay declared
        if m.get("role") == "assistant":
            for tc in m.get("tool_calls") or []:
                state.tools_shown.add((tc.get("function") or {}).get("name"))
    return [t for t in TOOLS
            if t["function"]["name"] not in CONTEXT_TOOLS
            or t["function"]["name"] in state.tools_shown]


# Pre-consolidation names → the current tool set. Stored conversations still
# carry them (migrated on load) and a resumed model may echo one.
def legacy_call(name: str, args: dict) -> tuple[str, dict]:
    if name in ("left_click", "right_click", "middle_click", "double_click"):
        button = "left" if name == "double_click" else name.split("_")[0]
        out = {**args, "button": button}
        if name == "double_click":
            out["count"] = 2
        return "click", out
    if name == "browser_navigate":
        return "open_url", args
    if name == "share_plan":
        return "plan", {"summary": args.get("plan", "")}
    if name == "update_todos":
        return "plan", {"steps": args.get("items")}
    return name, args


def migrate_legacy_tools(messages: list) -> None:
    """Rewrite legacy tool names/arguments inside a stored context so the
    history matches the tools the request declares."""
    for m in messages:
        if m.get("role") != "assistant":
            continue
        for tc in m.get("tool_calls") or []:
            fn = tc.get("function") or {}
            try:
                args = json.loads(fn.get("arguments") or "{}")
            except json.JSONDecodeError:
                args = {}
            name, new_args = legacy_call(fn.get("name", ""), args)
            if name != fn.get("name"):
                fn["name"] = name
                fn["arguments"] = json.dumps(new_args)


# Background helpers (spawn_agent) get a headless subset — nothing that
# touches the screen, the browser, or the user. They also can't spawn, so
# delegation never nests deeper than one level.
SUBAGENT_TOOL_NAMES = {"web_search", "fetch_url", "run_command",
                       "send_file", "task_complete"}
SUBAGENT_TOOLS = [t for t in TOOLS
                  if t["function"]["name"] in SUBAGENT_TOOL_NAMES]


def _resolve_path(path_str: str) -> Path | str:
    """Resolve a model-supplied path against the desktop home dir."""
    p = Path(path_str).expanduser()
    if not p.is_absolute():
        p = HOME_DIR / p
    try:
        p = p.resolve(strict=True)
    except OSError:
        return f"no such file: {path_str}"
    if not p.is_file():
        return f"not a regular file: {p}"
    return p


def save_attachments(files) -> tuple[list[dict], str | None]:
    """Decode client attachments and write them under ~/uploads.

    Returns (saved, error). Each saved entry: {name, size, mime, path, b64};
    on error nothing is delivered and the caller rejects the whole message so
    the client keeps its draft.
    """
    if not isinstance(files, list):
        return [], "bad attachments payload"
    if len(files) > MAX_ATTACHMENTS:
        return [], f"too many attachments — {MAX_ATTACHMENTS} max"
    try:
        UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    except OSError as e:
        return [], f"upload dir unavailable: {e}"
    saved = []
    total = 0
    for f in files or []:
        if not isinstance(f, dict):
            return [], "bad attachments payload"
        name = Path(str(f.get("name") or "file")).name.lstrip(".") or "file"
        try:
            raw = base64.b64decode(str(f.get("data") or ""), validate=True)
        except Exception:
            return [], f"could not decode attachment “{name}”"
        total += len(raw)
        if len(raw) > SEND_FILE_MAX_BYTES:
            return [], (f"“{name}” is over the "
                        f"{SEND_FILE_MAX_BYTES // int(1e6)} MB limit")
        if total > ATTACH_TOTAL_MAX_BYTES:
            return [], ("attachments total over the "
                        f"{ATTACH_TOTAL_MAX_BYTES // int(1e6)} MB limit")
        p = UPLOAD_DIR / name
        stem, suffix = p.stem, p.suffix
        n = 1
        while p.exists():
            p = UPLOAD_DIR / f"{stem}-{n}{suffix}"
            n += 1
        try:
            p.write_bytes(raw)
        except OSError as e:
            return [], f"could not save “{name}”: {e}"
        state.delivered[str(p.resolve())] = p.stat().st_mtime
        saved.append({
            "name": p.name, "size": len(raw), "path": str(p),
            "mime": str(f.get("mime") or
                        mimetypes.guess_type(p.name)[0] or
                        "application/octet-stream"),
            "b64": base64.b64encode(raw).decode()})
    return saved, None


def attach_note(saved: list[dict]) -> str:
    """Tell the model where its attachments landed on the desktop."""
    listing = "\n".join(f"- {f['path']} ({f['size']} bytes)" for f in saved)
    return f"[attached files, saved on the desktop:\n{listing}]"


def event_files(saved: list[dict]) -> list[dict]:
    """File metadata for the transcript event (no payload — the file itself
    already lives on the device)."""
    return [{"name": f["name"], "size": f["size"], "mime": f["mime"]}
            for f in saved]


def file_digest(p: Path) -> str:
    """One line on what a file actually contains — sheets and filled rows
    for a workbook, paragraphs for a document, records for JSON, a head
    for text. The mechanical counterpart to "I made the spreadsheet": it
    rides the send_file result so the agent sees an empty artifact before
    the user does, and the wrap-up ledger so the checker can too."""
    try:
        size = p.stat().st_size
    except OSError as e:
        return f"unreadable ({e})"
    suf = p.suffix.lower()
    try:
        if suf in (".xlsx", ".xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
            parts = []
            for ws in wb.worksheets:
                rows = [r for r in ws.iter_rows(values_only=True)
                        if any(c not in (None, "") for c in r)]
                cells = sum(1 for r in rows for c in r if c not in (None, ""))
                head = " | ".join(str(c) for c in (rows[0] if rows else ())
                                  if c not in (None, ""))[:80]
                parts.append(f"'{ws.title}': {len(rows)} filled rows, "
                             f"{cells} cells" + (f", header: {head}"
                                                 if head else ""))
            return f"{size} B, " + "; ".join(parts)
        if suf == ".docx":
            import docx
            d = docx.Document(str(p))
            paras = [x.text for x in d.paragraphs if x.text.strip()]
            words = sum(len(x.split()) for x in paras)
            return (f"{size} B, {len(paras)} paragraphs, {words} words, "
                    f"{len(d.tables)} tables")
        if suf == ".json":
            data = json.loads(p.read_text(errors="replace"))
            shape = (f"list of {len(data)}" if isinstance(data, list)
                     else f"object with {len(data)} keys" if isinstance(data, dict)
                     else type(data).__name__)
            return f"{size} B, {shape}: {json.dumps(data)[:120]}"
        if suf in (".csv", ".tsv", ".md", ".txt", ".html", ".xml", ".py"):
            text = p.read_text(errors="replace")
            lines = [ln for ln in text.splitlines() if ln.strip()]
            return (f"{size} B, {len(lines)} non-empty lines, head: "
                    + " ".join(text[:120].split()))
    except Exception as e:
        return f"{size} B, unreadable as {suf or 'text'} ({e})"
    return f"{size} B"


async def send_file(path_str: str, note: str = "",
                    conv: str | None = None) -> str:
    p = _resolve_path(path_str or "")
    if isinstance(p, str):
        return p
    size = p.stat().st_size
    if size > SEND_FILE_MAX_BYTES:
        return (f"file is {size / 1e6:.1f} MB — over the "
                f"{SEND_FILE_MAX_BYTES // int(1e6)} MB send limit; compress it "
                "or share a smaller artifact")
    await broadcast_conv({
        "type": "file", "name": p.name, "size": size,
        "mime": mimetypes.guess_type(p.name)[0] or "application/octet-stream",
        "note": note, "data": base64.b64encode(p.read_bytes()).decode()},
        conv)
    state.sent_files[str(p)] = p.stat().st_mtime
    digest = await asyncio.to_thread(file_digest, p)
    return f"sent {p.name} to the user — contents: {digest}"


async def send_image(path_str: str | None, caption: str = "") -> str:
    if path_str:
        p = _resolve_path(path_str)
        if isinstance(p, str):
            return p
        mime = mimetypes.guess_type(p.name)[0] or ""
        if not mime.startswith("image/"):
            return f"{p.name} is not an image — use send_file instead"
        if p.stat().st_size > SEND_FILE_MAX_BYTES:
            return f"image over the {SEND_FILE_MAX_BYTES // int(1e6)} MB send limit"
        b64, name = base64.b64encode(p.read_bytes()).decode(), p.name
    else:
        b64, _ = capture_frame()
        mime, name = "image/jpeg", "screenshot.jpg"
    await broadcast({"type": "image", "name": name, "caption": caption,
                     "mime": mime, "data": b64})
    return "image sent to the user"


async def ask_user(question: str) -> str:
    loop = asyncio.get_running_loop()
    state.pending_answer = loop.create_future()
    await broadcast({"type": "question", "text": question})
    await push_status()
    try:
        return await asyncio.wait_for(state.pending_answer, ASK_USER_TIMEOUT)
    except asyncio.TimeoutError:
        return "(no reply from user — proceed with best judgment or end the task)"
    finally:
        state.pending_answer = None
        await push_status()


# Intermediates the agent writes to *produce* deliverables — listing them
# as "unsent" just invites shipping build tooling nobody asked for.
_UNSENT_SKIP_EXT = {".py", ".sh", ".js", ".ts", ".rb", ".pl", ".ps1", ".r",
                    ".sql", ".ipynb"}


def unsent_outputs() -> list[Path]:
    """Files changed since the run started that were never sent back — the
    backstop for "agent saved a file but only told the user the path".
    Uploads are scanned recursively; home, Desktop and Downloads only
    shallowly so caches and app dirs don't count. Hidden names are skipped.
    """
    if not state.run_start:
        return []
    seen, out = set(), []
    for root, deep in ((UPLOAD_DIR, True), (HOME_DIR, False),
                     (HOME_DIR / "Desktop", False),
                     (HOME_DIR / "Downloads", False)):
        if not root.is_dir():
            continue
        try:
            paths = list(root.rglob("*")) if deep else list(root.iterdir())
        except OSError:
            continue
        for p in paths:
            if any(part.startswith(".")
                   for part in p.relative_to(root).parts):
                continue
            try:
                rp = p.resolve()
                if not rp.is_file() or rp.suffix.lower() in _UNSENT_SKIP_EXT:
                    continue
                mtime = rp.stat().st_mtime
            except OSError:
                continue
            key = str(rp)
            baseline = max(state.run_start, state.delivered.get(key, 0))
            if (key in seen or mtime <= baseline
                    or state.sent_files.get(key, 0) >= mtime):
                continue
            seen.add(key)
            out.append(rp)
    return sorted(out)


_TODO_MARKS = {"done": "☑", "in_progress": "◐", "pending": "☐"}


def render_todos(items: list[dict]) -> str:
    return "\n".join(
        f"{_TODO_MARKS.get(i['status'], '☐')} [{i.get('kind', 'research')}] "
        f"{i['content']}" for i in items)


# Evidence kinds that satisfy a step of each kind. A desktop step also
# accepts build evidence — downloads, installs and file moves are screen
# work as often as shell work, and a run_command that did the job must not
# bounce the step. decide needs nothing: it is a judgment, not an action.
_EVIDENCE_FOR = {
    "research": ("research",),
    "build": ("build",),
    "desktop": ("desktop", "build"),
    "deliver": ("deliver",),
    "decide": (),
}
_EVIDENCE_HINT = {
    "research": "spawn_agent a worker for it and collect_agent its report",
    "build": "run the script/command that produces it (run_command or "
             "office_eval) and check its output",
    "desktop": "act on the screen (or run_command) first",
    "deliver": "send_file the file first",
}


def note_evidence(kind: str) -> None:
    """Record one evidence-bearing event of `kind` for the running run."""
    if kind in state.evidence:
        state.evidence[kind] += 1


def _has_evidence(kind: str, mark: dict[str, int]) -> bool:
    return not _EVIDENCE_FOR.get(kind) or any(
        state.evidence.get(k, 0) > mark.get(k, 0)
        for k in _EVIDENCE_FOR[kind])


def _norm_step(content: str) -> str:
    return " ".join(content.lower().split())


async def update_todos(raw_items) -> str:
    """Replace the running conversation's checklist — persisted next to the
    context and pushed to clients as a live card (not a transcript event:
    the card only ever shows latest state, and replayed updates would spam
    history).

    The checklist is a contract, not a notepad, and this is where the
    daemon holds the model to it:
    * lint — a research step that enumerates three or more targets is
      bounced whole: one target per step, because each research step
      becomes one focused worker;
    * evidence gate — a step may only turn done once an evidence event of
      its kind has happened since it went in_progress (a worker report, a
      command that ran, a send_file, a screen action); otherwise it is put
      back in_progress and the result says what is missing.
    Both are bounded by GATE_MAX_BOUNCES so a model that cannot satisfy
    them still gets its list accepted, with the caveat spelled out."""
    items = []
    for it in (raw_items if isinstance(raw_items, list) else []):
        if not isinstance(it, dict):
            continue
        content = str(it.get("content") or "").strip()
        if not content:
            continue
        status = str(it.get("status") or "pending").lower()
        if status == "completed":
            status = "done"
        if status not in _TODO_MARKS:
            status = "pending"
        kind = str(it.get("kind") or "").lower()
        if kind not in STEP_KINDS:
            kind = infer_step_kind(content)
        items.append({"content": content[:200], "status": status,
                      "kind": kind})
    items = items[:TODO_MAX_ITEMS]
    enforce = state.gate_bounces < GATE_MAX_BOUNCES
    notes: list[str] = []

    # Lint: bundled research steps. The whole update is rejected — the
    # model re-posts with the step split (its next reply is pinned to
    # plan by the reconcile flag), so the checklist never holds a step
    # that one worker cannot own.
    if enforce:
        for it in items:
            targets = (enumerated_targets(it["content"])
                       if it["kind"] == "research" and it["status"] != "done"
                       else [])
            if targets:
                state.gate_bounces += 1
                state.todo_reconcile = True
                quoted = ", ".join(f"'{t}'" for t in targets)
                return (f"plan NOT accepted — the research step "
                        f"'{it['content']}' bundles {len(targets)} targets "
                        f"({quoted}). One research step per target — each "
                        "becomes one worker with one goal. Re-post the full "
                        "`steps` list with that step split.")

    # Evidence gate. Steps are matched to the previous list by wording;
    # a reworded or new step counts from the last plan call.
    prev = {_norm_step(i["content"]): i for i in state.todos}
    held: list[dict] = []
    for it in items:
        key = _norm_step(it["content"])
        before = prev.get(key)
        was = before["status"] if before else None
        if it["status"] == "done" and was != "done":
            mark = state.step_marks.get(key, state.plan_mark)
            if enforce and not _has_evidence(it["kind"], mark):
                it["status"] = "in_progress"
                held.append(it)
        if it["status"] == "in_progress" and was != "in_progress":
            state.step_marks[key] = dict(state.evidence)
    if held:
        state.gate_bounces += 1
        for it in held:
            hint = _EVIDENCE_HINT.get(it["kind"], "do the work first")
            notes.append(f"⟲ '{it['content']}' is NOT done — no "
                         f"{it['kind']} evidence since it started ({hint}); "
                         "kept in_progress")
    elif not enforce and state.gate_bounces == GATE_MAX_BOUNCES:
        state.gate_bounces += 1  # say it once
        notes.append("(plan checks exhausted for this run — the list is "
                     "accepted as posted; the wrap-up audit still applies)")

    state.todos = items
    state.plan_mark = dict(state.evidence)
    state.steps_since_todo = 0
    if state.conversation_id:
        conv_save_todos(state.conversation_id, state.todos)
    await broadcast({"type": "todos",
                     "conversation_id": state.conversation_id,
                     "items": state.todos})
    # Orchestrator mode: a research step in progress with no worker on it
    # is the step the model is about to do itself — it can't (no web
    # tools), so point it at the only path.
    if AGENT_ORCHESTRATE and any(
            i["status"] == "in_progress" and i["kind"] == "research"
            for i in items) and not any(
            e["status"] == "running" for e in state.subagents.values()):
        notes.append("→ a research step is in progress and no worker is "
                     "running: spawn_agent one now with that step's target "
                     "as its goal (several independent research steps → "
                     "several workers at once), then collect_agent.")
    out = ("checklist updated:\n"
           + (render_todos(state.todos) or "(empty — all steps done?)"))
    return out + ("\n" + "\n".join(notes) if notes else "")


async def execute_tool(name: str, args: dict,
                       agent: str | None = None) -> tuple[list | str, bool]:
    """Return (tool_result_content, task_done). `agent` set = the caller is a
    restricted agent — a background helper (headless SUBAGENT_TOOL_NAMES) or
    the cleanup janitor (desktop-only JANITOR_TOOL_NAMES)."""
    if agent is not None:
        allowed = (JANITOR_TOOL_NAMES if agent == JANITOR_NAME
                   else SUBAGENT_TOOL_NAMES)
        if name not in allowed:
            return f"{name} isn't available to the {agent} agent", False
    name, args = legacy_call(name, args)
    try:
        if name == "screenshot":
            region = args.get("region")
            if region:
                try:
                    rx, ry, rw, rh = (int(v) for v in region[:4])
                    x0, y0 = to_real_xy([rx, ry])
                    x1, y1 = to_real_xy([rx + rw, ry + rh])
                    b64, _ = capture_frame((x0, y0, x1, y1))
                except (ValueError, TypeError, IndexError,
                        subprocess.SubprocessError) as e:
                    return f"screenshot region failed: {e}", False
                return [{"type": "text", "text":
                         f"crop of region [{rx},{ry} {rw}x{rh}], enlarged "
                         "for reading — positions in this image are "
                         "crop-relative; clicks still use normal "
                         "screenshot coordinates"},
                        {"type": "image_url", "image_url": {
                            "url": f"data:image/jpeg;base64,{b64}"}}], False
            return [{"type": "text", "text": "captured screenshot"},
                    screenshot_block(force=True)], False
        if name == "wait":
            secs = max(0.1, min(float(args.get("seconds") or 1), 60))
            slept = 0.0
            while slept < secs and not state.stop:
                step = min(0.25, secs - slept)
                await asyncio.sleep(step)
                slept += step
            result = f"waited {slept:.1f}s"
        elif name == "click":
            button = str(args.get("button") or "left").lower()
            if button not in ("left", "right", "middle"):
                button = "left"
            result = await asyncio.to_thread(
                _click_at, args, button, int(args.get("count") or 1) >= 2)
        elif name == "mouse_move":
            x, y, pos = _pos(args)
            info = await asyncio.to_thread(_a11y_at, x, y)
            pyautogui.moveTo(x, y)
            note = (f" — over {_a11y_desc(info['hit'])}"
                    if info and info.get("hit") else "")
            result = f"moved to {pos}{note}"
        elif name == "scroll":
            if args.get("coordinate"):
                x, y, _ = _pos(args); pyautogui.moveTo(x, y)
            asked = abs(int(args.get("amount") or 3))
            clicks = max(1, min(asked, SCROLL_MAX_CLICKS))
            direction = args.get("direction", "down")
            if direction in ("up", "down"):
                pyautogui.scroll(clicks if direction == "up" else -clicks)
            else:
                pyautogui.hscroll(clicks if direction == "right" else -clicks)
            result = f"scrolled {direction} {clicks}"
            if asked > clicks:
                result += (f" — amount is wheel clicks, not pixels; capped "
                           f"{asked} to {clicks} (~10 clicks is a screenful). "
                           "To read a whole page use browser_text instead of "
                           "scrolling.")
        elif name == "type_text":
            result = type_text(str(args.get("text", "")))
            warn = _unfocused_warning()
            if warn:
                result = _note(result, warn)
        elif name == "key":
            spec = str(args.get("keys", ""))
            if agent == JANITOR_NAME and _session_combo(spec):
                return ("key combo blocked — session-level shortcuts are "
                        "off limits during cleanup", False)
            result = press_keys(spec)
            warn = _unfocused_warning()
            if warn:
                result = _note(result, warn)
        elif name == "run_command":
            # to_thread keeps a long command from freezing the event loop —
            # matters now that helper agents share it with the main loop.
            result = await asyncio.to_thread(
                run_command, str(args.get("command", "")),
                agent is not None, agent or "")
        elif name == "web_search":
            queries = [str(q) for q in args.get("queries") or []]
            if str(args.get("query") or "").strip():
                queries.insert(0, str(args["query"]))
            queries = list(dict.fromkeys(queries))
            if not queries:
                result = "empty query — pass `query` or `queries`"
            else:
                over = len(queries) - 8
                parts = await asyncio.gather(*(
                    web_search(q, int(args.get("max_results") or 8),
                               str(args.get("lang") or ""),
                               str(args.get("region") or ""))
                    for q in queries[:8]))
                result = "\n\n".join(parts)
                if over > 0:
                    result += (f"\n\n[{over} quer{'ies' if over > 1 else 'y'}"
                               " dropped — the cap is 8 per call]")
        elif name == "fetch_url":
            result = await fetch_url(str(args.get("url", "")),
                                     int(args.get("max_chars") or 6000))
        elif name == "browser_dom":
            result = await browser_dom()
            state.dom_seen = True
        elif name == "browser_text":
            result = await browser_text()
        elif name == "browser_click":
            result = await browser_click(int(args.get("ref", 0)))
        elif name == "browser_type":
            result = await browser_type(int(args.get("ref", 0)),
                                        str(args.get("text", "")))
        elif name == "browser_eval":
            result = await browser_eval(str(args.get("expression", "")))
        elif name == "open_url":
            result = await open_url(str(args.get("url", "")),
                                    bool(args.get("new_tab")))
        elif name == "list_windows":
            result = list_windows()
        elif name == "focus_window":
            result = focus_window(str(args.get("match", "")))
        elif name == "desktop_tree":
            result = await asyncio.to_thread(desktop_tree,
                                             str(args.get("app", "")))
            state.tree_seen = True
        elif name == "desktop_act":
            result = await asyncio.to_thread(desktop_act, args.get("ref"),
                                             str(args.get("action", "")))
        elif name == "desktop_click":
            result = await asyncio.to_thread(desktop_click, args.get("ref"))
        elif name == "desktop_type":
            result = await asyncio.to_thread(desktop_type, args.get("ref"),
                                             str(args.get("text", "")))
        elif name == "office_eval":
            result = await asyncio.to_thread(office_eval,
                                             str(args.get("code", "")))
        elif name == "send_message":
            text = str(args.get("text", "")).strip()
            if not text:
                return "empty message — nothing sent", False
            await broadcast({"type": "agent_msg", "text": text})
            result = "message sent to the user"
        elif name == "send_file":
            hconv = (state.subagents.get(agent) or {}).get("conv") \
                if agent else None
            result = await send_file(str(args.get("path", "")),
                                     str(args.get("note", "")), conv=hconv)
        elif name == "send_image":
            result = await send_image(args.get("path"),
                                      str(args.get("caption", "")))
        elif name == "spawn_agent":
            result = spawn_agent(str(args.get("task", "")),
                                 str(args.get("name", "")),
                                 str(args.get("model", "")))
        elif name == "collect_agent":
            result = await collect_agent(str(args.get("name", "")))
        elif name == "ask_user":
            return await ask_user(str(args.get("question", ""))), False
        elif name == "plan":
            summary = str(args.get("summary") or "").strip()
            if summary:
                state.plan_shared = True
                await broadcast({"type": "plan", "text": summary})
            if "steps" in args:
                result = await update_todos(args.get("steps"))
            elif summary:
                result = "checklist unchanged"
            else:
                return "empty plan — pass `steps` (and `summary`)", False
            if summary:
                result = "plan shared with the user; " + result
        elif name == "task_complete":
            undone = ([i for i in state.todos if i.get("status") != "done"]
                      if agent is None else [])
            if undone and not state.todo_nudge_done:
                state.todo_nudge_done = True
                state.todo_reconcile = True
                listing = "\n".join(f"- {i['content']}"
                                    for i in undone[:10])
                return (f"checklist still lists {len(undone)} unfinished "
                        f"item(s):\n{listing}\nPost the final list now with "
                        "plan (full `steps` — mark done what's done, keep "
                        "or drop the rest honestly), then finish any real "
                        "remaining work and call task_complete again. To "
                        "end with the list as-is, call plan unchanged "
                        "first."), False
            unsent = unsent_outputs() if agent is None else []
            if unsent and not state.output_nudge_done:
                state.output_nudge_done = True
                listing = "\n".join(f"- {p}" for p in unsent[:10])
                return (f"files changed during this task but were never "
                        f"sent to the user:\n{listing}\nThe user can't "
                        f"browse this filesystem — send_file any they need, "
                        f"then call task_complete again; or call it again "
                        f"now if none of these are deliverables"), False
            if agent is not None and agent != JANITOR_NAME:
                # Findings gate: a worker's prose report is a summary; the
                # findings file is the record the orchestrator's assembly
                # script reads. A report that cites figures/URLs with an
                # empty file forces transcription through prose — exactly
                # how URLs and decimals get dropped. Bounce once so the
                # worker writes the records it already found.
                entry = state.subagents.get(agent)
                summary = str(args.get("summary") or "")
                if entry is not None and not entry.get("findings_nudge") \
                        and _REPORT_CLAIM_RE.search(summary) \
                        and _findings_empty(worker_dir(agent)):
                    entry["findings_nudge"] = True
                    return ("your report cites figures/URLs but no "
                            "substantive findings file exists under "
                            f"{worker_dir(agent)} — the file is the record "
                            "the orchestrator builds from, your report only "
                            "summarizes it. Write each finding to "
                            "findings.json as a JSON object (value, unit, "
                            "currency, source_url, note), then call "
                            "task_complete again."), False
            return str(args.get("summary", "done")), True
        else:
            return f"unknown tool: {name}", False
    except Exception as e:
        return f"{name} failed: {e}", False

    if agent is None:
        _account_evidence(name, result)
    # A fresh screenshot is attached once per turn by agent_loop (deduped on
    # unchanged frames) — not per tool call.
    return result, False


# Tool → the step kind its successful result is evidence for. Screen
# actions other than these count as desktop evidence; reads (screenshot,
# browser_dom, desktop_tree, list_windows) prove nothing and count for
# nothing. Helpers' calls never reach here — a worker's work becomes
# orchestrator evidence only through its delivered report.
_EVIDENCE_TOOLS = {
    "run_command": "build", "office_eval": "build",
    "send_file": "deliver",
    "web_search": "research", "fetch_url": "research",
    "browser_text": "research",
}
# A command that crashed built nothing; a tool result that opens with an
# error phrase found nothing. Page text is exempt from the crash patterns —
# a fetched StackOverflow thread full of tracebacks is still research.
_BUILD_FAILED_RE = re.compile(
    r"Traceback \(most recent call last\)|command not found|"
    r"No such file or directory|ModuleNotFoundError|SyntaxError|"
    r"^\(exit [1-9]", re.M)
_TOOL_ERR_RE = re.compile(
    r"\A(?:BLOCKED:|url must start|search failed|fetch failed|no results|"
    r"empty |[^\n]{0,40}\b(?:failed|isn't available|not found)\b)")


def _account_evidence(name: str, result) -> None:
    text = result if isinstance(result, str) else " ".join(
        str(b.get("text", "")) for b in result if b.get("type") == "text")
    kind = _EVIDENCE_TOOLS.get(name)
    if kind is None:
        if name in SCREEN_TOOLS and name != "wait":
            kind = "desktop"
        else:
            return
    if name == "send_file" and not text.startswith("sent "):
        return
    if _TOOL_ERR_RE.search(text):
        return
    if kind == "build" and _BUILD_FAILED_RE.search(text[:2000]):
        return
    note_evidence(kind)


# ── Agent loop ───────────────────────────────────────────────────────────────

def assistant_to_dict(msg: dict) -> dict:
    out = {"role": "assistant", "content": msg.get("content")}
    if msg.get("tool_calls"):
        out["tool_calls"] = msg["tool_calls"]
    return out


def message_text(msg: dict) -> str:
    """Visible reply text — a plain string, or text blocks joined."""
    t = msg.get("content") or ""
    if not isinstance(t, str):
        t = " ".join(str(b.get("text", "")) for b in t
                     if isinstance(b, dict))
    return t.strip()


def reasoning_text(msg: dict) -> str:
    """Thinking produced outside the reply — LiteLLM's normalized
    reasoning_content, or thinking blocks inside content. Broadcast live and
    never persisted; keep the tail, where the freshest reasoning sits."""
    t = msg.get("reasoning_content")
    if isinstance(t, list):
        t = " ".join(str(b.get("text", "")) for b in t
                     if isinstance(b, dict))
    if not isinstance(t, str) or not t.strip():
        t = " ".join(str(b.get("thinking", ""))
                     for b in (msg.get("thinking_blocks") or [])
                     if isinstance(b, dict))
    if not t.strip():
        c = msg.get("content")
        if isinstance(c, list):
            t = " ".join(str(b.get("thinking", "")) for b in c
                         if isinstance(b, dict)
                         and b.get("type") == "thinking")
    t = t.strip()
    return t[-800:] if len(t) > 800 else t


def prune_images(messages: list, keep: int | None = None) -> None:
    """Replace all but the newest `keep` screenshots with text placeholders.

    Without this the transcript carries every frame ever taken and each
    request re-pays for all of them — input cost grows quadratically with the
    step count, and stale frames actively confuse the model. The text action
    log stays intact, so the narrative is preserved; the model can call
    screenshot anytime for a fresh look.
    """
    if keep is None:
        keep = SCREENSHOT_HISTORY
    seen = 0
    for msg in reversed(messages):
        content = msg.get("content")
        if not isinstance(content, list):
            continue
        for block in reversed(content):
            if block.get("type") != "image_url":
                continue
            seen += 1
            if seen > keep:
                block.clear()
                block.update({"type": "text",
                              "text": "[earlier image omitted]"})


PRUNE_EXEMPT_TOOLS = {"web_search"}


def prune_tool_results(messages: list, keep: int | None = None) -> None:
    """Truncate tool results older than the newest `keep` to a stub.

    browser_dom / fetch_url / desktop_tree / run_command outputs are the
    largest payloads in the transcript (up to ~12k chars each) and go stale
    within a step or two, yet every request re-bills them. The head of each
    result survives (usually enough to recall what it was); the model can
    re-run the tool if it needs the full text again.
    """
    if keep is None:
        keep = TOOL_RESULT_HISTORY
    # Search results stay whole: they are the run's leads, small (a numbered
    # URL list), and a 300-char stub keeps exactly the first hit — a run
    # once found the right page as result #2, lost it to the stub eight
    # steps later and never got back to it.
    exempt = {tc.get("id") for m in messages if m.get("role") == "assistant"
              for tc in (m.get("tool_calls") or [])
              if (tc.get("function") or {}).get("name") in PRUNE_EXEMPT_TOOLS}
    seen = 0
    for msg in reversed(messages):
        if msg.get("role") != "tool":
            continue
        seen += 1
        if seen <= keep or msg.get("tool_call_id") in exempt:
            continue
        content = msg.get("content")
        if isinstance(content, str):
            if len(content) > TOOL_RESULT_STUB_CHARS:
                msg["content"] = (content[:TOOL_RESULT_STUB_CHARS]
                                  + "… [truncated — re-run the tool for "
                                    "the full output]")
        elif isinstance(content, list):
            for b in content:
                if (b.get("type") == "text"
                        and len(b.get("text", "")) > TOOL_RESULT_STUB_CHARS):
                    b["text"] = (b["text"][:TOOL_RESULT_STUB_CHARS]
                                 + "… [truncated]")


def cache_friendly(model: str | None = None) -> bool:
    """Only Anthropic-family models understand cache_control — for everyone
    else LiteLLM may forward it and the provider may reject the request."""
    m = (model or state.model).lower()
    return "claude" in m or "anthropic" in m


def apply_cache_control(messages: list) -> None:
    """Mark the system prompt + the newest message as cache breakpoints.

    Anthropic bills cached prefixes at ~10%, so with the breakpoint riding the
    tail each step pays full input price only for the newest messages and the
    recent screenshots — the whole prior history comes back as cache reads.
    """
    sysmsg = messages[0]
    if isinstance(sysmsg.get("content"), str):
        sysmsg["content"] = [{"type": "text", "text": sysmsg["content"],
                              "cache_control": {"type": "ephemeral"}}]
    for msg in messages[1:]:
        content = msg.get("content")
        if isinstance(content, list):
            for block in content:
                block.pop("cache_control", None)
    last = messages[-1]
    content = last.get("content")
    if isinstance(content, str):
        last["content"] = [{"type": "text", "text": content,
                            "cache_control": {"type": "ephemeral"}}]
    elif isinstance(content, list) and content:
        content[-1]["cache_control"] = {"type": "ephemeral"}


async def maybe_escalate(messages: list, conv_id: str | None,
                         reason: str) -> bool:
    """Switch the run to ESCALATION_MODEL once — the cheap model is stuck or
    erroring. Persists the choice on the conversation so a resume doesn't
    drop back to the model that was failing. Returns True if it switched."""
    if (not ESCALATION_MODEL or state.model == ESCALATION_MODEL
            or getattr(state, "escalated", False)):
        return False
    # Don't escalate into a deployment LiteLLM doesn't have — a typo'd or
    # unkeyed target would just fail the same way.
    try:
        deployed = {m.get("model_name") for m in await litellm_deployments()}
        if ESCALATION_MODEL not in deployed:
            print(f"[gut] escalation target {ESCALATION_MODEL} not deployed")
            return False
    except httpx.HTTPError:
        pass  # couldn't check — try anyway
    old = state.model
    state.model = ESCALATION_MODEL
    state.escalated = True
    swap_coord_prompt(messages, old)
    if conv_id:
        try:
            conv_set_model(conv_id, ESCALATION_MODEL)
        except ValueError:
            pass
    await broadcast({"type": "agent_msg",
                     "text": f"(escalating to {ESCALATION_MODEL} — {reason})"})
    await push_status()
    return True


def _note(content, text: str):
    """Append a text note to tool-result content (str or block list)."""
    if isinstance(content, str):
        return content + "\n" + text
    for block in content:
        if block.get("type") == "text":
            block["text"] += "\n" + text
            return content
    content.append({"type": "text", "text": text})
    return content


def tool_result_images_ok() -> bool:
    """Anthropic-family models consume image blocks inside tool results
    natively. The OpenAI-format providers LiteLLM fronts for everything
    else (OpenRouter, Ollama, vLLM, …) drop them — frames must ride a user
    message instead or the model never sees the screen."""
    return cache_friendly()


def _hoist_tool_images(messages: list, extra: dict | None = None) -> None:
    """Move image blocks out of the trailing tool results into a user
    message — the shape every vision provider accepts. `extra` is a frame
    that would otherwise have been attached to the last tool result."""
    imgs = []
    for msg in reversed(messages):
        if msg.get("role") != "tool":
            break
        content = msg.get("content")
        if not isinstance(content, list):
            continue
        moved = [b for b in content if b.get("type") == "image_url"]
        if not moved:
            continue
        imgs = moved + imgs
        keep = [b for b in content if b.get("type") != "image_url"]
        msg["content"] = keep or [
            {"type": "text", "text": "(image attached below)"}]
    if extra is not None:
        imgs.append(extra)
    if imgs:
        messages.append({"role": "user", "content": [
            {"type": "text", "text": "(latest screen)"}] + imgs})


def _unstick_note(reason: str) -> str:
    """Directive injected when the agent is looping without progress.

    Rather than asking the user, push the model to brainstorm alternatives
    and try a different approach. Only a human-only blocker (credentials,
    2FA, CAPTCHA, a decision) justifies ask_user.
    """
    return (f"STUCK: {reason}. Do not keep doing the same thing. Before your "
            "next tool call, brainstorm at least 5 fundamentally different "
            "ways to reach your goal — different tools, keyboard vs mouse, "
            "menus, run_command, browser_* tools, another app entirely, "
            "reading docs or --help output — then try the most promising one "
            "you haven't tried. Keep generating new approaches — never give "
            "up. Only call ask_user if you're blocked on something only a "
            "human can provide.")


STUCK_ASK_USER = ("I've brainstormed and tried several different approaches "
                  "and I'm still stuck. Any guidance?")

# Tools that only read — re-running them with identical arguments after
# every change is correct behavior, so they never count as stalling and are
# never disabled.
OBSERVE_TOOLS = {"screenshot", "browser_dom", "browser_text", "desktop_tree",
                 "list_windows", "collect_agent"}

# Pure network reads — no screen, no shared state, no ordering between
# calls. A turn whose calls are all in this set runs them concurrently:
# a batch of lookups used to cost a full model roundtrip each.
PARALLEL_TOOLS = {"web_search", "fetch_url"}

_BLOCKED_MSG = ("BLOCKED: this exact call was already repeated with no "
                "progress and is disabled for the rest of this run — use a "
                "different tool or different arguments.")

# Appended to the result when the run has spent RESEARCH_NUDGE_STREAK
# consecutive turns on one-at-a-time research calls.
RESEARCH_NUDGE_MSG = (
    "{n} research calls in a row, one per turn — the slow path. "
    "Parallelize: web_search's `queries` list runs several lookups in one "
    "call; several independent web_search/fetch_url calls in one reply run "
    "concurrently; spawn_agent puts a helper on each target while you keep "
    "working. And if you already have enough for a defensible answer, stop "
    "researching — flag what's unverified and deliver.")


class StallDetector:
    """Loop detection for one run, on tool signatures (name + sorted args).

    Two signals:
    * exact repeats — the same call 3 times in the last 10 earns a warning,
      5 times disables that exact call for the rest of the run;
    * cycles — CYCLE_WINDOW acting calls in a row that were each seen
      earlier in the run. A long loop (open A, scroll, open B, open C, wait,
      …) repeats every element once per lap and never trips the repeat
      counter, yet nothing new is being tried. Every call in the window is
      disabled at once, so the model has to change approach.

    `rescues` counts the unstick directives spent; the loop escalates or
    checks in with the user on it. It is not reset by a repainted screen —
    a cycle changes the screen every step and still goes nowhere.
    """

    def __init__(self) -> None:
        self.recent: deque = deque(maxlen=10)
        self.seen: set[str] = set()
        # AGENT_CYCLE_WINDOW <= 1 turns cycle detection off.
        self.novelty: deque | None = (deque(maxlen=CYCLE_WINDOW)
                                      if CYCLE_WINDOW > 1 else None)
        self.blocked: set[str] = set()
        self.rescues = 0

    @staticmethod
    def signature(name: str, args: dict) -> str:
        if name in ("ask_user", "task_complete"):
            return ""
        if name == "run_command":
            # The weak-model failure mode is respawning the same failing
            # script with cosmetic edits (whitespace, quoting) — strip
            # those so a semantic repeat still counts as a repeat. This is
            # a loop key, not the command that runs, so over-normalizing
            # only ever means a stuck loop is caught sooner.
            cmd = re.sub(r"[\s\"']+", "",
                         str(args.get("command", "")).lower())
            return name + " " + cmd
        return name + " " + json.dumps(args, sort_keys=True, default=str)

    def observe(self, name: str, sig: str) -> tuple[str, bool]:
        """Record an executed call. Returns (note, stuck): `note` is text
        to append to the tool result, `stuck` means a rescue was spent."""
        note, stuck = "", False
        reads = name in OBSERVE_TOOLS
        self.recent.append(sig)
        repeats = sum(1 for s in self.recent if s == sig)
        if repeats == 3:
            note = ("WARNING: you've done this exact action 3 times in your "
                    f"last {len(self.recent)} steps with no progress — "
                    "switch tactics now.")
        elif repeats >= 5 and not reads:
            note = _unstick_note(f"you've repeated this exact action "
                                 f"{repeats} times with no progress")
            self.blocked.add(sig)
            self.recent.clear()
            stuck = True
        if not reads and self.novelty is not None:
            self.novelty.append((sig, sig not in self.seen))
            self.seen.add(sig)
            if (len(self.novelty) == self.novelty.maxlen
                    and not any(new for _, new in self.novelty)):
                cycle = {s for s, _ in self.novelty}
                self.blocked |= cycle
                self.novelty.clear()
                self.recent.clear()
                cycle_note = _unstick_note(
                    f"your last {self.novelty.maxlen} actions were all "
                    "repeats of things you already did earlier in this run "
                    f"— you are cycling through {len(cycle)} dead ends. All "
                    "of them are now disabled")
                note = (note + "\n" + cycle_note) if note else cycle_note
                stuck = True
        if stuck:
            self.rescues += 1
        return note, stuck


async def spend_rescue(stall: StallDetector, messages: list, conv_id: str,
                       reason: str) -> str:
    """After an unstick directive: switch to ESCALATION_MODEL once the
    starting model has burnt ESCALATION_RESCUES of them, and check in with
    the user when even that (or no escalation model) leaves the run stuck.
    Returns a note to append to the tool result, "" when nothing happened."""
    if (stall.rescues >= ESCALATION_RESCUES
            and await maybe_escalate(messages, conv_id, reason)):
        stall.rescues = 0
        return ""
    if stall.rescues >= (3 if ESCALATION_MODEL else 2):
        stall.rescues = 0
        answer = await ask_user(STUCK_ASK_USER)
        return f"[user replied]: {answer}"
    return ""


async def llm_request_forcing(http: httpx.AsyncClient, messages: list,
                              force: str | None,
                              tools: list) -> httpx.Response:
    """llm_request with the reply pinned to tool `force` — how the checklist
    nudge and the final-step task_complete stop being suggestions the model
    answers in prose. A provider that rejects tool_choice (some OpenRouter
    routes, Ollama) gets the plain request instead and is remembered, so
    the 400 is paid once per model, not once per nudge."""
    if force and state.model not in state.no_tool_choice:
        try:
            return await llm_request(http, messages, tools=tools, tool_choice={
                "type": "function", "function": {"name": force}})
        except httpx.HTTPStatusError as e:
            if e.response.status_code != 400 or is_context_overflow(e):
                raise
            state.no_tool_choice.add(state.model)
            print(f"[gut] {state.model} rejected tool_choice "
                  f"({force}): {e.response.text[:200]}")
    if force and state.model not in state.no_tool_required:
        # Servers that only take the string forms (LM Studio, llama.cpp)
        # still honour "required": the reply is *some* tool call rather
        # than the one we wanted — but never prose, which is the failure
        # the pin exists to prevent.
        try:
            return await llm_request(http, messages, tools=tools,
                                     tool_choice="required")
        except httpx.HTTPStatusError as e:
            if e.response.status_code != 400 or is_context_overflow(e):
                raise
            state.no_tool_required.add(state.model)
    return await llm_request(http, messages, tools=tools)


def _inject_note(messages: list, text: str) -> None:
    """Put a system-side note where the next request will read it: onto the
    last tool result or user message, else as a user message of its own."""
    if messages and messages[-1].get("role") in ("tool", "user"):
        messages[-1]["content"] = _note(messages[-1]["content"], text)
    else:
        messages.append({"role": "user", "content": text})


def wrap_up_note(remaining: int) -> str:
    """Countdown to the step cap: a run that spends its last steps exploring
    ends with nothing delivered, when a flagged partial result was there
    for the taking."""
    if remaining <= 1:
        return ("FINAL STEP: the run ends after this call. Call task_complete "
                "now — say what you found, what is unverified or missing, "
                "and where any files are (send_file first if you have one).")
    return (f"NOTICE: only {remaining} steps remain before this run is cut "
            "off. Stop exploring and deliver: write up what you have — mark "
            "unverified values as estimates — send_file every deliverable, "
            "then call task_complete. A flagged partial result beats an "
            "empty run.")


async def llm_request(http: httpx.AsyncClient, messages: list,
                      model: str | None = None,
                      tools: list | None = None,
                      tool_choice: dict | str | None = None) -> httpx.Response:
    """One chat-completion call with retry on transient failures.

    A single 429/5xx mid-task used to kill the run and waste all prior spend.
    Retries with backoff; non-retryable 4xx propagates immediately.
    `tool_choice` pins the reply to one tool (OpenAI shape — LiteLLM
    translates it for Anthropic/Gemini).
    """
    payload = {"model": model or state.model, "messages": messages,
               "tools": tools if tools is not None else TOOLS}
    if isinstance(tool_choice, str) and payload["tools"]:
        payload["tool_choice"] = tool_choice  # "required" / "auto"
    elif tool_choice and any(t["function"]["name"]
                             == tool_choice["function"]["name"]
                             for t in payload["tools"]):
        payload["tool_choice"] = tool_choice
    if not payload["tools"]:
        del payload["tools"]  # no tools wanted (compaction) — an empty
                              # array is rejected by some providers
    delay, last_exc = 2.0, None
    for attempt in range(LLM_MAX_RETRIES):
        try:
            r = await http.post(
                f"{LITELLM_URL}/v1/chat/completions",
                headers={"Authorization": f"Bearer {state.litellm_key}"},
                json=payload)
            r.raise_for_status()
            return r
        except httpx.HTTPStatusError as e:
            last_exc = e
            code = e.response.status_code
            # 402 earns the backoff too: OpenRouter raises it while
            # in-flight requests' spend crosses the balance — it settles
            # once they finish, so an instant raise kills a run a short
            # wait would have saved. A real empty balance still fails,
            # just after the retries.
            if code not in (402, 408, 409, 429) and code < 500:
                raise
        except httpx.TransportError as e:
            last_exc = e
        if attempt + 1 < LLM_MAX_RETRIES:
            await broadcast({"type": "agent_msg",
                             "text": f"(model request failed, retrying in "
                                     f"{delay:.0f}s: {last_exc})"})
            await asyncio.sleep(delay)
            delay = min(delay * 2, 30)
    if last_exc is not None:
        raise last_exc
    raise RuntimeError("llm_request exhausted retries without a response")


# ── Context compaction ─────────────────────────────────────────────────────
# Long runs would otherwise grow the context array until the provider
# rejects the request. When the last response's prompt_tokens approach the
# model's window, one extra call summarizes the history into a handoff note
# and the loop continues on [system, handoff + checklist, recent tail].

COMPACT_PROMPT = """Summarize this conversation so far as a handoff note to yourself — the older messages are about to be dropped from context.

Cover, compactly:
1. Goal — the user's task, in one or two sentences.
2. Progress — what's done, what worked, what's left.
3. Key facts — file paths, URLs, names, values, decisions: anything the remaining steps still need.
4. Screen state — what's open / running right now.
5. Next — the immediate next action.

Be terse but complete: this note plus the last few messages are all you keep."""

# Summarizer input cap — the compact call itself must fit the window even
# when history is already past it, so the middle of a huge history is cut.
COMPACT_INPUT_MAX_CHARS = int(
    os.environ.get("AGENT_COMPACT_INPUT_CHARS", "400000"))


_CTX_CAP_RE = re.compile(r"(\d+)\s*([km]?)", re.I)


def model_context_cap(model: str) -> int:
    """Configured context ceiling for `model` from MODEL_CONTEXT_LIMITS —
    the model's own entry, else the "*" wildcard. 0 = uncapped."""
    caps = {}
    for part in str(MODEL_CONTEXT_LIMITS).split(","):
        name, _, val = part.partition("=")
        m = _CTX_CAP_RE.fullmatch(val.strip())
        if name.strip() and m:
            caps[name.strip()] = (int(m.group(1)) * 1000 ** {
                                  "": 0, "k": 1, "m": 2}[m.group(2).lower()])
    return caps.get(model) or caps.get("*") or 0


async def model_context_limit(http: httpx.AsyncClient) -> int:
    """Usable context window for the active model: max_input_tokens from
    LiteLLM's /model/info, cached per model name; AGENT_CONTEXT_LIMIT when
    it reports nothing. A MODEL_CONTEXT_LIMITS cap then shrinks the
    result — never grows it — so compaction fires early enough to keep a
    run under its set context."""
    cached = state.ctx_limit.get(state.model)
    if cached:
        return cached
    cap = model_context_cap(state.model)
    try:
        r = await http.get(
            f"{LITELLM_URL}/model/info",
            headers={"Authorization": f"Bearer {state.litellm_key}"})
        data = r.json()
        items = data.get("data", []) if isinstance(data, dict) else data
        for m in items:
            if m.get("model_name") == state.model:
                info = m.get("model_info") or {}
                lim = info.get("max_input_tokens") or info.get("max_tokens")
                if lim:
                    lim = min(int(lim), cap) if cap else int(lim)
                    state.ctx_limit[state.model] = lim
                    return lim
    except Exception:
        pass
    return min(COMPACT_CONTEXT_LIMIT, cap) if cap else COMPACT_CONTEXT_LIMIT


def _text_only(messages: list) -> list[dict]:
    """Copy of messages with image blocks replaced by markers — the
    summarizer call is text-only."""
    out = []
    for msg in messages:
        content = msg.get("content")
        if isinstance(content, list):
            content = [{"type": "text",
                        "text": "[screenshot]"
                        if b.get("type") == "image_url"
                        else str(b.get("text", ""))}
                       for b in content]
        out.append({**msg, "content": content})
    return out


def _msg_size(msg: dict) -> int:
    return len(json.dumps(msg, default=str))


def _image_token_est() -> int:
    """Pixel-based weight of one image block — providers bill vision frames
    by pixels (≈ w*h/750), not by the base64 length a chars/4 guess would
    use."""
    w, h = state.shot_size or (1024, 768)
    return max(85, int(w * h / 750))


def ctx_breakdown(messages: list) -> dict:
    """Raw context-size weights per category: {system, shots, tools, chat}.

    Text weighs as serialized chars/4 (same ruler as _msg_size); image
    blocks get the pixel estimate wherever they sit — user messages, or
    tool results on vision-native providers. Callers normalize the weights
    to the billed prompt_tokens, so only the proportions need to be right.
    """
    parts = {"system": 0, "shots": 0, "tools": 0, "chat": 0}
    for m in messages:
        content = m.get("content")
        if isinstance(content, list):
            imgs = [b for b in content if b.get("type") == "image_url"]
            if imgs:
                parts["shots"] += len(imgs) * _image_token_est()
                m = {**m, "content":
                     [b for b in content if b.get("type") != "image_url"]}
        cat = {"system": "system", "tool": "tools"}.get(m.get("role"),
                                                        "chat")
        parts[cat] += _msg_size(m) / 4
    return parts


async def compact_context(http: httpx.AsyncClient, conv_id: str,
                          messages: list, checklist: bool = True,
                          save: bool = True) -> bool:
    """Summarize older history into a handoff note and rebuild `messages`
    in place as [system, handoff + checklist, recent tail].

    `checklist` appends the run's live checklist to the handoff — workers
    pass False: the parent's plan is not theirs. `save` writes the
    rebuilt context to the conversation — workers pass False: their
    `messages` are their own, and conv_id is the parent's file. Returns
    False with `messages` untouched when summarization fails — the run
    just continues on the full history.
    """
    view = _text_only(messages[1:])
    total = sum(_msg_size(m) for m in view)
    if total > COMPACT_INPUT_MAX_CHARS and len(view) > 4:
        tail, acc = [], 0
        for m in reversed(view[1:]):
            acc += _msg_size(m)
            if acc > COMPACT_INPUT_MAX_CHARS:
                break
            tail.insert(0, m)
        view = [view[0],
                {"role": "user",
                 "content": "[... middle of the history omitted ...]"},
                *tail]
    try:
        r = await llm_request(
            http, view + [{"role": "user", "content": COMPACT_PROMPT}],
            model=COMPACT_MODEL or None, tools=[])
    except Exception as e:
        await broadcast({"type": "error",
                         "text": f"context compaction failed: {e}"})
        return False
    usd, tin, tout = track_cost(r)
    if (usd or tin or tout) and conv_id:
        conv_add_usage(conv_id, usd, tin, tout)
    await push_cost()
    summary = r.json()["choices"][0]["message"].get("content") or ""
    if not isinstance(summary, str):
        summary = " ".join(str(b.get("text", "")) for b in summary
                           if isinstance(b, dict))
    summary = summary.strip()
    if not summary:
        return False
    # Keep the freshest few messages — never messages[0] (the system
    # prompt), drop leading tool orphans, then let sanitize_context repair
    # a tool-call pair cut at the boundary.
    tail = messages[max(1, len(messages) - COMPACT_KEEP):]
    while tail and tail[0].get("role") == "tool":
        tail = tail[1:]
    dropped = len(messages) - 1 - len(tail)
    handoff = ("[The earlier conversation was compacted into this handoff "
               "summary; the messages after it are the recent tail.]\n\n"
               + summary)
    if checklist and state.todos:
        handoff += "\n\nCurrent checklist:\n" + render_todos(state.todos)
    messages[1:] = [{"role": "user", "content": handoff}, *tail]
    sanitize_context(messages)
    state.steps_since_compact = 0
    if save and conv_id:
        await asyncio.to_thread(conv_save_context, conv_id, messages)
    await broadcast({"type": "compact",
                     "text": f"context compacted — {dropped} older messages "
                             "summarized into a handoff note"})
    return True


# Provider error bodies for a blown context window vary — match the usual
# phrasings so an overflow 400 triggers compaction instead of killing the
# run (the ratio trigger normally fires first; this catches unknown-window
# models and single huge tool results).
_OVERFLOW_RE = re.compile(
    r"context|too many tokens|maximum.{0,20}length|token.{0,12}limit|"
    r"prompt is too long|reduce the length", re.I)


def is_context_overflow(e: httpx.HTTPError) -> bool:
    resp = getattr(e, "response", None)
    if resp is None or resp.status_code not in (400, 413):
        return False
    return bool(_OVERFLOW_RE.search(resp.text or ""))


# ── Wrap-up verification ───────────────────────────────────────────────────
# The summary is the only thing the user reads, and nothing checked it
# against what the tools actually did — a confident fabrication sailed
# straight through. Before task_complete is accepted, one cheap text-only
# call audits the claim against a ledger of the run's calls; unsupported
# claims bounce back as a critique the model must fix (or explicitly
# dispute) and resubmit.

VERIFY_PROMPT = """You audit an autonomous desktop agent's wrap-up message against the log of what its tools actually did and returned. Be strict about facts, generous about style.

The wrap-up PASSES when:
- every fact stated as certain (numbers, prices, dates, names, URLs, quotes, file contents) appears in the tool log — including inside subagent reports, whose [evidence — …] footer lists the pages actually fetched: a figure from a report whose footer says no page was fetched is unsupported, and
- every claimed deliverable matches a send_file call or a file the log shows being created — in the form asked for: a spreadsheet request delivered as CSV, a document that only exists as chat text, or a file the log shows failing to build counts as a mismatch; the "actual contents" digest is authoritative — a workbook with 0 filled rows or a document with 0 paragraphs is not the deliverable described, and
- claimed actions match calls that ran without an error result.

Claims the agent itself flags as unverified, estimated or approximate are fine — flagged doubt is honest. Fail ONLY for material claims stated as fact that the log doesn't support or directly contradicts — never over omissions, tone, or hedged language.

Ledger lines ending "… [truncated]" had their tails cut for length — a claim needing the cut part counts as unsupported, and prefer the digests over guessing what a cut page said.

A delivered file's digest (bytes, paragraphs, filled rows) is authoritative that it exists and is nonempty — fail missing, empty or wrong-type deliverables. The wrap-up describing its own document loosely ("two-paragraph summary" while the digest counts heading and caption lines too) is style, not an unsupported claim.

Reply with exactly one line:
PASS
or
FAIL
- <unsupported claim> — what the log shows instead (or "nothing in the log supports this")
- … one short bullet per problem."""


# Ledger slice multiplier per tool (× VERIFY_RESULT_CHARS): worker reports
# and the searches/fetches behind them carry the run's evidence — numbers
# and source URLs the wrap-up repeats — while a screenshot result or a
# screen action needs only its head.
_LEDGER_CAP_MULT = {"collect_agent": 10, "fetch_url": 4, "web_search": 4,
                    "run_command": 2, "send_file": 4}


def run_ledger(messages: list) -> list[str]:
    """One line per completed tool call: name(args…) → result head. The
    evidence the checker audits the wrap-up against — built mechanically,
    so unlike a self-written summary it can't confabulate."""
    pending: dict = {}
    entries: list[str] = []
    for msg in messages:
        role = msg.get("role")
        if role == "assistant":
            for tc in msg.get("tool_calls") or []:
                fn = tc.get("function") or {}
                pending[tc.get("id")] = (
                    fn.get("name", "?"),
                    str(fn.get("arguments") or "")[:160])
        elif role == "tool":
            name, argstr = pending.pop(msg.get("tool_call_id"), ("?", ""))
            res = msg.get("content")
            if isinstance(res, list):
                res = " ".join(str(b.get("text", "")) for b in res
                               if b.get("type") == "text")
            res = " ".join(str(res or "").split())
            # Evidence-heavy results get a bigger slice — the checker fails
            # claims whose only support fell past the cap, and research
            # figures live inside worker reports and fetched pages. Mark
            # cut tails so truncation reads as truncation, not absence.
            cap = VERIFY_RESULT_CHARS * _LEDGER_CAP_MULT.get(name, 1)
            entries.append(f"{name}({argstr}) → {res[:cap]}"
                           + (" … [truncated]" if len(res) > cap else ""))
        elif role == "user" and isinstance(msg.get("content"), str) \
                and msg["content"].startswith("[subagent '"):
            # Reports auto-delivered between steps ride user turns, not
            # tool results — without this line the checker never sees
            # the evidence behind a delegated finding.
            res = " ".join(msg["content"].split())
            cap = VERIFY_RESULT_CHARS * _LEDGER_CAP_MULT["collect_agent"]
            entries.append(f"subagent report → {res[:cap]}"
                           + (" … [truncated]" if len(res) > cap else ""))
    return entries


def _ledger_window(entries: list[str]) -> str:
    """Fit the ledger under VERIFY_INPUT_CHARS — keep the run's start and
    the recent tail, drop the middle."""
    if sum(len(e) + 1 for e in entries) <= VERIFY_INPUT_CHARS:
        return "\n".join(entries)
    half = VERIFY_INPUT_CHARS // 2
    head, tail, size = [], [], 0
    for e in entries:
        if size + len(e) + 1 > half:
            break
        head.append(e)
        size += len(e) + 1
    rest = entries[len(head):]
    size = 0
    for e in reversed(rest):
        if size + len(e) + 1 > half:
            break
        tail.append(e)
        size += len(e) + 1
    tail.reverse()
    omitted = len(rest) - len(tail)
    mid = [f"[… {omitted} calls omitted …]"] if omitted > 0 else []
    return "\n".join(head + mid + tail)


async def verify_wrap_up(http: httpx.AsyncClient, conv_id: str,
                         summary: str, messages: list) -> str | None:
    """Audit the wrap-up against the run ledger. Returns the critique the
    model sees as task_complete's tool result, or None to let the
    completion through — also the fail-open path: a checker that can't
    reach its model must never wedge a run."""
    if not AGENT_VERIFY:
        return None
    entries = run_ledger(messages)
    if len(entries) < VERIFY_MIN_CALLS:
        return None
    if state.verify_rejects >= VERIFY_MAX_REJECTS:
        return None  # cap spent — the stored caveat rides the done text
    # Real user turns are block lists; daemon nudges are plain strings.
    user_texts = []
    for m in messages:
        c = m.get("content")
        if m.get("role") == "user" and isinstance(c, list):
            t = " ".join(str(b.get("text", "")) for b in c
                         if b.get("type") == "text").strip()
            if t:
                user_texts.append(t[:1500])
    users = "\n---\n".join(user_texts)
    if len(users) > 6000:
        users = (users[:2500] + "\n[… earlier messages trimmed …]\n"
                 + users[-3500:])
    unsent = [str(p) for p in unsent_outputs()]

    def _digests(paths) -> str:
        return "\n".join(f"- {p} — {file_digest(Path(p))}"
                         for p in list(paths)[:15]) or "(none)"
    sent_d, unsent_d = await asyncio.to_thread(
        lambda: (_digests(sorted(state.sent_files)), _digests(unsent)))
    prompt = (
        "USER MESSAGES (oldest→newest):\n" + (users or "(none)") + "\n\n"
        "CHECKLIST STATE:\n" + (render_todos(state.todos) or "(none)")
        + "\n\n"
        "send_file DELIVERED (with actual contents):\n" + sent_d + "\n"
        "files changed this run, NOT delivered:\n" + unsent_d + "\n\n"
        "WRAP-UP UNDER REVIEW:\n" + summary + "\n\n"
        "TOOL LEDGER (oldest→newest):\n" + _ledger_window(entries))
    try:
        r = await llm_request(
            http,
            [{"role": "system", "content": VERIFY_PROMPT},
             {"role": "user", "content": prompt}],
            model=VERIFY_MODEL or COMPACT_MODEL or None, tools=[])
    except Exception as e:
        print(f"[gut] verifier call failed, accepting wrap-up: {e}")
        return None
    usd, tin, tout = track_cost(r)
    if conv_id and (usd or tin or tout):
        conv_add_usage(conv_id, usd, tin, tout)
    await push_cost()
    verdict = message_text(r.json()["choices"][0]["message"]).lstrip()
    if verdict[:4].upper() != "FAIL":
        state.verify_caveat = None
        return None
    state.verify_rejects += 1
    critique = re.sub(r"^FAIL\w*\s*[:\-]?\s*", "", verdict).strip() or \
        "claims not supported by the tool log"
    state.verify_caveat = critique[:1500]
    return ("VERIFICATION FAILED — the wrap-up states things the tool log "
            "doesn't back up:\n" + critique[:4000] +
            "\nFix the flagged claims — re-check them with tools or mark "
            "them unverified — then call task_complete again. If a flagged "
            "claim IS in the log, say where in the new summary.")


# ── Subagents ─────────────────────────────────────────────────────────────
# spawn_agent runs a helper loop in the background — headless tools only, so
# parallel helpers can't fight over the screen, the browser or the user.
# Reports reach the parent two ways: auto-injected as a user message at the
# top of the next agent_loop step (subagent_inbox), or pulled on demand via
# collect_agent. `delivered` dedupes the two paths; `conv` scopes injection
# to the conversation that spawned the helper.

def _subagent_name(name: str) -> str:
    return re.sub(r"[^a-z0-9-]+", "-", name.lower()).strip("-")[:24]


def agents_status() -> str:
    if not state.subagents:
        return "no helpers yet"
    return "helpers: " + ", ".join(
        f"{n} ({e['status']})" for n, e in state.subagents.items())


def subagent_report(name: str, e: dict) -> str:
    return (f"[subagent '{name}' {e['status']} — {e.get('steps', 0)} steps, "
            f"${e.get('usd', 0):.4f}]\n{e.get('result') or '(no report)'}")


# Claims in a worker's own report text: a fetched URL or a figure with a
# unit/currency. Bare prose ("tried 3 searches", "220 people") doesn't
# match — the gate targets reports that assert findings.
_REPORT_CLAIM_RE = re.compile(
    r"https?://|\d[\d.,]*\s*(?:kr\.?|dkk|usd|eur|gbp|sek|nok|\$|€|£|"
    r"kg|g\b|stk|cl\b|ml\b|l\b|liter|%)", re.I)


def _findings_empty(wd: Path) -> bool:
    """True when a worker's dir holds no substantive findings file —
    missing, or only a trivially empty `[]`/`{}`/stub."""
    try:
        files = [p for p in wd.iterdir() if p.is_file()] \
            if wd.is_dir() else []
    except OSError:
        files = []
    for p in files:
        try:
            body = p.read_text(errors="replace").strip()
        except OSError:
            continue
        if len(body) > 20 and body not in ("[]", "{}", "null"):
            return False
    return True


def worker_footer(messages: list, workdir: Path, report: str = "") -> str:
    """Mechanical evidence trailer for a worker's report: what its tools
    actually did (searches run, pages fetched and which, commands run) and
    what it left on disk. Built from the worker's own message log, so it
    cannot confabulate — a report full of prices whose footer shows no
    page fetched is exactly the report the orchestrator must not trust."""
    counts: dict[str, int] = {}
    fetched: list[str] = []
    pending: dict[str, tuple[str, dict]] = {}
    for m in messages:
        if m.get("role") == "assistant":
            for tc in m.get("tool_calls") or []:
                fn = tc.get("function") or {}
                try:
                    args = json.loads(fn.get("arguments") or "{}")
                except json.JSONDecodeError:
                    args = {}
                pending[tc.get("id")] = (fn.get("name", "?"), args)
        elif m.get("role") == "tool":
            name, args = pending.pop(m.get("tool_call_id"), ("?", {}))
            res = m.get("content")
            if isinstance(res, list):
                res = " ".join(str(b.get("text", "")) for b in res
                               if b.get("type") == "text")
            res = str(res or "")
            if name == "task_complete":
                continue
            counts[name] = counts.get(name, 0) + 1
            if name == "fetch_url" and not _TOOL_ERR_RE.search(res):
                fetched.append(str(args.get("url", ""))[:120])
    parts = [f"{n} {t}" for t, n in sorted(counts.items())]
    if fetched:
        shown = ", ".join(dict.fromkeys(fetched))
        parts.append(f"pages fetched: {shown[:600]}")
    files = []
    try:
        for p in sorted(workdir.iterdir()) if workdir.is_dir() else []:
            if p.is_file():
                files.append(f"{p.relative_to(SCRATCH_DIR)} — "
                             f"{file_digest(p)}")
    except OSError:
        pass
    parts.append("files: " + ("; ".join(files) if files else "none"))
    foot = "[evidence — " + "; ".join(parts) + "]"
    if not fetched and (counts.get("web_search") or counts.get("run_command")):
        foot += ("\n⚠ no page was fetched — any figure above rests on search "
                 "snippets or memory and is UNVERIFIED")
    elif not counts:
        foot += "\n⚠ no tools were used — this report is UNVERIFIED"
    if _REPORT_CLAIM_RE.search(report) and _findings_empty(workdir):
        foot += ("\n⚠ the report cites figures but no findings file was "
                 "written — treat its values as UNVERIFIED")
    return foot


def worker_dir(name: str) -> Path:
    """Where a worker writes its findings — one dir per worker under
    scratch, so the orchestrator's assembly script has a known place to
    read from and the footer has a known place to list."""
    return SCRATCH_DIR / name


def spawn_agent(task: str, name: str = "", model: str = "") -> str:
    task = task.strip()
    if not task:
        return "empty task — nothing spawned"
    if state.spawned >= SUBAGENT_MAX_TOTAL:
        return (f"spawn cap reached ({SUBAGENT_MAX_TOTAL} helpers this run) "
                "— work with the reports you have; flag what's missing as "
                "unverified.")
    # GC: drop finished entries whose report already reached the parent.
    for n in [n for n, e in state.subagents.items()
              if e["status"] != "running" and e.get("delivered")]:
        del state.subagents[n]
    running = [n for n, e in state.subagents.items()
               if e["status"] == "running"]
    if len(running) >= SUBAGENT_MAX_CONCURRENT:
        return (f"at the helper cap ({SUBAGENT_MAX_CONCURRENT}) — running: "
                + ", ".join(running)
                + ". collect_agent or keep working until one reports.")
    name = _subagent_name(name)
    if name and state.subagents.get(name, {}).get("status") == "running":
        return f"'{name}' is already running — pick another name"
    if not name:
        n = 1
        while f"agent-{n}" in state.subagents:
            n += 1
        name = f"agent-{n}"
    model = model.strip() or SUBAGENT_MODEL or state.model
    # The goal rides inside the daemon's contract (SUBAGENT_TASK) — what
    # the worker must deliver and where is never left to the caller.
    wd = worker_dir(name)
    prompt = SUBAGENT_TASK.format(goal=task, workdir=wd)
    state.spawned += 1
    entry = {"name": name, "desc": task[:200], "conv": state.conversation_id,
             "status": "running", "result": None, "model": model,
             "usd": 0.0, "steps": 0, "delivered": False, "workdir": str(wd)}
    entry["task"] = asyncio.create_task(
        subagent_loop(name, prompt, model, entry))
    state.subagents[name] = entry
    return (f"spawned '{name}' (model {model}) — it works in the background; "
            "its report arrives as a message. collect_agent(name) blocks "
            f"for it. Its findings land under {wd}/ "
            f"({state.spawned}/{SUBAGENT_MAX_TOTAL} spawns used).")


async def collect_agent(name: str = "") -> str:
    """Wait for a helper's report. No name = the next one to finish."""
    name = _subagent_name(name)
    if name:
        e = state.subagents.get(name)
        if e is None:
            return f"no helper '{name}' — {agents_status()}"
        while e["status"] == "running" and not state.stop:
            await asyncio.sleep(0.3)
        _deliver_report(e)
        return subagent_report(name, e)
    while not state.stop:
        done = [(n, e) for n, e in state.subagents.items()
                if e["status"] != "running" and not e.get("delivered")]
        if done:
            n, e = done[0]
            _deliver_report(e)
            return subagent_report(n, e)
        if not any(e["status"] == "running"
                   for e in state.subagents.values()):
            return f"nothing to collect — {agents_status()}"
        await asyncio.sleep(0.4)
    return "(stopped)"


async def subagent_loop(name: str, task_text: str, model: str,
                        entry: dict) -> None:
    """Headless helper loop: text tools only, own step cap and model. The
    final report queues for injection into the parent conversation."""
    conv_id = entry["conv"]
    usd0 = state.session_usd
    messages = [
        {"role": "system", "content": SUBAGENT_PROMPT.format(
            home=HOME_DIR, name=name, workdir=worker_dir(name))},
        {"role": "user", "content": task_text}]
    status, result = "done", "(ended without a report)"
    await broadcast_conv({"type": "subagent", "name": name,
                          "state": "running", "model": model,
                          "task": task_text[:300]}, conv_id)
    searches, fetches, overflowed = 0, 0, False
    try:
        async with httpx.AsyncClient(
                timeout=httpx.Timeout(300, connect=30)) as http:
            for step in range(SUBAGENT_MAX_STEPS):
                if state.stop:
                    status, result = "stopped", "(stopped by user)"
                    break
                entry["steps"] = step + 1
                # Worker discipline, daemon-side: a worker that only ever
                # searches has read nothing — after a few searches with no
                # page fetched, say so; two steps before the cap, demand
                # the report — a capped worker returns nothing at all.
                remaining = SUBAGENT_MAX_STEPS - step
                if remaining <= 2:
                    _inject_note(messages, (
                        "FINAL STEP: call task_complete NOW with what you "
                        "have — values you read off pages, the URLs, and "
                        "an UNVERIFIED list for the rest."))
                elif searches >= 3 and not fetches and searches % 3 == 0:
                    _inject_note(messages, (
                        f"you have run {searches} searches and fetched no "
                        "page — a snippet is not a source. fetch_url the "
                        "best 2 results now and read the figure off the "
                        "page, or task_complete with UNVERIFIED and what "
                        "you tried."))
                # Workers hold fetched pages in context — prune everything
                # but the newest few results every step (their record is
                # findings.json, not the transcript), and on an overflow
                # 400 prune hard and compact rather than dying.
                prune_tool_results(messages, keep=4)
                if cache_friendly(model):
                    apply_cache_control(messages)
                try:
                    r = await llm_request(http, messages, model=model,
                                          tools=SUBAGENT_TOOLS)
                except httpx.HTTPError as e:
                    body = getattr(e.response, "text", "") \
                        if hasattr(e, "response") else ""
                    if is_context_overflow(e) and not overflowed:
                        overflowed = True
                        prune_tool_results(messages, keep=2)
                        if await compact_context(http, conv_id, messages,
                                                 checklist=False,
                                                 save=False):
                            continue
                    status, result = "error", \
                        f"model request failed: {e} {body[:300]}"
                    break
                usd, tin, tout = track_cost(r)
                if conv_id and (usd or tin or tout):
                    conv_add_usage(conv_id, usd, tin, tout)
                await push_cost()

                msg = r.json()["choices"][0]["message"]
                messages.append(assistant_to_dict(msg))
                think = reasoning_text(msg)
                if think:
                    await broadcast_conv({"type": "thinking", "text": think,
                                          "agent": name}, conv_id)
                reply = message_text(msg)
                if reply:
                    await broadcast_conv({"type": "thought", "text": reply,
                                          "agent": name}, conv_id)
                tool_calls = msg.get("tool_calls") or []
                if not tool_calls:
                    # A plain reply IS the report — no idle nudges here.
                    result = reply or result
                    break
                finished = False
                parsed: list[tuple[dict, str, dict]] = []
                for tc in tool_calls:
                    fn = tc.get("function") or {}
                    try:
                        targs = json.loads(fn.get("arguments") or "{}")
                    except json.JSONDecodeError:
                        targs = {}
                    parsed.append((tc, fn.get("name", ""), targs))
                # Same rule as the main loop: a turn of nothing-but
                # PARALLEL_TOOLS calls runs concurrently.
                pres: list | None = None
                if len(parsed) > 1 and all(n in PARALLEL_TOOLS
                                           for _, n, _ in parsed):
                    for _, tname, targs in parsed:
                        await broadcast_conv({"type": "action", "tool": tname,
                                              "args": targs, "agent": name},
                                             conv_id)
                    pres = await asyncio.gather(*(
                        execute_tool(n, a, agent=name)
                        for _, n, a in parsed))
                for i, (tc, tname, targs) in enumerate(parsed):
                    if state.stop:
                        break
                    if pres is None:
                        await broadcast_conv({"type": "action", "tool": tname,
                                              "args": targs, "agent": name},
                                             conv_id)
                        res, fin = await execute_tool(tname, targs,
                                                      agent=name)
                    else:
                        res, fin = pres[i]
                    messages.append({"role": "tool",
                                     "tool_call_id": tc.get("id"),
                                     "content": res})
                    prev = res if isinstance(res, str) else next(
                        (b.get("text", "") for b in res
                         if b.get("type") == "text"), "")
                    if tname == "web_search":
                        searches += 1
                    elif tname == "fetch_url" and \
                            not _TOOL_ERR_RE.search(prev):
                        fetches += 1
                    ev = {"type": "action_result", "tool": tname,
                          "result": prev.strip()[:500], "agent": name}
                    urls = _result_urls(tname, prev)
                    if urls:
                        ev["urls"] = urls
                    await broadcast_conv(ev, conv_id)
                    if fin:  # task_complete — the summary is the report
                        result = res if isinstance(res, str) else prev
                        finished = True
                        break
                if finished:
                    break
            else:
                result = f"{result} (hit the {SUBAGENT_MAX_STEPS}-step cap)"
    except asyncio.CancelledError:
        status, result = "stopped", "(stopped)"
    except Exception as e:
        status, result = "error", f"helper error: {e}"
    entry["usd"] = round(state.session_usd - usd0, 6)
    if status == "done":
        footer = worker_footer(messages, worker_dir(name), str(result))
        entry["footer"] = footer
        result = f"{result}\n{footer}"
    entry.update(status=status, result=result)
    # Even a stopped helper's report is queued — if the run was stopped the
    # drained-on-resume message tells the parent the helper died with it.
    state.subagent_inbox.append(entry)
    await broadcast_conv({"type": "subagent", "name": name, "state": status,
                          "result": str(result)[:600], "usd": entry["usd"],
                          "steps": entry["steps"], "model": model}, conv_id)
    if conv_id and not (state.running and state.conversation_id == conv_id):
        # The helper's conversation isn't running — the inbox copy won't
        # drain into a live context and dies with the daemon. Append the
        # report to the stored context so the next run sees it even across
        # a restart (a rare duplicate beats a lost report).
        try:
            msgs = conv_load_context(conv_id)
            if msgs is not None:
                msgs.append({"role": "user",
                             "content": subagent_report(name, entry)})
                conv_save_context(conv_id, msgs)
        except Exception as e:
            print(f"[gut] helper report persist failed: {e}")


def drain_subagent_inbox(conv_id: str, messages: list) -> None:
    """Deliver finished helper reports into the run's context as user
    messages — scoped to the conversation that spawned them, so a helper of
    another conversation stays queued until that conv next runs."""
    kept = []
    while state.subagent_inbox:
        e = state.subagent_inbox.popleft()
        if e.get("delivered"):
            continue
        if e.get("conv") != conv_id:
            kept.append(e)
            continue
        _deliver_report(e)
        messages.append({"role": "user",
                         "content": subagent_report(e["name"], e)})
    state.subagent_inbox.extend(kept)


def _deliver_report(e: dict) -> None:
    """Mark a helper's report as delivered to the orchestrator. A report
    that finished with findings is the research evidence the checklist
    gate looks for; one that errored, was stopped or came back empty is
    not — a step can't be done on a worker that did nothing."""
    if e.get("delivered"):
        return
    e["delivered"] = True
    if e.get("status") == "done" and e.get("result") \
            and not str(e["result"]).startswith("("):
        note_evidence("research")


async def drain_user_msgs(conv_id: str, messages: list, mode: str) -> bool:
    """Deliver queued user messages into the run's context as user turns.

    "steer" entries land mid-run at the next step boundary; "queue"
    entries are delivered when the run would otherwise finish — a
    follow-up the user typed while the agent was still working, picked up
    before the end-of-run cleanup tears down what the run left behind.
    """
    delivered, kept = [], []
    while state.user_msgs:
        e = state.user_msgs.popleft()
        if e.get("conv") == conv_id and e.get("mode") == mode:
            delivered.append(e)
        else:
            kept.append(e)
    state.user_msgs.extend(kept)
    if not delivered:
        return False
    for e in delivered:
        text = e["text"]
        if mode == "steer":
            text = ("[the user sent you this message while you were "
                    f"working]\n{text}")
        content = [{"type": "text", "text": text}]
        for f in e.get("files") or []:
            if f["mime"].startswith("image/"):
                content.append({"type": "image_url", "image_url": {
                    "url": f"data:{f['mime']};base64,{f['b64']}"}})
        if mode == "queue":
            # A follow-up turn starts from the screen the last task left.
            shot = screenshot_block(force=True)
            if shot is not None:
                content.append(shot)
        messages.append({"role": "user", "content": content})
    items = [{"conv": conv_id, "seq": e["seq"]} for e in delivered
             if e.get("seq") is not None]
    if items:
        await broadcast({"type": "dequeue", "items": items})
    return True


# ── Run cleanup ──────────────────────────────────────────────────────────────
# End-of-run teardown — daemon-driven, never model-invoked. First a
# deterministic sweep of whatever the run created: browser tabs (CDP target
# diff), windows (wmctrl id diff) and processes (pid+starttime diff), plus a
# wipe of ~/scratch. Then a bounded janitor LLM pass with a fresh context
# cleans up residue the sweep can't enumerate (modal dialogs, wedged apps).
# Cookies and logins persist — tabs are closed, the profile is never touched.
# The baseline is persisted at run start so a daemon restart mid-run still
# lets the next boot sweep that run's leftovers.

def _snapshot_procs() -> dict[str, str]:
    """pid -> starttime for every process. Starttime (jiffies since boot)
    disambiguates PID reuse: same pid + different starttime = new process."""
    out = {}
    try:
        entries = list(Path("/proc").iterdir())
    except OSError:
        return out  # no procfs (bare-metal dev) — nothing to diff
    for p in entries:
        if not p.name.isdigit():
            continue
        try:
            st = (p / "stat").read_text().rsplit(")", 1)[1].split()
            # Zombies are already dead — a pid-1 uvicorn never reaps
            # orphans, so they linger and would be swept every time.
            if st[0] == "Z":
                continue
            out[p.name] = st[19]
        except (OSError, IndexError):
            continue
    return out


def _snapshot_windows() -> list[str]:
    try:
        p = subprocess.run(["wmctrl", "-l"], capture_output=True, text=True,
                           timeout=10)
    except (OSError, subprocess.SubprocessError):
        return []
    return [ln.split(None, 1)[0] for ln in p.stdout.splitlines()
            if ln.strip()]


def _snapshot_tabs() -> list[str]:
    try:
        r = httpx.get(f"{CDP_HTTP}/json", timeout=5)
        return [t["id"] for t in r.json() if t.get("type") == "page"]
    except Exception:
        return []


def _boot_key() -> str:
    """Starttime of pid 1 — changes when the container/host reboots, so a
    persisted baseline can tell "the daemon restarted" (same boot: sweep the
    dead run's leftovers) from "the whole desktop restarted" (every process
    is new — the baseline is meaningless and sweeping it would kill Xvfb,
    websockify, the panels…)."""
    try:
        return Path("/proc/1/stat").read_text().rsplit(")", 1)[1].split()[19]
    except (OSError, IndexError):
        return ""


def capture_baseline() -> dict:
    try:
        SCRATCH_DIR.mkdir(parents=True, exist_ok=True)
    except OSError:
        pass
    return {"started": time.time(), "boot": _boot_key(),
            "pids": _snapshot_procs(),
            "windows": _snapshot_windows(),
            "tabs": _snapshot_tabs()}


def _persist_baseline(b: dict | None) -> None:
    try:
        if b is None:
            RUN_STATE_FILE.unlink(missing_ok=True)
        else:
            RUN_STATE_FILE.parent.mkdir(parents=True, exist_ok=True)
            RUN_STATE_FILE.write_text(json.dumps(b))
    except OSError:
        pass


# Desktop/session plumbing the sweep must never kill, even when it respawned
# mid-run — a new pid+starttime reads as run-created, but SIGTERM to
# xfce4-panel takes the top bar and dock with it, and losing the session,
# VNC or model stack is worse. comm strings are capped at 15 chars.
PROTECTED_PROCS = {
    "Xvfb", "Xorg", "x11vnc", "xfce4-session", "xfce4-panel", "xfdesktop",
    "xfwm4", "xfsettingsd", "xfconfd", "Thunar", "wrapper-2.0",
    "dbus-daemon", "dbus-launch", "dbus-run-sessio", "gpg-agent",
    "ssh-agent", "startxfce4", "gut-session.sh", "xfce4-power-man",
    "light-locker", "litellm", "openserp",
}
PROTECTED_PREFIXES = ("gvfs", "at-spi", "xfsm-", "postgres")
PROTECTED_CMDLINES = ("websockify", "tcpmux", "litellm", "openserp")


def _proc_protected(pid: int) -> bool:
    try:
        comm = Path(f"/proc/{pid}/comm").read_text().strip()
    except OSError:
        return False
    if comm in PROTECTED_PROCS or comm.startswith(PROTECTED_PREFIXES):
        return True
    try:
        cl = (Path(f"/proc/{pid}/cmdline").read_bytes()
              .replace(b"\0", b" ").decode("utf-8", "replace"))
    except OSError:
        return False
    return any(s in cl for s in PROTECTED_CMDLINES)


def _helper_claimed(pid: int) -> bool:
    """True while pid is tagged GUT_HELPER=<name> (set on helper run_command
    shells, inherited by their children) and that helper is still running —
    the sweep must not kill a live helper's processes out from under it.
    Once the helper reports, its leftovers become fair game again."""
    try:
        env = Path(f"/proc/{pid}/environ").read_bytes()
    except OSError:
        return False
    for kv in env.split(b"\0"):
        k, _, v = kv.partition(b"=")
        if k == b"GUT_HELPER":
            e = state.subagents.get(v.decode("utf-8", "replace"))
            return bool(e and e.get("status") == "running")
    return False


def _window_owners() -> dict[str, int]:
    """window-id -> owner pid via `wmctrl -lp` (0 when unresolvable)."""
    try:
        p = subprocess.run(["wmctrl", "-l", "-p"], capture_output=True,
                           text=True, timeout=10)
    except (OSError, subprocess.SubprocessError):
        return {}
    out = {}
    for ln in p.stdout.splitlines():
        parts = ln.split(None, 4)
        if len(parts) >= 3:
            try:
                out[parts[0]] = int(parts[2])
            except ValueError:
                out[parts[0]] = 0
    return out


def sweep_desktop(baseline: dict) -> dict:
    """Close tabs/windows and kill processes created since baseline — never
    anything that predates the run or is desktop infrastructure."""
    stats = {"tabs": 0, "windows": 0, "procs": 0, "scratch": 0}

    # Closing a tab never logs anyone out — cookies live in the profile.
    old_tabs = set(baseline.get("tabs") or [])
    try:
        cur = {t["id"] for t in
               httpx.get(f"{CDP_HTTP}/json", timeout=5).json()
               if t.get("type") == "page"}
        for tid in cur - old_tabs:
            try:
                httpx.get(f"{CDP_HTTP}/json/close/{tid}", timeout=5)
                stats["tabs"] += 1
            except httpx.HTTPError:
                pass
    except Exception:
        pass

    # Graceful close so apps shut down clean (no soffice recovery prompt).
    # Windows owned by protected processes (a panel that respawned mid-run,
    # the desktop itself) read as "new" but are not residue — skip them.
    owners = _window_owners()
    wids = set(owners) or set(_snapshot_windows())
    for wid in wids - set(baseline.get("windows") or []):
        if _proc_protected(owners.get(wid) or 0):
            continue
        subprocess.run(["wmctrl", "-i", "-c", wid],
                       check=False, capture_output=True)
        stats["windows"] += 1
    if stats["windows"]:
        time.sleep(1.5)  # let apps quit before the process sweep

    # Whatever's still alive that wasn't at baseline — backgrounded shells,
    # GUI apps, timed-out commands. TERM, then KILL the stubborn.
    old_pids = baseline.get("pids") or {}
    me = str(os.getpid())

    def new_procs() -> list[int]:
        return [int(p) for p, s in _snapshot_procs().items()
                if old_pids.get(p) != s and p != me
                and not _proc_protected(int(p))
                and not _helper_claimed(int(p))]

    for pid in new_procs():
        try:
            os.kill(pid, signal.SIGTERM)
            stats["procs"] += 1
        except OSError:
            pass
    if stats["procs"]:
        time.sleep(1.0)
        for pid in new_procs():
            try:
                os.kill(pid, signal.SIGKILL)
            except OSError:
                pass

    if SCRATCH_DIR.is_dir():
        for child in SCRATCH_DIR.iterdir():
            try:
                if child.is_dir() and not child.is_symlink():
                    shutil.rmtree(child)
                else:
                    child.unlink()
                stats["scratch"] += 1
            except OSError:
                pass
    return stats


def _proc_alive(comm: str) -> bool:
    """pgrep -x, minus zombies: orphaned procs reparent to pid 1 — the
    daemon itself — which never reaps, so the dead still list."""
    try:
        pids = subprocess.run(["pgrep", "-x", comm], capture_output=True,
                              text=True, timeout=5).stdout.split()
    except (OSError, subprocess.SubprocessError):
        return False
    for pid in pids:
        try:
            st = Path(f"/proc/{pid}/stat").read_text() \
                .rsplit(")", 1)[1].split()[0]
            if st != "Z":
                return True
        except (OSError, IndexError):
            continue
    return False


def _session_env() -> dict:
    """The daemon's env lacks the session bus — dbus-run-session creates a
    private one for gut-session.sh — so borrow it from the session manager
    for relaunched components to land on the right bus and display."""
    env = dict(os.environ)
    try:
        pids = subprocess.run(["pgrep", "-x", "xfce4-session"],
                              capture_output=True, text=True,
                              timeout=5).stdout.split()
    except (OSError, subprocess.SubprocessError):
        return env
    for pid in pids:
        try:
            raw = Path(f"/proc/{pid}/environ").read_bytes()
        except OSError:
            continue
        for kv in raw.split(b"\0"):
            k, _, v = kv.partition(b"=")
            if k in (b"DBUS_SESSION_BUS_ADDRESS", b"DISPLAY", b"XAUTHORITY"):
                env[k.decode()] = v.decode()
        if env.get("DBUS_SESSION_BUS_ADDRESS"):
            break
    return env


def heal_desktop() -> list[str]:
    """Relaunch desktop components missing after cleanup — a panel or WM
    that died mid-run (or was swept before protection existed) leaves a
    bare screen. No-op when the session itself is gone (real logout), so
    this never fights a shutdown."""
    if not _proc_alive("xfce4-session"):
        return []
    env = _session_env()
    revived = []
    for proc, cmd in (("xfwm4", ["xfwm4"]),
                      ("xfsettingsd", ["xfsettingsd"]),
                      ("xfdesktop", ["xfdesktop"]),
                      ("xfce4-panel", ["xfce4-panel"])):
        if _proc_alive(proc):
            continue
        try:
            subprocess.Popen(["setsid", *cmd], env=env,
                             stdin=subprocess.DEVNULL,
                             stdout=subprocess.DEVNULL,
                             stderr=subprocess.DEVNULL)
            revived.append(proc)
        except (OSError, subprocess.SubprocessError):
            continue
    if revived:
        time.sleep(1.5)  # let them register before the next baseline
    return revived


# The janitor gets desktop tools only — no run_command (too broad for an
# unsupervised pass), no browser_* (tabs were just closed), no ask_user
# (cleanup must never block) and no send_* (the wrap-up already went out).
JANITOR_NAME = "janitor"
JANITOR_TOOL_NAMES = {
    "screenshot", "wait", "list_windows", "focus_window", "click",
    "mouse_move", "scroll", "type_text", "key", "desktop_tree",
    "desktop_act", "desktop_click", "task_complete",
}
JANITOR_TOOLS = [t for t in TOOLS
                 if t["function"]["name"] in JANITOR_TOOL_NAMES]

JANITOR_PROMPT = """You are the cleanup pass on a Gut Linux desktop ({res}). \
The task that was running just ended; an automatic sweep already closed the \
apps, windows and browser tabs it tracked. Look at the screen and the window \
list — if something the task left behind is still open (dialogs, save or \
discard prompts, stray windows), close it. Escape or alt-F4 for dialogs.

Rules:
- Discard unsaved work when asked — deliverables were already sent to the user.
- Never delete files, clear browser data, or log out of anything.
- Leave the panels, wallpaper and desktop icons alone — never click power,
  session, "Log Out" or "Shut Down" controls (including the panel's corner
  button), and never send Ctrl+Alt+Delete. If a logout or shutdown dialog
  is already open, cancel it with Escape or its Cancel button.
- When nothing is left to close — or nothing needed closing — call \
task_complete with a one-line summary of what you closed.
"""


async def janitor_pass(conv_id: str) -> str:
    """Bounded post-sweep tidy with a fresh context — its chatter never
    touches the conversation's model context."""
    # JANITOR_MODEL runs cleanup on a cheaper vision model: swap state.model
    # for the pass so cache_friendly() and the coordinate convention match
    # the model actually answering. cleanup_after_run restores it.
    prev_model = state.model
    if JANITOR_MODEL:
        state.model = JANITOR_MODEL
    try:
        res = "%dx%d" % tuple(pyautogui.size())
    except Exception:
        res = RESOLUTION
    wins = await asyncio.to_thread(list_windows)
    content = [{"type": "text", "text":
                f"Open windows:\n{wins}\n\nClose whatever is left, then call "
                "task_complete."}]
    try:
        shot = screenshot_block(force=True)
    except Exception:
        shot = None
    if shot:
        content.append(shot)
    messages = [
        {"role": "system", "content": JANITOR_PROMPT.format(res=res)},
        {"role": "user", "content": content}]
    try:
        async with httpx.AsyncClient(
                timeout=httpx.Timeout(300, connect=30)) as http:
            for _ in range(JANITOR_MAX_STEPS):
                if state.stop:
                    return "(stopped)"
                prune_images(messages)
                if cache_friendly():
                    apply_cache_control(messages)
                try:
                    r = await llm_request(http, messages,
                                          model=JANITOR_MODEL or None,
                                          tools=JANITOR_TOOLS)
                except httpx.HTTPError as e:
                    return f"model request failed: {e}"
                usd, tin, tout = track_cost(r)
                if conv_id and (usd or tin or tout):
                    conv_add_usage(conv_id, usd, tin, tout)
                await push_cost()

                msg = r.json()["choices"][0]["message"]
                messages.append(assistant_to_dict(msg))
                tool_calls = msg.get("tool_calls") or []
                if not tool_calls:
                    reply = msg.get("content") or ""
                    if not isinstance(reply, str):
                        reply = " ".join(
                            str(b.get("text", "")) for b in reply
                            if isinstance(b, dict))
                    return reply.strip()[:300] or "done"
                acted = False
                for tc in tool_calls:
                    if state.stop:
                        return "(stopped)"
                    fn = tc.get("function") or {}
                    tname = fn.get("name", "")
                    try:
                        targs = json.loads(fn.get("arguments") or "{}")
                    except json.JSONDecodeError:
                        targs = {}
                    await broadcast({"type": "action", "tool": tname,
                                     "args": targs, "agent": JANITOR_NAME})
                    res2, fin = await execute_tool(tname, targs,
                                                   agent=JANITOR_NAME)
                    acted = acted or tname in SCREEN_TOOLS
                    messages.append({"role": "tool",
                                     "tool_call_id": tc.get("id"),
                                     "content": res2})
                    prev = res2 if isinstance(res2, str) else next(
                        (b.get("text", "") for b in res2
                         if b.get("type") == "text"), "")
                    await broadcast({"type": "action_result", "tool": tname,
                                     "result": prev.strip()[:500],
                                     "agent": JANITOR_NAME})
                    if fin:
                        return str(res2 if isinstance(res2, str)
                                   else prev)[:300]
                shot = screenshot_block() if acted else None
                if not tool_result_images_ok():
                    # Same hoist as the main loop — tool-result images
                    # never reach OpenAI-format providers.
                    _hoist_tool_images(messages, shot)
                elif shot is not None:
                    t = messages[-1]
                    c = t["content"]
                    if isinstance(c, str):
                        t["content"] = [{"type": "text", "text": c}, shot]
                    elif not any(b.get("type") == "image_url"
                                 for b in c):
                        c.append(shot)
    except Exception as e:
        return f"janitor error: {e}"
    return f"hit the {JANITOR_MAX_STEPS}-step cap"


async def cleanup_after_run(conv_id: str, baseline: dict | None,
                            janitor: bool = True) -> None:
    """Sweep what the run created, then let the janitor handle residue, then
    sweep once more for anything the janitor itself opened."""
    try:
        if not GUT_CLEANUP:
            return
        stats = (await asyncio.to_thread(sweep_desktop, baseline)
                 if baseline else
                 {"tabs": 0, "windows": 0, "procs": 0, "scratch": 0})
        parts = [f"{stats[k]} {label}" for k, label in (
            ("tabs", "tabs"), ("windows", "windows"),
            ("procs", "processes"), ("scratch", "scratch items"))
            if stats[k]]
        report = ""
        if janitor:
            await broadcast({"type": "status", "state": "cleanup",
                             "model": state.model,
                             "run_started": state.run_start,
                             "conversation_id": conv_id})
            # janitor_pass swaps state.model to JANITOR_MODEL for the pass;
            # restore whatever the conversation was using.
            prev_model = state.model
            report = await janitor_pass(conv_id)
            state.model = prev_model
            if baseline:
                await asyncio.to_thread(sweep_desktop, baseline)
        revived = await asyncio.to_thread(heal_desktop)
        text = "tidy-up"
        if parts:
            text += ": closed " + ", ".join(parts)
        if report:
            text += f" — janitor: {report}"
        if revived:
            text += " — restarted " + ", ".join(revived)
        if parts or report or revived:
            await broadcast({"type": "cleanup", "text": text})
    except Exception as e:
        print(f"[gut] cleanup failed: {e}")
    finally:
        _persist_baseline(None)


async def agent_loop(conv_id: str, task_text: str,
                     attachments: list[dict] | None = None) -> None:
    state.stop = False
    state.paused = False
    state.run_start = time.time()
    state.sent_files = {}
    state.output_nudge_done = False
    state.todo_nudge_done = False
    state.todo_reconcile = False
    state.verify_rejects = 0
    state.verify_caveat = None
    state.escalated = False
    run_usd0 = state.session_usd  # session spend baseline for AGENT_MAX_USD
    # One try wraps setup AND the step loop: a crash anywhere in the run is
    # reported by the except below, then the finally cleans up and drops
    # status to idle. A failed task must never just go quiet — or sit on
    # "running" — with no explanation for the user.
    baseline = None
    messages: list[dict] = []
    try:
        # Snapshot the desktop before the run touches it — the end-of-run
        # sweep only removes what appears after this point, so user-opened
        # windows and tabs survive. Persisted so a crash mid-run still
        # cleans up at boot.
        baseline = await asyncio.to_thread(capture_baseline)
        state.baseline_tabs = set(baseline.get("tabs") or [])
        if GUT_CLEANUP:
            await asyncio.to_thread(_persist_baseline, baseline)
        await broadcast({"type": "status", "state": "running",
                         "model": state.model,
                         "run_started": state.run_start,
                         "conversation_id": conv_id})
        # A revisited conversation resumes its own stored context; a fresh
        # one starts from just the system prompt.
        # The client resizes the display to fit its pane, so RESOLUTION (the
        # Xvfb startup size / max) may be stale — report the live screen size.
        try:
            res = "%dx%d" % tuple(pyautogui.size())
        except Exception:
            res = RESOLUTION
        coords = coord_prompt_for(state.model)
        messages = conv_load_context(conv_id) or [
            {"role": "system", "content": system_prompt(res, coords)}]
        sanitize_context(messages)
        content = [{"type": "text", "text": task_text}]
        # Attached images go inline so the model sees them directly; other
        # files are referenced by path in the task text (attach_note).
        for f in attachments or []:
            if f["mime"].startswith("image/"):
                content.append({"type": "image_url", "image_url": {
                    "url": f"data:{f['mime']};base64,{f['b64']}"}})
        content.append(screenshot_block(force=True))
        messages.append({"role": "user", "content": content})
        done = False
        stall, unchanged_streak = StallDetector(), 0
        idle_replies = 0
        state.todos = conv_load_todos(conv_id)
        for t in state.todos:  # lists stored before kinds existed
            t.setdefault("kind", infer_step_kind(t.get("content", "")))
        state.plan_shared = False
        state.steps_since_todo = 0
        state.evidence = dict.fromkeys(STEP_KINDS, 0)
        state.step_marks = {}
        state.plan_mark = dict.fromkeys(STEP_KINDS, 0)
        state.gate_bounces = 0
        state.spawned = 0
        state.steps_since_compact = 99
        state.url_fail_streak = 0
        state.solo_research = 0
        state.dom_seen = False
        state.tree_seen = False
        state.tools_shown = set()
        last_tin = 0    # prompt_tokens of the last request — compaction signal
        overflow_retries = 0
        async with httpx.AsyncClient(timeout=httpx.Timeout(600, connect=30)) as http:
            for step in range(MAX_STEPS):
                while state.paused and not state.stop:
                    await asyncio.sleep(0.4)
                if state.stop:
                    break
                if done:
                    # A follow-up the user queued while the agent worked
                    # becomes the next turn here — ahead of the end-of-run
                    # cleanup, so the desktop is still as the run left it.
                    if await drain_user_msgs(conv_id, messages, "queue"):
                        done = False
                        continue
                    break

                drain_subagent_inbox(conv_id, messages)
                await drain_user_msgs(conv_id, messages, "steer")

                # Long-horizon: compact history once the last request's
                # billed input approaches the model's context window. The
                # >=3-step gap bounds compaction-call spend if the context
                # stays over the ratio (e.g. one huge tool result).
                state.steps_since_compact += 1
                if last_tin and state.steps_since_compact >= 3:
                    limit = await model_context_limit(http)
                    if last_tin > limit * COMPACT_RATIO:
                        await compact_context(http, conv_id, messages)

                # Checklist nudges on long tasks: bootstrap one when the run
                # is deep with none, or poke a stale one back into view. The
                # note alone got answered in prose ("Del 1… Del 2…") and
                # ignored six times over in one run, so the next request
                # also pins the reply to plan — the checklist gets
                # made, not merely suggested.
                force: str | None = None
                state.steps_since_todo += 1
                undone = [i for i in state.todos
                          if i.get("status") != "done"]
                threshold = (TODO_REMIND_STEPS if state.todos
                             else TODO_NUDGE_STEPS)
                if state.steps_since_todo >= threshold and messages:
                    state.steps_since_todo = 0
                    note = None
                    if undone:
                        note = (
                            f"your checklist hasn't changed in "
                            f"{TODO_REMIND_STEPS} steps — post the updated "
                            "list now with plan (full `steps` list, exactly "
                            "one item in_progress), then carry on.")
                    elif not state.todos:
                        if state.plan_shared:
                            note = (
                                f"{TODO_NUDGE_STEPS} steps in and no "
                                "checklist — post it now with plan (full "
                                "`steps` list, one item in_progress), then "
                                "carry on.")
                        else:
                            note = (
                                f"{TODO_NUDGE_STEPS} steps in and no plan — "
                                "this task has outgrown a quick action. "
                                "Call plan now with `summary` (a short plan "
                                "for the user) and `steps` (full checklist, "
                                "one item in_progress), then carry on.")
                    # All items done: the list is already truthful — don't
                    # spend a pinned call re-posting it.
                    if note:
                        _inject_note(messages, note)
                        force = "plan"

                # A bounced task_complete owes the checklist one update —
                # pin the next reply to plan so ending a run always passes
                # through a final reconcile, not just a suggestion.
                if state.todo_reconcile:
                    state.todo_reconcile = False
                    force = "plan"

                # Step-cap countdown: one heads-up to switch from exploring
                # to delivering, then task_complete pinned on the last step
                # so a capped run still ends with a summary to the user.
                remaining = MAX_STEPS - step
                if WRAP_UP_STEPS > 0 and remaining in (WRAP_UP_STEPS, 1):
                    _inject_note(messages, wrap_up_note(remaining))
                    if remaining == 1:
                        # The last call must land — completion nudges it can
                        # no longer act on must not bounce it, and neither
                        # may the checker: a rejected wrap-up here is no
                        # wrap-up at all, so it's accepted with the caveat.
                        state.output_nudge_done = True
                        state.todo_nudge_done = True
                        state.verify_rejects = max(state.verify_rejects,
                                                   VERIFY_MAX_REJECTS)
                        force = "task_complete"

                prune_images(messages)
                prune_tool_results(messages)
                # A configured cap is a ceiling, not just the trigger above:
                # last_tin saw the *previous* request, so one huge tool
                # result could push this call over it. Estimate the
                # post-prune context and compact before it does (the same
                # 3-step gap bounds compaction spend when a single result
                # alone is bigger than the cap).
                if (model_context_cap(state.model)
                        and state.steps_since_compact >= 3):
                    limit = await model_context_limit(http)
                    if (sum(ctx_breakdown(messages).values())
                            > limit * COMPACT_RATIO):
                        await compact_context(http, conv_id, messages)
                if cache_friendly():
                    apply_cache_control(messages)
                tools = tools_for_run(messages)
                try:
                    r = await llm_request_forcing(http, messages, force, tools)
                except httpx.HTTPError as e:
                    body = getattr(e.response, "text", "") if hasattr(e, "response") else ""
                    # Context-window overflow: compact and retry rather than
                    # killing the run. The ratio check above normally fires
                    # first; this catches unknown-window models and single
                    # tool results big enough to overshoot between checks.
                    if (overflow_retries < 2 and is_context_overflow(e)
                            and await compact_context(http, conv_id,
                                                      messages)):
                        overflow_retries += 1
                        continue
                    # Provider-side rejection (bad images, decommissioned
                    # deployment, …) — escalate rather than kill the run.
                    # Transport errors skip this: a down proxy fails the
                    # same whatever the model.
                    if (isinstance(e, httpx.HTTPStatusError)
                            and await maybe_escalate(
                                messages, conv_id,
                                f"{state.model} request failed")):
                        continue
                    await broadcast({"type": "error",
                                     "text": f"LiteLLM request failed: {e} {body[:300]}"})
                    break

                usd, tin, tout = track_cost(r)
                last_tin = tin
                if tin:
                    # The billed prompt size IS the context occupancy — keep
                    # a proportional category split so the client can meter
                    # the window live (and show it drop when we compact).
                    limit = await model_context_limit(http)
                    raw = ctx_breakdown(messages)
                    scale = tin / max(1, sum(raw.values()))
                    state.ctx_used = tin
                    state.ctx_parts = {
                        k: round(v * scale) for k, v in raw.items()}
                if usd or tin or tout:
                    conv_add_usage(conv_id, usd, tin, tout,
                                   ctx_tokens=tin or None)
                await push_cost()
                if (AGENT_MAX_USD
                        and state.session_usd - run_usd0 >= AGENT_MAX_USD):
                    await broadcast({"type": "error", "text":
                        f"run hit the ${AGENT_MAX_USD:.2f} spend cap "
                        f"(this run: ${state.session_usd - run_usd0:.4f}); "
                        "stopping"})
                    break

                msg = r.json()["choices"][0]["message"]
                messages.append(assistant_to_dict(msg))
                # Reasoning rides an ephemeral channel — the client flashes
                # it in the status line; it never lands in the transcript.
                think = reasoning_text(msg)
                if think:
                    await broadcast({"type": "thinking", "text": think})
                # Model narration goes to the verbose log, not the chat —
                # the user only hears what the agent deliberately sends.
                reply_text = message_text(msg)
                if reply_text:
                    await broadcast({"type": "thought",
                                     "text": reply_text})

                tool_calls = msg.get("tool_calls") or []
                if not tool_calls:
                    idle_replies += 1
                    if idle_replies >= IDLE_REPLY_LIMIT:
                        # The model insists on prose instead of calling
                        # task_complete — end the run and deliver its last
                        # reply as the wrap-up so it reaches the user.
                        done = True
                        await broadcast({"type": "done",
                                         "text": (reply_text or "done")[:8000]})
                        break
                    nudge = ("Continue with tool calls, or call task_complete "
                             "when done.") if idle_replies == 1 else (
                        "Plain text doesn't reach the user and doesn't end "
                        "the task. If you're finished, call task_complete now "
                        "with your answer as `summary`; otherwise make your "
                        "next tool call.")
                    messages.append({"role": "user", "content": nudge})
                    continue
                idle_replies = 0

                acted_on_screen = acted_visibly = False
                parsed: list[tuple[dict, str, dict]] = []
                for tc in tool_calls:
                    fn = tc.get("function") or {}
                    try:
                        args = json.loads(fn.get("arguments") or "{}")
                    except json.JSONDecodeError:
                        args = {}
                    # A resumed context or a model echoing pre-consolidation
                    # names still works — left_click → click, etc.
                    name, args = legacy_call(fn.get("name", ""), args)
                    parsed.append((tc, name, args))

                # A turn of nothing-but PARALLEL_TOOLS calls has no ordering
                # and touches no shared state — run the calls concurrently
                # instead of paying a serial await for each.
                pres: list | None = None
                if len(parsed) > 1 and all(n in PARALLEL_TOOLS
                                           for _, n, _ in parsed):
                    for _, name, args in parsed:
                        await broadcast({"type": "action", "tool": name,
                                         "args": args})

                    async def _par(n, a):
                        sig = stall.signature(n, a)
                        if sig and sig in stall.blocked:
                            return _BLOCKED_MSG, False
                        return await execute_tool(n, a)
                    pres = await asyncio.gather(
                        *(_par(n, a) for _, n, a in parsed))

                # Exactly-one-research-call turns, over and over, is the
                # pattern the nudge exists to break; a batched turn (or a
                # `queries` list) and any non-research call reset it.
                solo_research = (
                    len(parsed) == 1 and parsed[0][1] in PARALLEL_TOOLS
                    and isinstance(parsed[0][2], dict)
                    and not (isinstance(parsed[0][2].get("queries"), list)
                             and len(parsed[0][2]["queries"]) > 1))
                state.solo_research = \
                    state.solo_research + 1 if solo_research else 0

                for i, (tc, name, args) in enumerate(parsed):
                    # Honor takeover/stop between calls, not just between turns.
                    while state.paused and not state.stop:
                        await asyncio.sleep(0.4)
                    if state.stop:
                        break
                    # Back-to-back screen actions in one batch were planned
                    # against the same frame — let the last one's effect
                    # render before the next click lands.
                    if acted_on_screen and name in SCREEN_TOOLS:
                        await asyncio.sleep(INTER_ACTION_DELAY)
                    if pres is None:
                        await broadcast({"type": "action", "tool": name,
                                         "args": args})
                    sig = stall.signature(name, args)
                    if pres is not None:
                        result, finished = pres[i]
                    elif sig and sig in stall.blocked:
                        # A call that already earned a STUCK note is a proven
                        # dead-end — refuse to run it again so the model is
                        # forced to change approach instead of ignoring the note.
                        result, finished = _BLOCKED_MSG, False
                    else:
                        result, finished = await execute_tool(name, args)
                        acted_on_screen = acted_on_screen \
                            or name in SCREEN_TOOLS
                        acted_visibly = acted_visibly \
                            or name in VISUAL_TOOLS

                    # Stall detector: exact repeats within the recent window
                    # and longer cycles of previously-seen actions — see
                    # StallDetector. A spent rescue escalates the model or
                    # checks in with the user once the run has had enough.
                    if sig:
                        note, stuck = stall.observe(name, sig)
                        if note:
                            result = _note(result, note)
                        if stuck:
                            note = await spend_rescue(
                                stall, messages, conv_id,
                                "kept repeating dead-end actions")
                            if note:
                                result = _note(result, note)

                    # A turn that was a single research call extends the
                    # serial-research streak — nudge it toward batching or
                    # delegation once it's clearly a pattern.
                    if solo_research \
                            and state.solo_research >= RESEARCH_NUDGE_STREAK:
                        result = _note(result, RESEARCH_NUDGE_MSG.format(
                            n=state.solo_research))
                        state.solo_research = 0

                    messages.append({
                        "role": "tool",
                        "tool_call_id": tc.get("id"),
                        "content": result,
                    })
                    preview = result if isinstance(result, str) else next(
                        (b.get("text", "") for b in result
                         if b.get("type") == "text"), "")
                    ev = {"type": "action_result", "tool": name,
                          "result": preview.strip()[:500]}
                    urls = _result_urls(name, preview)
                    if urls:
                        ev["urls"] = urls
                    await broadcast(ev)
                    if finished and name == "task_complete":
                        critique = await verify_wrap_up(
                            http, conv_id, str(result), messages)
                        if critique is not None:
                            result, finished = critique, False
                            # The just-appended tool result holds the
                            # rejected summary — swap it for the critique
                            # or the model thinks it completed.
                            messages[-1]["content"] = critique
                            await broadcast({"type": "agent_msg",
                                             "text": critique[:600]})
                    if finished:
                        done = True
                        text = str(result)
                        if state.verify_caveat:
                            # Cap the summary so the caveat can't be
                            # truncated away below — it's the part the
                            # user most needs to see.
                            text = (text[:6300] + "\n\n⚠ Checker flagged "
                                    "claims it couldn't verify in the "
                                    "tool log:\n" + state.verify_caveat)
                        await broadcast({"type": "done",
                                         "text": text[:8000]})
                        break

                # One fresh screenshot per turn on the last tool result —
                # skipped when the frame is byte-identical to the last sent.
                shot = None
                if acted_on_screen and not done and not state.stop:
                    shot = screenshot_block()
                    # An unchanged frame right after an action usually means
                    # the app hasn't repainted yet rather than that nothing
                    # happened — poll for a changed frame before telling the
                    # model the screen is unchanged (otherwise it re-clicks
                    # into a UI that has since moved on).
                    if shot is None and ACTION_SETTLE_SECS > 0:
                        deadline = (asyncio.get_running_loop().time()
                                    + ACTION_SETTLE_SECS)
                        while shot is None and not state.stop:
                            left = deadline - asyncio.get_running_loop().time()
                            if left <= 0:
                                break
                            await asyncio.sleep(min(0.3, left))
                            shot = await asyncio.to_thread(screenshot_block)
                    target = messages[-1]
                    if shot is None and not acted_visibly:
                        # Shell/UNO-only turn: an unchanged screen is the
                        # expected outcome, not a stalled one.
                        target["content"] = _note(target["content"],
                                                  "(screen unchanged)")
                    elif shot is None:
                        unchanged_streak += 1
                        target["content"] = _note(target["content"],
                                                  "(screen unchanged)")
                        if unchanged_streak == 3:
                            target["content"] = _note(target["content"],
                                "WARNING: the screen has not changed across "
                                "your last 3 actions — they are having no "
                                "effect. Switch tactics.")
                        elif unchanged_streak >= 6:
                            target["content"] = _note(
                                target["content"], _unstick_note(
                                    "the screen has not changed across your "
                                    f"last {unchanged_streak} actions — they "
                                    "are having no effect"))
                            unchanged_streak = 0
                            stall.rescues += 1
                            note = await spend_rescue(
                                stall, messages, conv_id,
                                "actions had no visible effect")
                            if note:
                                target["content"] = _note(target["content"],
                                                          note)
                    # A repainted screen resets the no-effect streak only —
                    # not the rescue count: a cycle of dead ends changes the
                    # screen every step and is exactly what the count exists
                    # to escalate on.
                    elif tool_result_images_ok():
                        unchanged_streak = 0
                        content = target["content"]
                        if isinstance(content, str):
                            target["content"] = [
                                {"type": "text", "text": content}, shot]
                        elif not any(b.get("type") == "image_url"
                                     for b in content):
                            content.append(shot)
                        # else: last result already carries a frame (e.g. the
                        # turn ended on the screenshot tool) — don't double up.
                    else:
                        unchanged_streak = 0
                # Providers that can't read images inside tool results get
                # this turn's frames — the auto-shot plus any tool-returned
                # images (screenshot/crop) — as a user message instead.
                if not tool_result_images_ok():
                    _hoist_tool_images(messages, shot)

                # Persist context each step so a follow-up message — or a
                # daemon restart — resumes mid-conversation.
                await asyncio.to_thread(conv_save_context, conv_id, messages)
            else:
                await broadcast({"type": "error",
                                 "text": f"hit step cap ({MAX_STEPS}); stopping"})
    except Exception as e:
        # Whatever killed the run — a malformed model reply, a blown tool,
        # dead storage — the user hears why it stopped instead of watching
        # the agent silently go idle.
        await broadcast({"type": "error", "text": f"agent crashed: {e}"})
    finally:
        # messages is empty when setup died before the context loaded —
        # saving then would clobber the stored conversation with nothing.
        if messages:
            try:
                await asyncio.to_thread(conv_save_context, conv_id, messages)
            except Exception as e:
                print(f"[gut] context save failed: {e}")
        # The sweep runs on every exit path; the janitor only on natural
        # endings — an explicit stop hands the desktop back as-is.
        await cleanup_after_run(conv_id, baseline, janitor=not state.stop)
        await broadcast({"type": "status", "state": "idle",
                         "model": state.model, "conversation_id": None})
        await push_cost()
        # Queued follow-ups outlive the run they waited behind — a stop
        # ends the current task, not what the user lined up after it.
        while state.user_msgs:
            nxt = state.user_msgs.popleft()
            try:
                meta = _read_meta(nxt["conv"])
            except ValueError:
                meta = None
            if meta is None:
                continue  # conversation deleted while the message waited
            state.conversation_id = nxt["conv"]
            state.model = str(meta.get("model") or state.model)
            if nxt.get("seq") is not None:
                await broadcast({"type": "dequeue", "items": [
                    {"conv": nxt["conv"], "seq": nxt["seq"]}]})
            state.task = asyncio.create_task(
                agent_loop(nxt["conv"], nxt["text"], nxt["files"]))
            break


async def handle_client_msg(ws: WebSocket, msg: dict) -> None:
    mtype = msg.get("type")
    if mtype == "task":
        cid = str(msg.get("conversation_id") or "")
        try:
            meta = _read_meta(cid)
        except ValueError:
            meta = None
        if meta is None:
            await ws.send_text(json.dumps(
                {"type": "error",
                 "text": "no such conversation — create one first"}))
            return
        text = str(msg.get("text", ""))
        saved, err = save_attachments(msg.get("files") or [])
        if err:
            await ws.send_text(json.dumps({"type": "error", "text": err}))
            return
        if not text.strip() and not saved:
            return
        event = {"type": "user", "text": text}
        if saved:
            event["files"] = event_files(saved)
        if saved:
            text = f"{text}\n\n{attach_note(saved)}" if text \
                else attach_note(saved)
        if state.running:
            # One agent per desktop environment: a message sent mid-run
            # queues instead of starting. "steer" is injected into the
            # running conversation's context at the next step; anything
            # else is picked up when the run ends — same-conversation
            # entries even before the cleanup sweep.
            steer = bool(msg.get("steer")) and cid == state.conversation_id
            seq = conv_append_event(cid, event)
            await broadcast({**event, "conversation_id": cid, "seq": seq,
                             "queued": "steer" if steer else "queue"})
            state.user_msgs.append({
                "conv": cid, "text": text, "files": saved,
                "mode": "steer" if steer else "queue", "seq": seq})
            return
        state.conversation_id = cid
        state.model = str(meta.get("model") or state.model)
        seq = conv_append_event(cid, event)
        await broadcast({**event, "conversation_id": cid, "seq": seq})
        state.task = asyncio.create_task(agent_loop(cid, text, saved))
    elif mtype == "answer":
        if state.pending_answer is not None and not state.pending_answer.done():
            text = str(msg.get("text", ""))
            saved, err = save_attachments(msg.get("files") or [])
            if err:
                await ws.send_text(json.dumps({"type": "error", "text": err}))
                return
            cid = state.conversation_id
            if cid and state.running:
                event = {"type": "user", "text": text}
                if saved:
                    event["files"] = event_files(saved)
                seq = conv_append_event(cid, event)
                await broadcast({**event, "conversation_id": cid, "seq": seq})
            if saved:
                paths = ", ".join(f["path"] for f in saved)
                text = f"{text}\n[attached files: {paths}]" if text \
                    else f"[attached files: {paths}]"
            state.pending_answer.set_result(text)
    elif mtype == "control":
        action = msg.get("action")
        if action == "pause":
            state.paused = True
        elif action == "resume":
            state.paused = False
        elif action == "stop":
            state.stop = True
            state.paused = False
            if state.pending_answer is not None and not state.pending_answer.done():
                state.pending_answer.set_result("(task stopped by user)")
            # Helpers honor state.stop at their next step; cancel hurries
            # ones blocked in an LLM request or a long run_command.
            for e in state.subagents.values():
                t = e.get("task")
                if t and e.get("status") == "running":
                    t.cancel()
        elif action == "deliver":
            # Promote queued message(s) of the running conversation to
            # steer — they're injected at the next step boundary.
            seq = msg.get("seq")
            for e in state.user_msgs:
                if (e["conv"] == state.conversation_id
                        and (seq is None or e.get("seq") == seq)):
                    e["mode"] = "steer"
        elif action == "dequeue":
            seq = msg.get("seq")
            dropped = [{"conv": e["conv"], "seq": e.get("seq")}
                       for e in state.user_msgs
                       if seq is None or e.get("seq") == seq]
            if dropped:
                seqs = {d["seq"] for d in dropped}
                state.user_msgs = deque(
                    e for e in state.user_msgs
                    if e.get("seq") not in seqs)
                await broadcast({"type": "dequeue", "items": dropped,
                                 "dropped": True})
        await push_status()
    elif mtype == "set_model":
        model = str(msg.get("model") or state.model)
        cid = str(msg.get("conversation_id") or "")
        if cid:
            try:
                conv_set_model(cid, model)
            except ValueError:
                pass
        # Apply live when it targets the running conversation or when the
        # device is idle. A cid-less set_model mid-run is a boot-time
        # default sync from a just-connected client — applying it would
        # swap the live agent's model (and its coordinate convention)
        # underneath it.
        if cid == state.conversation_id or not state.running:
            state.model = model
        await push_status()
    elif mtype == "ping":
        await ws.send_text(json.dumps({"type": "pong"}))


# Shell payload spawned by POST /api/update on gut-bot (deb) installs. It
# must run outside gut-bot.service — the deb's postinst restarts the service
# and its whole cgroup, so the updater goes through systemd-run (or setsid
# on non-systemd hosts). Args: <status-file> <version> <arch> <repo>.
SELF_UPDATE_SCRIPT = """#!/bin/sh
STATUS="$1"; TARGET="$2"; ARCH="$3"; REPO="${4:-valteryde/gut}"
say() {
  err=""
  [ $# -gt 1 ] && err=$(printf '%s' "$2" | tr -dc 'a-zA-Z0-9 ._-')
  printf '{"state":"%s","target":"%s","error":"%s","ts":%s}\\n' "$1" "$TARGET" "$err" "$(date +%s)" > "$STATUS"
}
say downloading
TMP=$(mktemp -d)
cd "$TMP" || { say failed "mktemp failed"; exit 1; }
BASE="https://github.com/$REPO/releases/download/v$TARGET"
DEB="gut-bot_${TARGET}_${ARCH}.deb"
curl -fsSL -o pkg.deb "$BASE/$DEB" || { say failed "download failed"; exit 1; }
if curl -fsSL -o sums.txt "$BASE/SHA256SUMS.txt"; then
  want=$(awk -v f="$DEB" '$2 == f {print $1}' sums.txt)
  if [ -n "$want" ]; then
    printf '%s  %s\\n' "$want" pkg.deb | sha256sum -c - >/dev/null 2>&1 || { say failed "checksum mismatch"; exit 1; }
  fi
fi
say installing
# Noninteractive: a changed conffile (e.g. litellm.yaml) otherwise makes
# dpkg prompt on a dead stdin and leaves the package half-configured with
# the services stopped. confdef/confold keep the installed config; the new
# template lands alongside as .dpkg-new.
sudo -n env DEBIAN_FRONTEND=noninteractive apt-get install -y \
  -o Dpkg::Options::=--force-confdef -o Dpkg::Options::=--force-confold \
  "$TMP/pkg.deb" >/dev/null 2>&1 || { say failed "install failed"; exit 1; }
say done
rm -rf "$TMP"
"""


# ── TLS termination ─────────────────────────────────────────────────────────

def _ensure_tls_cert() -> None:
    """Self-signed cert for the pinned-TLS listener. start.sh already does
    this for docker/gut-bot (websockify needs the files too); this fallback
    covers bare `uvicorn agent_daemon:app` dev runs."""
    global TLS_ERR
    if TLS_CERT.exists() and TLS_KEY.exists():
        return
    cn = re.sub(r"[^A-Za-z0-9._-]", "", DEVICE_NAME) or "device"
    try:
        TLS_CERT.parent.mkdir(parents=True, exist_ok=True)
        subprocess.run(
            ["openssl", "req", "-x509", "-newkey", "ec",
             "-pkeyopt", "ec_paramgen_curve:prime256v1", "-nodes",
             "-days", "3650", "-subj", f"/CN=gut-{cn}",
             "-keyout", str(TLS_KEY), "-out", str(TLS_CERT)],
            check=True, capture_output=True)
        TLS_KEY.chmod(0o600)
    except (OSError, subprocess.CalledProcessError) as e:
        TLS_ERR = f"cert generation failed ({e})"
        print(f"[gut] TLS cert generation failed ({e}) — TLS disabled")


async def _pipe(reader: asyncio.StreamReader,
                writer: asyncio.StreamWriter) -> None:
    try:
        while data := await reader.read(65536):
            writer.write(data)
            await writer.drain()
    except (ConnectionError, OSError, asyncio.IncompleteReadError):
        pass
    finally:
        try:
            writer.close()
        except OSError:
            pass


async def _tls_bridge(client_r: asyncio.StreamReader,
                      client_w: asyncio.StreamWriter) -> None:
    try:
        local_r, local_w = await asyncio.open_connection(
            "127.0.0.1", BACKEND_PORT)
    except OSError:
        client_w.close()
        return
    await asyncio.gather(_pipe(client_r, local_w), _pipe(local_r, client_w))


async def _start_tls_proxy():
    """TLS-terminating forwarder onto the plain uvicorn socket."""
    global TLS_FP, TLS_ERR
    _ensure_tls_cert()
    try:
        der = ssl.PEM_cert_to_DER_cert(TLS_CERT.read_text())
        ctx = ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER)
        ctx.load_cert_chain(TLS_CERT, TLS_KEY)
    except (OSError, ValueError, ssl.SSLError) as e:
        TLS_ERR = f"cert unusable ({e})"
        print(f"[gut] TLS unavailable ({e}) — plain HTTP only")
        return None
    TLS_FP = hashlib.sha256(der).hexdigest()
    try:
        srv = await asyncio.start_server(_tls_bridge, "0.0.0.0", TLS_PORT,
                                         ssl=ctx)
    except OSError as e:
        TLS_ERR = f"cannot bind :{TLS_PORT} ({e})"
        print(f"[gut] TLS listener failed ({e}) — plain HTTP only")
        return None
    print(f"[gut] TLS listener on :{TLS_PORT} (self-signed, pinned by the app)")
    return srv


# ── HTTP API ─────────────────────────────────────────────────────────────────

@asynccontextmanager
async def lifespan(_app: FastAPI):
    try:
        CONV_DIR.mkdir(parents=True, exist_ok=True)
    except OSError as e:
        print(f"[gut] conversation dir {CONV_DIR} unavailable: {e}")
    # A leftover run-state file means the previous run died before its
    # end-of-run cleanup — sweep what it left before taking new work.
    try:
        stale = json.loads(RUN_STATE_FILE.read_text())
    except (OSError, json.JSONDecodeError):
        stale = None
    if stale:
        if stale.get("boot") and stale["boot"] != _boot_key():
            # Whole environment restarted since — nothing to sweep.
            print("[gut] stale run-state from before a restart — discarded")
        elif GUT_CLEANUP:
            try:
                stats = await asyncio.to_thread(sweep_desktop, stale)
                print(f"[gut] swept leftovers of interrupted run: {stats}")
                revived = await asyncio.to_thread(heal_desktop)
                if revived:
                    print(f"[gut] restarted desktop components: {revived}")
            except Exception as e:
                print(f"[gut] startup sweep failed: {e}")
        _persist_baseline(None)
    tls_srv = await _start_tls_proxy()
    # Port multiplexers (see tcpmux): uvicorn and plain websockify sit on
    # loopback internals; the public ports carry plain AND TLS so encrypted
    # transport works wherever the plain ports already reach.
    mux_srvs = []
    if MUX:
        if tcpmux is None:
            # uvicorn is parked on loopback — without the mux nothing can
            # reach the API at all, so die loudly instead of limping.
            raise RuntimeError("tcpmux.py missing but GUT_UVICORN_PORT is "
                               "set — broken install")
        try:
            mux_srvs = [
                await tcpmux.start_mux(HTTP_PORT, TLS_PORT, BACKEND_PORT),
                await tcpmux.start_mux(NOVNC_PORT, NOVNC_TLS_PORT,
                                       NOVNC_PLAIN_PORT)]
        except OSError as e:
            print(f"[gut] port mux failed ({e}) — public ports unproxied")
            mux_srvs = []
    state.litellm_key = await provision_key()
    sync_task = asyncio.create_task(model_sync_loop())
    print(f"[gut] agent ready, model={state.model}")
    try:
        yield
    finally:
        sync_task.cancel()
        if tls_srv:
            tls_srv.close()
        for srv in mux_srvs:
            srv.close()


app = FastAPI(lifespan=lifespan)
app.add_middleware(CORSMiddleware, allow_origins=["*"],
                   allow_methods=["*"], allow_headers=["*"])


def _token_ok(token: str) -> bool:
    return bool(GUT_API_TOKEN) and token == GUT_API_TOKEN


@app.middleware("http")
async def require_token(request: Request, call_next):
    # /api/version and /api/hello stay open — the Electron app probes them
    # to detect a Gut backend and to pair TLS before it has credentials.
    # OPTIONS is a CORS preflight.
    if (not GUT_API_TOKEN or request.method == "OPTIONS"
            or request.url.path in ("/api/version", "/api/hello")):
        return await call_next(request)
    auth = request.headers.get("authorization", "")
    if (auth == f"Bearer {GUT_API_TOKEN}"
            or _token_ok(request.query_params.get("token", ""))):
        return await call_next(request)
    return JSONResponse({"detail": "unauthorized"}, status_code=401)


@app.websocket("/ws/chat")
async def chat_ws(ws: WebSocket):
    if GUT_API_TOKEN:
        auth = ws.headers.get("authorization", "")
        tok = ws.query_params.get("token", "")
        if auth != f"Bearer {GUT_API_TOKEN}" and tok != GUT_API_TOKEN:
            await ws.close(code=4401)
            return
    await ws.accept()
    state.clients.add(ws)
    await ws.send_text(json.dumps({
        "type": "hello", "state": state.phase, "model": state.model,
        "device": DEVICE_NAME,
        "run_started": state.run_start if state.running else None,
        "running_conversation":
            state.conversation_id if state.running else None,
        # Queued mid-run messages — lets a rejoining client restore the
        # "queued" badges on transcript events it fetches afterwards.
        "queue": [{"conv": e["conv"], "seq": e.get("seq"),
                   "mode": e["mode"]} for e in state.user_msgs],
    }))
    await push_cost()
    try:
        while True:
            try:
                await handle_client_msg(
                    ws, json.loads(await ws.receive_text()))
            except (WebSocketDisconnect, json.JSONDecodeError, RuntimeError):
                raise
            except Exception as e:
                # A handler bug must not kill the socket and swallow the
                # user's message in silence — report it back to them.
                try:
                    await ws.send_text(json.dumps(
                        {"type": "error",
                         "text": f"internal error: {e}"}))
                except Exception:
                    raise WebSocketDisconnect()
    # Abrupt TCP drops (tab reload, network blip) surface as RuntimeError
    # ("WebSocket is not connected") rather than WebSocketDisconnect — the
    # peer is gone either way.
    except (WebSocketDisconnect, json.JSONDecodeError, RuntimeError):
        pass
    finally:
        state.clients.discard(ws)


@app.get("/api/version")
async def api_version():
    """Unauthenticated handshake — the client probes this to detect a Gut
    backend and check compatibility before presenting credentials."""
    return {"version": GUT_VERSION, "device": DEVICE_NAME,
            "auth": bool(GUT_API_TOKEN), "install": INSTALL_KIND}


@app.get("/api/hello")
async def api_hello(n: str = ""):
    """Unauthenticated TLS pairing handshake, like /api/version.

    The app reads the pinned listener's cert fingerprint here (over plain
    HTTP is fine — it leaks nothing), opens a provisional TLS connection to
    see the cert actually presented, and pins it when the fingerprint
    matches and `mac` verifies: HMAC(device password, fp + client nonce).
    A MITM can't forge the MAC without the password, and a pure relay only
    ever forwards the real cert — so first connect has no trust gap. `n`
    makes every answer single-use so replays can't pin stale certs.
    """
    if not TLS_FP:
        raise HTTPException(
            404, TLS_ERR or "this backend has no TLS listener")
    mac = None
    if GUT_API_TOKEN:
        mac = hmac.new(GUT_API_TOKEN.encode(),
                       f"gut-tls-pin:{TLS_FP}:{n}".encode(),
                       hashlib.sha256).hexdigest()
    # With the mux running, TLS rides the same public ports as plain HTTP —
    # advertise those so pairing works wherever :8000/:6080 already reach.
    return {"cert_sha256": TLS_FP, "mac": mac,
            "port": HTTP_PORT if MUX else TLS_PORT,
            "vnc_port": NOVNC_PORT if MUX else NOVNC_TLS_PORT}


def _update_status() -> dict:
    try:
        s = json.loads(UPDATE_STATUS_FILE.read_text())
        if isinstance(s, dict):
            return s
    except (OSError, json.JSONDecodeError):
        pass
    return {"state": "idle"}


@app.get("/api/update")
async def api_update_status():
    return {"supported": INSTALL_KIND == "deb", "install": INSTALL_KIND,
            **_update_status()}


@app.post("/api/update")
async def api_update(body: dict = Body(default={})):
    """Self-update for gut-bot deb installs: download the release deb and
    install it detached — the package's postinst restarts this service, so
    the updater must live outside the service's cgroup."""
    if INSTALL_KIND != "deb":
        raise HTTPException(400, "self-update needs a gut-bot install "
                                 f"(this backend is '{INSTALL_KIND}')")
    target = str(body.get("version") or "")
    if not re.fullmatch(r"\d+\.\d+\.\d+", target):
        raise HTTPException(400, "version must look like 1.2.3")
    if state.running:
        raise HTTPException(409, "agent is working — stop it first")
    if _update_status().get("state") in ("starting", "downloading",
                                         "installing"):
        raise HTTPException(409, "update already in progress")
    try:
        arch = subprocess.check_output(
            ["dpkg", "--print-architecture"], text=True).strip()
    except (OSError, subprocess.CalledProcessError):
        arch = ""
    if arch not in ("amd64", "arm64"):
        raise HTTPException(400, f"no gut-bot build for arch '{arch or '?'}'")

    script = GUT_DATA_DIR / "self-update.sh"
    script.write_text(SELF_UPDATE_SCRIPT)
    script.chmod(0o755)
    cmd = ["/bin/sh", str(script), str(UPDATE_STATUS_FILE), target,
           arch, GITHUB_REPO]
    UPDATE_STATUS_FILE.write_text(json.dumps(
        {"state": "starting", "target": target, "ts": int(time.time())}))
    launched = False
    if Path("/run/systemd/system").exists():
        launched = subprocess.run(
            ["sudo", "-n", "systemd-run", "-q", "--collect",
             "--unit=gut-self-update", *cmd],
            capture_output=True).returncode == 0
    if not launched:
        subprocess.Popen(["setsid", *cmd], stdin=subprocess.DEVNULL,
                         stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    return {"ok": True, "target": target}


def _keys_state() -> dict:
    eff = effective_provider_keys()
    return {k: {"set": k in eff,
                "source": eff[k][1] if k in eff else None}
            for k in PROVIDER_KEYS}


@app.get("/api/keys")
async def api_keys():
    """Which provider keys this device can serve models for — set flags and
    where each came from, never the values."""
    return _keys_state()


@app.post("/api/keys")
async def api_keys_set(body: dict = Body(...)):
    """Upsert provider keys pushed from the app; an empty value removes.

    Keys are stored on this device (PROVIDER_KEYS_FILE, next to the
    conversation store) and synced into LiteLLM's model store — no restart
    needed. Unknown key names are rejected so this can't write arbitrary
    environment.
    """
    updates = body.get("keys")
    if not isinstance(updates, dict):
        raise HTTPException(400, 'expected {"keys": {NAME: value}}')
    bad = sorted(k for k in updates if k not in PROVIDER_KEYS)
    if bad:
        raise HTTPException(400, "unknown provider keys: " + ", ".join(bad))
    saved = load_provider_keys()
    for k, v in updates.items():
        v = str(v or "").strip()
        # Server providers take an address, not a secret — store it
        # normalized so the UI and the reconciler read the same canonical
        # value ('192.168.1.5' → 'http://192.168.1.5:11434').
        if v and k in SERVER_BASE_KEYS:
            norm, hint = SERVER_BASE_KEYS[k]
            v = norm(v)
            if not v:
                raise HTTPException(400, f"{k} wants a host — {hint}")
        if v:
            saved[k] = v
        else:
            saved.pop(k, None)
    store_provider_keys(saved)
    state.models_synced = False  # the sync loop retries on failure
    try:
        await reconcile_models()
        state.models_synced = True
    except Exception as e:
        return {"keys": _keys_state(), "applied": False,
                "error": f"saved on device but LiteLLM rejected it "
                         f"(retrying): {e}"}
    return {"keys": _keys_state(), "applied": True}


@app.get("/api/ollama")
async def api_ollama(base: str = ""):
    """Live probe of the Ollama server — ?base= tests an unsaved address.

    The settings card calls this to answer "does this AI work?": whether the
    server answers from this device, which models it has pulled, and which
    of them can see (the agent is blind on a text-only model).
    """
    configured = ollama_base()
    if base.strip():
        target = normalize_ollama_base(base)
        if not target:
            return {"set": bool(configured), "base": base.strip(),
                    "reachable": False, "models": [],
                    "error": "doesn't look like a host:port"}
    else:
        target = configured
    if not target:
        return {"set": False, "base": "", "reachable": False,
                "models": [], "error": "no server configured"}
    try:
        models = await ollama_tags(target)
    except Exception as e:
        return {"set": bool(configured), "base": target,
                "reachable": False, "models": [],
                "error": _probe_err(e, "an Ollama server")}
    # Reachable but the managed deployments drifted (server was down at the
    # last sync, or models were pulled since) — let the sync loop heal it.
    if target == configured:
        try:
            deployed = {str(m.get("model_name"))
                        for m in await litellm_deployments()}
            live = {f"ollama/{m['name']}" for m in models}
            if live != {n for n in deployed if n.startswith("ollama/")}:
                state.models_synced = False
        except Exception:
            pass
    return {"set": bool(configured), "base": target, "reachable": True,
            "models": models}


@app.post("/api/compat")
async def api_compat(body: dict = Body(default={})):
    """Live probe of an OpenAI-compatible server (vLLM, LM Studio, …).

    {"base": "…"} tests an unsaved address; {"key": "…"} overrides the
    stored key for the probe ('' = anonymous, absent = use the stored one).
    """
    configured = compat_base()
    raw = str(body.get("base") or "")
    if raw.strip():
        target = normalize_compat_base(raw)
        if not target:
            return {"set": bool(configured), "base": raw.strip(),
                    "reachable": False, "models": [],
                    "error": "doesn't look like a host:port"}
    else:
        target = configured
    if not target:
        return {"set": False, "base": "", "reachable": False,
                "models": [], "error": "no server configured"}
    key = (str(body["key"]).strip() if "key" in body else compat_key())
    try:
        models = await compat_models(target, key)
    except Exception as e:
        return {"set": bool(configured), "base": target,
                "reachable": False, "models": [],
                "error": _probe_err(e, "an OpenAI-compatible server",
                                    auth_hint=True)}
    if target == configured:
        try:
            deployed = {str(m.get("model_name"))
                        for m in await litellm_deployments()}
            live = {f"compat/{m['name']}" for m in models}
            if live != {n for n in deployed if n.startswith("compat/")}:
                state.models_synced = False
        except Exception:
            pass
    return {"set": bool(configured), "base": target, "reachable": True,
            "models": models}


def _config_effective() -> dict:
    """Live value of every whitelisted knob — the module global when one
    exists, else the raw env/file string."""
    out = {}
    for k in sorted(CONFIG_KEYS):
        g = globals().get(CONFIG_GLOBALS.get(k, k))
        out[k] = g if g is not None else os.environ.get(k, "")
    return out


@app.get("/api/config")
async def api_config():
    """Non-secret device config: what's stored in config.json (the managed
    store) and the live effective values. Never returns secrets — only the
    whitelisted CONFIG_KEYS can exist here."""
    return {"config": load_config_file(), "effective": _config_effective(),
            "restart_required": sorted(CONFIG_RESTART)}


@app.post("/api/config")
async def api_config_set(body: dict = Body(...)):
    """Upsert config pushed from the app; an empty value removes the key.
    Values hot-apply to the module globals where possible; keys in
    CONFIG_RESTART persist to the file but only take effect at next boot
    (they're consumed by start.sh)."""
    updates = body.get("config")
    if not isinstance(updates, dict):
        raise HTTPException(400, 'expected {"config": {NAME: value}}')
    bad = sorted(k for k in updates if k not in CONFIG_KEYS)
    if bad:
        raise HTTPException(400, "unknown config keys: " + ", ".join(bad))
    saved = load_config_file()
    applied, restart = [], []
    for k, v in updates.items():
        v = str(v or "").strip()
        if v:
            saved[k] = v
            os.environ[k] = v
        else:
            saved.pop(k, None)
            orig = _ENV_ORIG.get(k)
            if orig is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = orig
        gname = CONFIG_GLOBALS.get(k, k)
        if k in CONFIG_RESTART or gname not in globals():
            restart.append(k)
            continue
        try:
            # On removal, re-coerce the restored env value — if neither file
            # nor env holds one, non-string globals can't express the code
            # default and stay as-is until the next boot.
            globals()[gname] = _coerce_like(
                globals()[gname], v or os.environ.get(k, ""))
            applied.append(k)
            # DEFAULT_MODEL seeds state.model at boot — apply it there too
            # or the push would only matter after a restart.
            if k == "DEFAULT_MODEL" and v:
                state.model = v
        except (ValueError, TypeError):
            restart.append(k)
    if "MODEL_CONTEXT_LIMITS" in updates:
        state.ctx_limit.clear()  # limits re-resolve under the new caps
    store_config_file(saved)
    return {"config": saved, "effective": _config_effective(),
            "applied": applied, "restart_required": restart}


@app.get("/api/device")
async def api_device():
    return {"name": DEVICE_NAME}


@app.get("/api/status")
async def api_status():
    return {"state": state.phase, "model": state.model,
            "device": DEVICE_NAME,
            "conversation_id": state.conversation_id if state.running else None,
            "session_usd": state.session_usd,
            "tokens_in": state.tokens_in, "tokens_out": state.tokens_out}


@app.get("/api/conversations")
async def api_conversations():
    metas = conv_list()
    for m in metas:
        m["running"] = bool(
            state.running and state.conversation_id == m.get("id"))
    return metas


@app.post("/api/conversations", status_code=201)
async def api_conversation_create(body: dict = Body(default={})):
    meta = conv_create(str(body.get("model") or state.model),
                       str(body.get("title") or ""))
    await broadcast({"type": "conversations"})
    return meta


@app.get("/api/conversations/{cid}")
async def api_conversation(cid: str):
    try:
        conv = conv_get(cid)
    except ValueError:
        conv = None
    if conv is None:
        raise HTTPException(404, "no such conversation")
    conv["meta"]["running"] = bool(
        state.running and state.conversation_id == cid)
    conv["ctx"] = conv_ctx_estimate(cid, conv["meta"])
    return conv


@app.delete("/api/conversations/{cid}")
async def api_conversation_delete(cid: str):
    if state.running and state.conversation_id == cid:
        raise HTTPException(409, "conversation is running — stop it first")
    try:
        ok = conv_delete(cid)
    except ValueError:
        ok = False
    if not ok:
        raise HTTPException(404, "no such conversation")
    if state.conversation_id == cid:
        state.conversation_id = None
    await broadcast({"type": "conversations"})
    return {"ok": True}


@app.get("/api/cost")
async def api_cost():
    return {"session_usd": state.session_usd,
            "lifetime_usd": await litellm_spend(),
            "tokens_in": state.tokens_in, "tokens_out": state.tokens_out,
            "model": state.model}


@app.get("/api/models")
async def api_models():
    """Model list enriched with LiteLLM's /model/info metadata (price per
    token, context size, vision support) so the client can annotate the
    picker. Falls back to bare ids when info is unavailable."""
    try:
        async with httpx.AsyncClient(timeout=15) as c:
            r = await c.get(f"{LITELLM_URL}/model/info",
                            headers={"Authorization": f"Bearer {state.litellm_key}"})
            r.raise_for_status()
            data = r.json()
            items = data.get("data", []) if isinstance(data, dict) else data
            models = []
            for m in items:
                name = m.get("model_name")
                if not name:
                    continue
                info = m.get("model_info") or {}
                ctx = info.get("max_input_tokens") or info.get("max_tokens")
                # Report the usable window, not the raw one — a configured
                # cap is what the agent will actually stay under.
                cap = model_context_cap(name)
                if cap:
                    try:
                        ctx = min(int(ctx or COMPACT_CONTEXT_LIMIT), cap)
                    except (TypeError, ValueError):
                        pass  # a malformed ctx passes through as-is
                models.append({
                    "id": name,
                    "cost_in": info.get("input_cost_per_token"),
                    "cost_out": info.get("output_cost_per_token"),
                    "ctx": ctx,
                    "vision": info.get("supports_vision"),
                })
            if models:
                return models
    except Exception:
        pass
    try:
        async with httpx.AsyncClient(timeout=15) as c:
            r = await c.get(f"{LITELLM_URL}/v1/models",
                            headers={"Authorization": f"Bearer {state.litellm_key}"})
            r.raise_for_status()
            return [{"id": m["id"]} for m in r.json().get("data", [])]
    except Exception:
        return [{"id": state.model}]


@app.post("/api/model")
async def api_model(body: dict = Body(...)):
    state.model = str(body.get("model", state.model))
    await push_status()
    return {"model": state.model}


@app.post("/api/pause")
async def api_pause():
    state.paused = True
    await push_status()
    return {"state": state.phase}


@app.post("/api/resume")
async def api_resume():
    state.paused = False
    await push_status()
    return {"state": state.phase}
